"""The CLIENT's own BOQ reference number of a BOQ line, for the client-facing DPR ("BOQ Ref (client)").

Why this is not just the item code: the client numbers a PARENT item (8.1.2, 10.2, 39) and writes its lines as a, b, c / i, ii, iii /
"b) 100 mm dia" with no number of their own (pipes, fans, cable tray...). The reference of such a line is "<parent> (<label>)", e.g.
"10.2 (iii)" -- read from the client's own BOQ sheet (the saved boq.xlsx), using the line's `source_sheet` + `src_row` that the BOQ upload
already stores. Narrow deterministic rules only (a wrong reference on a client document is worse than "—"):

  * a number-like reference (1.1.4, 13.2.1, 3.03.5, 39) is used as it is; if that row's text ends with " a) ..." and the next row starts with
    "b) " the row is the first lettered line of that parent -> "39 (a)";
  * a line whose reference is a bare label (a-z, i/ii/iii...) or has none but its text starts with "x) " takes the nearest NUMBERED row above
    as its parent -- only if that row is a heading (no unit, no quantity) or a first lettered line itself -> "<parent> (<label>)";
  * anything else -> None (shown as "—").

Only ProjectBase-shaped sheets (they carry a "Client Ref No" column) are read this way. For any other BOQ the item's own code is the
reference when the BOQ gave one, else None. Nothing here is guessed; the sheet is only ever read, never changed."""
import functools
import re
from pathlib import Path

import pandas as pd

from . import boq as boq_mod
from . import schema

_NUM = re.compile(r"^\d+(?:\.\d+)*$")
_LABEL = re.compile(r"^(?:[a-z]|[ivx]+)$")
_FIRST_SUB = re.compile(r"(?:^|\s)a\)\s[^)]*$")
_SUB_START = re.compile(r"^\s*([a-z])\)\s")
_SUB_B = re.compile(r"^\s*b\)\s")


def _s(v):
    return "" if v is None or (isinstance(v, float) and v != v) else str(v).strip()


def _columns(raw):
    """Header-driven column positions of a ProjectBase sheet, or None."""
    hr = boq_mod._pb_header_row(raw)
    if hr is None:
        return None
    hdr = [schema.norm(v) for v in raw.iloc[hr].tolist()]
    def idx(name):
        return hdr.index(name) if name in hdr else None
    c = {"code": idx("CLIENT REF NO"), "desc": idx("BOQ"), "unit": idx("BOQ UNIT"), "qty": idx("DESIGN QUANTITY")}
    return c if None not in (c["code"], c["desc"], c["unit"], c["qty"]) else None


def derive(rows, cols, src_row):
    """(ref, parent_text, own_text) of the line at 1-based `src_row`. `rows` = {row_number: [cell text, ...]}. Everything None when the
    reference is not known. parent_text = the client's text of the parent item (shown grey above its lettered lines); own_text = the line's
    own wording when its row also carries the parent's text ("...pipes. a) 150 mm dia" -> "a) 150 mm dia")."""
    row = rows.get(src_row)
    if not row:
        return None, None, None
    cc, cd, cu, cq = cols["code"], cols["desc"], cols["unit"], cols["qty"]
    at = lambda r, c: (rows.get(r) or [])[c] if r in rows and c < len(rows[r]) else ""
    ref, desc = at(src_row, cc), at(src_row, cd)
    if _NUM.match(ref):
        nxt = at(src_row + 1, cd)
        m = _FIRST_SUB.search(desc)
        if m and _SUB_B.match(nxt):                                  # "...pipes. a) 150 mm dia" followed by "b) 100 mm dia"
            cut = desc.rfind("a) ")
            return f"{ref} (a)", desc[:cut].strip(), desc[cut:].strip()
        return ref, None, None
    m = _SUB_START.match(desc)
    label = ref if ref and _LABEL.match(ref) else (m.group(1) if (not ref and m) else None)
    if not label:
        return None, None, None
    r = src_row - 1
    while r >= 1:                                                    # the nearest numbered row above = the parent
        if r in rows and _NUM.match(at(r, cc)):
            pref, pdesc = at(r, cc), at(r, cd)
            heading = (at(r, cu) == "" and at(r, cq) == "")
            first_sub = bool(_FIRST_SUB.search(pdesc)) and bool(_SUB_B.match(at(r + 1, cd)))
            if not (heading or first_sub):
                return None, None, None
            if first_sub:
                pdesc = pdesc[:pdesc.rfind("a) ")].strip()
            return f"{pref} ({label})", pdesc, None
        r -= 1
    return None, None, None


@functools.lru_cache(maxsize=4)
def _load(path_str, mtime_ns):
    """{sheet: (cols or None, {row_number: [cell text...]})} for every sheet of boq.xlsx (read once per file version)."""
    import openpyxl
    out = {}
    wb = openpyxl.load_workbook(path_str, read_only=True, data_only=True)
    try:
        for name in wb.sheetnames:
            rows = {i: [_s(v) for v in row] for i, row in enumerate(wb[name].iter_rows(values_only=True), start=1)}
            raw = pd.DataFrame([rows[i] for i in sorted(rows)][:8]) if rows else pd.DataFrame()
            cols = _columns(raw) if len(raw) else None
            out[name] = (cols, rows)
    finally:
        wb.close()
    return out


def refs_for(project_dir, items):
    """{(service, item_code): {"ref", "parent", "text"}} for every row of the BOQ frame `items` that has a known reference. A missing/unreadable
    boq.xlsx, or a BOQ that was not read from sheet rows, simply gives the fallback (the BOQ's own code) or nothing -- never an error."""
    x = Path(project_dir) / "boq.xlsx"
    sheets = {}
    if x.exists():
        try:
            sheets = _load(str(x), x.stat().st_mtime_ns)
        except Exception:
            sheets = {}
    out = {}
    has_rows = "src_row" in items.columns and "source_sheet" in items.columns
    for t in items.itertuples():
        key = (t.service, str(t.item_code))
        ref = parent = text = None
        sr = getattr(t, "src_row", None) if has_rows else None
        sh = getattr(t, "source_sheet", None) if has_rows else None
        if sh in sheets and sheets[sh][0] and sr is not None and not pd.isna(sr):
            ref, parent, text = derive(sheets[sh][1], sheets[sh][0], int(sr))
        else:                                                         # another kind of BOQ: its own code, if it gave one
            raw = _s(getattr(t, "item_code_raw", None))
            if raw and "#" not in raw and raw.lower() != "nan":
                ref = raw
        if ref:
            out[key] = {"ref": ref, "parent": parent or None, "text": text or None}
    return out
