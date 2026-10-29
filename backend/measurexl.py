"""Excel side of the Measurement grid: build the workbook the grid exports, and read one back for the import preview.

Nothing here touches project data. The routes in siteprogress.py decide what a number in the file may do (locks, plan, whole numbers) and the
grid saves it through the same route typing uses, so an import can never do what typing a number cannot.

Layout of the "Measurement" sheet (the importer relies on it, so it is fixed):
    row 1  title                      row 2  how to use it
    row 3  hidden key row: "#KEY" in column A, and the id of each place above its column
    row 4  level / floor names        row 5  zone / room names
    row 6+ one BOQ line per row: A = code (text), B = description, C = unit, then one column per place, then BOQ qty / Site plan / To date / Pending
A hidden sheet "_meta" says which project and service the file was made for, so a file of another service is refused.
"""
import io
import math
import re
import datetime as dt

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

SHEET = "Measurement"
META = "_meta"
KEY_ROW, HEAD1, HEAD2, FIRST_ROW, FIRST_COL = 3, 4, 5, 6, 4
MAGIC = "measurement-grid v1"
MAX_BYTES = 8 * 1024 * 1024
MAX_PLACES, MAX_LINES = 600, 5000

_GREY = PatternFill("solid", fgColor="EEECE6")
_LOCK = PatternFill("solid", fgColor="E4E1D8")
_HEAD = PatternFill("solid", fgColor="F3F2EE")
_WARN = PatternFill("solid", fgColor="FBF1DA")
_thin = Side(style="thin", color="D9D6CD")
_BOX = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
NUM = "#,##0.##"


def _place_label(loc, name):
    return f"{loc} ▸ {name}" if loc and loc != name else name


# ------------------------------------------------------------------ export
CLIENT_SHEET = "Measurement sheet"


def build(project, service, slug, columns, grid, when=None, runs=None):
    """columns: [{"id","loc","name"}] in sheet order; grid: the measure-grid payload (rows / cells / locks).
    runs: [{"code","level","place","from","to","start","end","qty","remarks","kind"}] -- the From-To lines of the length items, for the vertical
    "Measurement sheet" (the client's format: item by item, place by place). Returns the .xlsx bytes.

    Lines and places with nothing recorded are HIDDEN (grouped, so Excel's + shows them again), never removed: the file still has every line
    and place, so it imports exactly as before. When nothing at all is recorded yet, nothing is hidden."""
    if len(columns) > MAX_PLACES:
        raise ValueError(f"this project has {len(columns)} places; the Excel sheet holds up to {MAX_PLACES}")
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    lastc = FIRST_COL + len(columns) - 1
    c_boq, c_plan, c_done, c_pend = lastc + 1, lastc + 2, lastc + 3, lastc + 4
    ws["A1"] = f"{project} · Measurement · {service}"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("Type the quantity recorded to date in each place (a total, not an addition). Grey boxes are not in that place; darker boxes are locked "
                "(their number comes from measurement entries or an area total) and are skipped on import. Leave a box empty to change nothing, type 0 to clear it.")
    ws["A2"].alignment = Alignment(wrap_text=False)
    ws["A2"].font = Font(italic=True, color="5F5E59")
    ws.cell(KEY_ROW, 1, "#KEY")
    for i, c in enumerate(columns):
        ws.cell(KEY_ROW, FIRST_COL + i, str(c["id"]))
    ws.row_dimensions[KEY_ROW].hidden = True
    # level headers, merged across their places
    i = 0
    while i < len(columns):
        j = i
        while j + 1 < len(columns) and columns[j + 1]["loc"] == columns[i]["loc"]:
            j += 1
        cell = ws.cell(HEAD1, FIRST_COL + i, columns[i]["loc"] or "")
        cell.font, cell.alignment, cell.fill = Font(bold=True), Alignment(horizontal="center"), _HEAD
        if j > i:
            ws.merge_cells(start_row=HEAD1, start_column=FIRST_COL + i, end_row=HEAD1, end_column=FIRST_COL + j)
        i = j + 1
    for h, text in ((1, "Sr"), (2, "BOQ line"), (3, "Unit")):
        ws.cell(HEAD2, h, text)
    for i, c in enumerate(columns):
        ws.cell(HEAD2, FIRST_COL + i, c["name"])
    for k, text in ((c_boq, "BOQ qty"), (c_plan, "Site plan"), (c_done, "To date"), (c_pend, "Pending")):
        ws.cell(HEAD2, k, text)
    for col in range(1, c_pend + 1):
        cell = ws.cell(HEAD2, col)
        cell.font, cell.fill, cell.border = Font(bold=True), _HEAD, _BOX
        cell.alignment = Alignment(horizontal="center" if col >= FIRST_COL else "left", wrap_text=True)
    cells, locks = grid["cells"], grid["locks"]
    r = FIRST_ROW
    for row in grid["rows"]:
        code = row["code"]
        a = ws.cell(r, 1, str(code))
        a.number_format = "@"
        ws.cell(r, 2, row["desc"])
        ws.cell(r, 3, row["unit"] or "")
        for i, c in enumerate(columns):
            cell = ws.cell(r, FIRST_COL + i)
            key = f"{code}|{c['id']}"
            if key not in cells:
                cell.fill = _GREY
            else:
                cell.value = grid.get("exprs", {}).get(key, cells[key])          # a box with a sum keeps it as the formula (=112+155)
                cell.number_format = NUM
                if key in locks:
                    cell.fill = _LOCK
                    cell.comment = Comment(locks[key]["why"], "Material Intelligence")
        first, last = get_column_letter(FIRST_COL), get_column_letter(lastc)
        if row["boq_qty"] is not None:
            ws.cell(r, c_boq, row["boq_qty"])
        ws.cell(r, c_plan, row["plan"])
        if row["plan_above_boq"]:
            ws.cell(r, c_plan).fill = _WARN
        ws.cell(r, c_done, f"=SUM({first}{r}:{last}{r})")
        boq = f"{get_column_letter(c_boq)}{r}"
        ws.cell(r, c_pend, f'=IF(ISNUMBER({boq}),MAX({boq}-{get_column_letter(c_done)}{r},0),"—")')
        for col in (c_boq, c_plan, c_done, c_pend):
            ws.cell(r, col).number_format = NUM
            ws.cell(r, col).alignment = Alignment(horizontal="right")
        r += 1
    hidden_rows = hidden_cols = 0
    live_rows = {i for i, row in enumerate(grid["rows"]) if any((cells.get(f"{row['code']}|{c['id']}") or 0) > 0 for c in columns)}
    if live_rows:
        for i, row in enumerate(grid["rows"]):
            if i not in live_rows:
                ws.row_dimensions[FIRST_ROW + i].outlineLevel = 1
                ws.row_dimensions[FIRST_ROW + i].hidden = True
                hidden_rows += 1
        for k, c in enumerate(columns):
            if not any((cells.get(f"{row['code']}|{c['id']}") or 0) > 0 for row in grid["rows"]):
                cd = ws.column_dimensions[get_column_letter(FIRST_COL + k)]
                cd.outlineLevel = 1
                cd.hidden = True
                hidden_cols += 1
    ws.freeze_panes = ws.cell(FIRST_ROW, FIRST_COL)
    ws.column_dimensions["A"].width = 9
    ws.column_dimensions["B"].width = 46
    ws.column_dimensions["C"].width = 8
    for col in range(FIRST_COL, c_pend + 1):
        ws.column_dimensions[get_column_letter(col)].width = 12
    # Summary sheet, with the sign-off box
    sm = wb.create_sheet("Summary")
    sm["A1"] = f"{project} · Measurement summary"
    sm["A1"].font = Font(bold=True, size=13)
    sm["A3"], sm["B3"] = "Service", service
    sm["A4"], sm["B4"] = "Exported on", (when or dt.datetime.now()).strftime("%d %b %Y %H:%M")
    sm["A5"], sm["B5"] = "BOQ lines", len(grid["rows"])
    sm["A6"], sm["B6"] = "Places", len(columns)
    sm["A8"] = "Quantities are the totals recorded to date in Site Progress at the time of export."
    sm["A8"].font = Font(italic=True, color="5F5E59")
    if hidden_rows or hidden_cols:
        sm["A9"] = (f"Only the lines and places with a recorded quantity are shown ({hidden_rows} lines and {hidden_cols} places are hidden). "
                    "Click the + in the margin of the Measurement sheet, or select all and Unhide, to fill the others.")
        sm["A9"].font = Font(italic=True, color="5F5E59")
    for k, who in enumerate(("Prepared by (site engineer)", "Checked by (project manager)", "Client representative")):
        rr = 11 + k * 4
        sm.cell(rr, 1, who).font = Font(bold=True)
        for cc in range(1, 4):
            sm.cell(rr + 1, cc).border = _BOX
            sm.cell(rr + 2, cc).border = _BOX
        sm.cell(rr + 1, 1, "Name / signature")
        sm.cell(rr + 2, 1, "Date")
        sm.cell(rr + 1, 1).font = sm.cell(rr + 2, 1).font = Font(color="8B8983")
    sm.column_dimensions["A"].width = 34
    sm.column_dimensions["B"].width = 28
    sm.column_dimensions["C"].width = 28
    _client_sheet(wb, project, service, columns, grid, runs or [], when)
    meta = wb.create_sheet(META)
    meta["A1"], meta["A2"], meta["A3"] = MAGIC, str(slug), str(service)
    meta.sheet_state = "hidden"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ the vertical sheet (the client's measurement sheet)
_PEACH = PatternFill("solid", fgColor="FCE4D6")
_GREEN = PatternFill("solid", fgColor="C6E0B4")
_BLK = Side(style="thin", color="000000")
_BX = Border(left=_BLK, right=_BLK, top=_BLK, bottom=_BLK)
_QTY = '#,##0.00;-#,##0.00;"-"'


def _item_lines(code, columns, grid, runs, length=False):
    """The lines of ONE item, top to bottom in structure order: [{"desc","from","to","start","end","qty","remarks","note"}].
    A place with measurement entries gives one line per entry (and one 'recorded without From-To detail' line for what no entry explains);
    a place without entries gives one line with its quantity; the places inside a whole-level total give ONE line for the level (never
    split into zones). `length`: a Mtr / RMT line -- a place holding a quantity that no entry explains says so. Whatever the lines do not explain of the grid's own total is one visible line, so the sheet always ties to the grid."""
    cells, locks = grid["cells"], grid["locks"]
    mine = [ln for ln in runs if ln["code"] == code]
    out, area_done = [], set()

    def line(desc, qty, frm="", to="", start=None, end=None, remarks="", note=False):
        out.append({"desc": desc, "from": frm, "to": to, "start": start, "end": end, "qty": qty, "remarks": remarks, "note": note})

    def from_entries(ln, desc):
        if ln.get("kind") == "run":
            line(desc, ln["qty"], ln.get("from") or "", ln.get("to") or "", ln.get("start"), ln.get("end"), ln.get("remarks") or "")
        else:
            line(desc, ln["qty"], ln.get("from") or "", note=True)

    for col in columns:
        key = f"{code}|{col['id']}"
        v = cells.get(key)
        if not v or v <= 0:
            continue
        kind = (locks.get(key) or {}).get("kind")
        loc = col["loc"] or ""
        if kind == "area":
            if loc in area_done:
                continue
            area_done.add(loc)
            desc = f"   {loc} — whole level / floor"
            ls = [ln for ln in mine if ln["level"] == loc and str(ln["place"]).startswith("Whole level")]
            if ls:
                for ln in ls:
                    from_entries(ln, desc)
            else:
                tot = sum(cells.get(f"{code}|{c2['id']}") or 0 for c2 in columns
                          if (c2["loc"] or "") == loc and (locks.get(f"{code}|{c2['id']}") or {}).get("kind") == "area")
                line(desc, round(tot, 4), remarks="Total for the whole area, not split zone-wise")
            continue
        if kind == "project":
            continue                                           # listed once below as the whole-project lines
        desc = f"   {loc} — {col['name']}" if loc else f"   {col['name']}"
        ls = [ln for ln in mine if ln["level"] == loc and ln["place"] == col["name"]]
        if ls:
            for ln in ls:
                from_entries(ln, desc)
        elif length:
            line(desc, v, frm="Recorded without From-To detail", note=True)
        else:
            line(desc, v)
    for ln in mine:
        if str(ln["place"]).startswith("Whole project"):
            from_entries(ln, "   Whole project (not by place)")
    return out


def _client_sheet(wb, project, service, columns, grid, runs, when):
    """The vertical 'Measurement sheet': the same quantities as the grid, laid out the way a client's measurement sheet reads -- one block per
    BOQ line that has something recorded, its places underneath, and for a length line the From / To and the drum readings behind each length.
    Read-only (the import ignores it). Nothing is added when nothing is recorded."""
    cells = grid["cells"]
    items = [(row, _item_lines(row["code"], columns, grid, runs, bool(row.get("runs_ok")))) for row in grid["rows"]]
    items = [(row, ls) for row, ls in items if ls or any((cells.get(f"{row['code']}|{c['id']}") or 0) > 0 for c in columns)]
    if not items:
        return
    ws = wb.create_sheet(CLIENT_SHEET)
    wb.move_sheet(ws, offset=-1)                               # right after the import sheet, before the summary
    BOLD, NORM = Font(name="Calibri", bold=True, size=11), Font(name="Calibri", size=11)
    ITAL, SMALL = Font(name="Calibri", italic=True, size=11, color="595959"), Font(name="Calibri", italic=True, size=10, color="595959")
    mid = Alignment(horizontal="center", vertical="center", wrap_text=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    NC = 9
    ws.sheet_view.showGridLines = False
    for col, wd in zip("ABCDEFGHI", (9, 44, 8, 30, 26, 11, 11, 13, 34)):
        ws.column_dimensions[col].width = wd
    ws["A1"] = project
    ws["A1"].font = Font(name="Calibri", bold=True, size=13)
    ws["A2"] = f"Work: {service} works"
    ws["A3"] = "Prepared on: " + (when or dt.datetime.now()).strftime("%d %b %Y %H:%M")
    ws["A2"].font = ws["A3"].font = NORM
    ws["A5"] = "Measurement Sheet — quantities recorded at site to date, place by place, with the From-To and drum readings of the length lines"
    ws["A5"].font = BOLD
    ws.merge_cells("H6:I6")
    ws["H6"] = "TO DATE"
    for c_ in ("H", "I"):
        ws[f"{c_}6"].fill, ws[f"{c_}6"].border = _GREEN, _BX
    ws["H6"].font, ws["H6"].alignment = BOLD, mid
    for j, t in enumerate(("Sr. No", "Description / Place", "Unit", "From", "To", "Drum reading START", "Drum reading END", "LENGTH / QTY", "Remarks"), start=1):
        c_ = ws.cell(7, j, t)
        c_.font, c_.fill, c_.border, c_.alignment = BOLD, _PEACH, _BX, mid
    ws.row_dimensions[7].height = 45
    r = 8

    def put(row, vals, font=None, fill=None):
        for j in range(1, NC + 1):
            c_ = ws.cell(row, j)
            c_.border, c_.font = _BX, font or NORM
            if fill:
                c_.fill = fill
        for j, v in vals.items():
            ws.cell(row, j, v)
        for j in (6, 7, 8):
            ws.cell(row, j).number_format = _QTY
            ws.cell(row, j).alignment = Alignment(horizontal="right", vertical="top")

    for row, ls in items:
        put(r, {1: row["code"], 2: row["desc"], 3: row["unit"] or ""}, font=BOLD)
        ws.cell(r, 1).number_format = "@"
        ws.cell(r, 1).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(r, 2).alignment = wrap
        ws.cell(r, 3).alignment = Alignment(horizontal="center", vertical="top")
        r += 1
        first, got = r, 0.0
        for ln in ls:
            put(r, {2: ln["desc"], 4: ln["from"] or None, 5: ln["to"] or None, 6: ln["start"], 7: ln["end"], 8: ln["qty"], 9: ln["remarks"] or None},
                font=ITAL if ln["note"] else None)
            for j in (2, 4, 5, 9):
                ws.cell(r, j).alignment = wrap
            ws.cell(r, 9).font = SMALL
            got += ln["qty"]
            r += 1
        gap = round(sum((cells.get(f"{row['code']}|{c['id']}") or 0) for c in columns) - got, 4)
        if abs(gap) > 0.005:                                   # never silently untied from the grid
            put(r, {2: "   Other quantity recorded for this line (see the Measurement sheet)", 8: gap}, font=ITAL)
            r += 1
        put(r, {2: "Total to date (as in the grid)", 3: row["unit"] or ""}, font=BOLD)
        ws.cell(r, 2).alignment = Alignment(horizontal="right")
        ws.cell(r, 3).alignment = Alignment(horizontal="center")
        ws.cell(r, 8, f"=SUM(H{first}:H{r - 1})" if r > first else 0)
        r += 1
    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=NC)
    ws.cell(r, 1, "Note: Each length line is one measurement between the two points shown, with the drum readings it was measured by. "
                  "Length = difference of the two readings. “Whole level / floor” means one quantity was recorded for the whole level, not split place by place. "
                  "A line struck off at site is not listed.").font = SMALL
    ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 30
    r += 2
    for a, b, t in ((1, 3, "Prepared by (Contractor)"), (4, 6, "Checked by (Contractor)"), (7, NC, "Certified by (Client / Consultant)")):
        ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
        ws.cell(r, a, t).font = BOLD
        ws.cell(r, a).alignment = mid
        ws.merge_cells(start_row=r + 1, start_column=a, end_row=r + 3, end_column=b)
        ws.merge_cells(start_row=r + 4, start_column=a, end_row=r + 4, end_column=b)
        ws.cell(r + 4, a, "Name / Sign / Date").font = SMALL
        for rr in range(r, r + 5):
            for cc in range(a, b + 1):
                ws.cell(rr, cc).border = Border(left=_BLK if cc == a else None, right=_BLK if cc == b else None,
                                                top=_BLK if rr in (r, r + 4) else None, bottom=_BLK if rr in (r, r + 3, r + 4) else None)
    ws.print_area = f"A1:{get_column_letter(NC)}{r + 4}"
    ws.page_setup.orientation, ws.page_setup.paperSize = "landscape", 9
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = "6:7"
    ws.freeze_panes = "A8"


# ------------------------------------------------------------------ reading a quantity typed in Excel
def _calc(s):
    """numbers, + and * and brackets only (a sum a quantity surveyor types: 10+2+30). No names, no cell references."""
    i = 0

    def ws():
        nonlocal i
        while i < len(s) and s[i] == " ":
            i += 1

    def num():
        nonlocal i
        ws()
        if i < len(s) and s[i] == "(":
            i += 1
            v = add()
            ws()
            if i >= len(s) or s[i] != ")":
                raise ValueError
            i += 1
            return v
        m = re.match(r"\d+(\.\d+)?|\.\d+", s[i:])
        if not m:
            raise ValueError
        i += len(m.group(0))
        return float(m.group(0))

    def mul():
        nonlocal i
        v = num()
        ws()
        while i < len(s) and s[i] == "*":
            i += 1
            v *= num()
            ws()
        return v

    def add():
        nonlocal i
        v = mul()
        ws()
        while i < len(s) and s[i] == "+":
            i += 1
            v += mul()
            ws()
        return v

    v = add()
    ws()
    if i != len(s):
        raise ValueError
    return v


def parse_qty(v, count=False):
    """-> ("blank",) | ("ok", number) | ("bad", reason). `v` is what the cell holds (a number, or text such as '10+2+30')."""
    if v is None or (isinstance(v, str) and v.strip() == ""):
        return ("blank",)
    if isinstance(v, bool):
        return ("bad", "not a number")
    if isinstance(v, (int, float)):
        x = float(v)
    else:
        s = str(v).strip()
        if s.startswith("="):
            s = s[1:].strip()
        s = s.replace("×", "*").replace("x", "*").replace("X", "*")
        if re.fullmatch(r"\d{1,3}(,\d{3})+(\.\d+)?", s):
            s = s.replace(",", "")
        if "-" in s:
            return ("bad", "a minus is not allowed")
        if not re.fullmatch(r"[0-9+*().\s]+", s):
            return ("bad", "not a number (formulas that point at other cells cannot be read; type the number)")
        try:
            x = _calc(s)
        except (ValueError, RecursionError):
            return ("bad", "the sum is incomplete")
    if math.isnan(x) or math.isinf(x):
        return ("bad", "not a number")
    if x < 0:
        return ("bad", "a quantity cannot be negative")
    x = round(x, 6)
    if x > 1e9:
        return ("bad", "a very large number, check it")
    if count and abs(x - round(x)) > 1e-9:
        return ("bad", "this item is counted in whole numbers")
    return ("ok", x)


# ------------------------------------------------------------------ import: read the file, make no decision
class XlError(ValueError):
    """The file is not one this grid made (message is shown to the person as it is)."""


def read(data):
    """-> {"slug","service","keys":{col:place_id},"rows":[{"row","code","cells":{place_id: raw},"formulas":{place_id: "=..."}}]}. Raw values are what the cells hold: a number, or the
    text / formula typed there; a formula without a stored result is passed on as its text so parse_qty can read simple sums."""
    if not data or len(data) > MAX_BYTES:
        raise XlError("The file is empty or larger than 8 MB.")
    if data[:2] != b"PK":
        raise XlError("This is not an Excel (.xlsx) file.")
    try:
        wv = load_workbook(io.BytesIO(data), data_only=True)
        wf = load_workbook(io.BytesIO(data), data_only=False)
    except Exception:
        raise XlError("This file could not be read as an Excel workbook.")
    if SHEET not in wv.sheetnames:
        raise XlError(f'There is no sheet named "{SHEET}". Use a file exported from this grid.')
    sv, sf = wv[SHEET], wf[SHEET]
    if str(sv.cell(KEY_ROW, 1).value or "").strip() != "#KEY":
        raise XlError("The hidden key row is missing. Use a file exported from this grid, and do not delete rows 1-5.")
    meta = {"slug": None, "service": None}
    if META in wv.sheetnames:
        m = wv[META]
        if str(m["A1"].value or "") == MAGIC:
            meta = {"slug": m["A2"].value, "service": m["A3"].value}
    keys = {}
    for col in range(FIRST_COL, min(sv.max_column, FIRST_COL + MAX_PLACES + 8) + 1):
        k = sv.cell(KEY_ROW, col).value
        if k is not None and str(k).strip() != "":
            keys[col] = str(k).strip()
    if not keys:
        raise XlError("No place keys found in row 3. Use a file exported from this grid.")
    rows, blank_run = [], 0
    for r in range(FIRST_ROW, min(sv.max_row, FIRST_ROW + MAX_LINES) + 1):
        code = sv.cell(r, 1).value
        if code is None or str(code).strip() == "":
            blank_run += 1
            if blank_run >= 3:
                break
            continue
        blank_run = 0
        cells, formulas = {}, {}
        for col, pid in keys.items():
            v = sv.cell(r, col).value
            f = sf.cell(r, col).value
            if isinstance(f, str) and f.startswith("="):
                formulas[pid] = f
                if v is None:
                    v = f
            cells[pid] = v
        if isinstance(code, float) and code.is_integer():
            code = int(code)                      # a code typed as a plain number in Excel (39, not 39.0)
        rows.append({"row": r, "code": str(code).strip(), "cells": cells, "formulas": formulas})
    return {"slug": meta["slug"], "service": meta["service"], "keys": keys, "rows": rows}
