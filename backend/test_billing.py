"""Billing (RA bills) -- end to end through the real FastAPI routers
(siteprogress upload -> progress -> stock links -> billing), on a ProjectBase-
shaped workbook built exactly like the real Thoth export: Internal Ref No in
column A, Client Ref No in B, rate-only ("RO") lines, "Cost of ..." sub-total
rows between annexure groups.

Project data lives in tmp_path (siteprogress.PROJECTS / RUNS are patched), so
nothing touches real data.
"""
import sys
import json
import os
import datetime as dt
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pandas as pd
import pytest
import openpyxl
from fastapi import FastAPI
from fastapi.testclient import TestClient

from backend import siteprogress as sp, billing, structure, boq

SLUG = "test-thoth"

HEAD = ["Internal Ref No", "Client Ref No", "BOQ", "BOQ Unit", "Order Quantity", "Design Quantity", "Rate",
        "Discount %", "HSN Code", "Tax", "Progress Method", "Billing Term", "Summary Item"]

# (internal, client, text, unit, qty, rate)  -- unit None => heading / sub-total row
FIRE_ROWS = [
    ("1", "A", "Fire Pump & Accessories", None, None, None),
    ("2", "1", "Diesel engine pump set", "Nos", 2, 100000),
    ("3", "2", "Butterfly valve 150mm", "Nos", 4, 5000),
    ("4", "3", "GI pipe 100mm", "Rmt", 100, 800),
    ("5", "4", "Rate only fitting", "Nos", "RO", 1200),
    ("6", None, "Cost of Fire Pump & Accessories", None, None, None),
    ("7", "B", "Hydrant System", None, None, None),
    ("8", "5", "Hydrant valve", "Nos", 10, 3000),
    ("9", None, "Cost of Hydrant System", None, None, None),
]
HVAC_ROWS = [
    ("1", "1", "Chiller", "Nos", 2, 5000000),
    ("2", "2", "Pump", "Nos", 4, 200000),
]


def _sheet(ws, rows):
    ws.append(HEAD)
    for internal, client, text, unit, qty, rate in rows:
        ws.append([internal, client, text, unit, qty, qty if qty != "RO" else None, rate,
                   None, None, "GST @ 18%", None, None, None])


def make_workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fire"
    _sheet(ws, FIRE_ROWS)
    _sheet(wb.create_sheet("HVAC"), HVAC_ROWS)
    wb.save(path)


@pytest.fixture
def env(tmp_path, monkeypatch):
    monkeypatch.setattr(sp, "PROJECTS", tmp_path / "projects")
    monkeypatch.setattr(sp, "RUNS", tmp_path / "runs")
    sp.PROJECTS.mkdir(parents=True)
    sp.RUNS.mkdir(parents=True)
    for c in (sp._read_forecast_parquet_cached, sp._forecast_pool_cached, sp._full_run_rows_cached,
              sp._linkage_match_cached):
        c.cache_clear()
    app = FastAPI()
    app.include_router(sp.router)
    app.include_router(billing.router)
    client = TestClient(app)
    xlsx = tmp_path / "boq.xlsx"
    make_workbook(xlsx)
    with open(xlsx, "rb") as fh:
        r = client.post(f"/api/siteprogress/{SLUG}/boq", files={"file": ("boq.xlsx", fh)})
    assert r.status_code == 200, r.text
    d = sp._dir(SLUG)
    s = structure.hotel("Thoth", floors=["1", "2"], room_labels=["a", "b"])
    (d / "structure.json").write_text(s.to_json())
    return client, d, tmp_path


def make_run(tmp_path, name, created, rows):
    rd = sp.RUNS / name
    rd.mkdir(parents=True, exist_ok=True)
    (rd / "meta.json").write_text(json.dumps({"project_slug": SLUG, "created": created}))
    pd.DataFrame([{"service": "Fire & HVAC", "rate_per_day": 0.0, "days_left": 99.0, "status": "OK",
                   **r} for r in rows]).to_parquet(rd / "forecast.parquet")
    for c in (sp._read_forecast_parquet_cached, sp._forecast_pool_cached, sp._full_run_rows_cached,
              sp._linkage_match_cached):
        c.cache_clear()


def link(client, code, material, factor=None):
    r = client.post(f"/api/siteprogress/{SLUG}/links",
                    json={"service": "Fire", "item_code": code, "materials": [{"material": material, "factor": factor}]})
    assert r.status_code == 200, r.text


def progress(client, code, frac):
    r = client.post(f"/api/siteprogress/{SLUG}/progress/item", json={"service": "Fire", "item_code": code, "frac": frac})
    assert r.status_code == 200, r.text


def setup_fire(client, mode="separate", sup=70, ins=30):
    r = client.post(f"/api/billing/{SLUG}/setup", json={"service": "Fire", "mode": mode, "supply_pct": sup, "install_pct": ins})
    assert r.status_code == 200, r.text


STOCK1 = [
    {"material": "BUTTERFLY VALVE 150MM", "unit": "Nos", "stock": 2.0, "total_consumed": 1.0},   # received 3
    {"material": "GI PIPE 100MM", "unit": "Rmt", "stock": 30.0, "total_consumed": 20.0},        # received 50
    {"material": "HYDRANT VALVE", "unit": "Nos", "stock": 10.0, "total_consumed": 0.0},         # received 10
]


def prepared(env):
    client, d, tmp = env
    make_run(tmp, "run1", "2026-09-01T10:00:00", STOCK1)
    link(client, "2", "BUTTERFLY VALVE 150MM")
    link(client, "3", "GI PIPE 100MM")
    link(client, "5", "HYDRANT VALVE")
    progress(client, "2", 0.5)      # butterfly installed 2 of 4
    progress(client, "3", 0.2)      # pipe installed 20 of 100
    setup_fire(client)
    return client, d, tmp


def line(bill, code):
    return next(x for x in bill["lines"] if x["item_code"] == code)


# --------------------------------------------------------------- parsing
def test_internal_ref_and_sheet_row_are_read_from_the_real_sheet(env):
    client, d, _ = env
    lines = billing._lines(d, "Fire")
    by = {l["item_code"]: l for l in lines}
    assert set(by) == {"1", "2", "3", "5"}                    # RO line "4" is not billable
    assert by["2"]["internal_ref"] == "3" and by["2"]["src_row"] == 4 and by["2"]["sheet"] == "Fire"
    assert by["5"]["internal_ref"] == "8" and by["5"]["src_row"] == 9
    assert by["2"]["rate"] == 5000 and by["2"]["boq_qty"] == 4


def test_old_project_without_ref_columns_gets_them_from_saved_boq_xlsx(env):
    client, d, _ = env
    df = pd.read_parquet(d / "boq.parquet").drop(columns=["internal_ref", "src_row"])
    df.to_parquet(d / "boq.parquet")
    by = {l["item_code"]: l for l in billing._lines(d, "Fire")}
    assert by["3"]["internal_ref"] == "4" and by["3"]["src_row"] == 5


# ---------------------------------------------------------------- setup
def test_service_without_payment_term_cannot_bill(env):
    client, _, _ = env
    r = client.post(f"/api/billing/{SLUG}/ra", json={"service": "HVAC", "series": "supply"})
    assert r.status_code == 409 and "setup" in r.json()["detail"].lower()


def test_unset_percent_blocks_only_that_series(env):
    client, _, _ = env
    client.post(f"/api/billing/{SLUG}/setup", json={"service": "HVAC", "mode": "separate", "supply_pct": 75})
    assert client.post(f"/api/billing/{SLUG}/ra", json={"service": "HVAC", "series": "supply"}).status_code == 200
    r = client.post(f"/api/billing/{SLUG}/ra", json={"service": "HVAC", "series": "install"})
    assert r.status_code == 409 and "Installation" in r.json()["detail"]


def test_percent_validation_and_sum(env):
    client, _, _ = env
    bad = client.post(f"/api/billing/{SLUG}/setup", json={"service": "Fire", "mode": "separate", "supply_pct": 70, "install_pct": 40})
    assert bad.status_code == 400
    assert client.post(f"/api/billing/{SLUG}/setup", json={"service": "Fire", "mode": "separate", "supply_pct": 0}).status_code == 400


def test_mode_is_only_the_default_it_never_forbids_a_bill_type_and_can_change_after_an_ra(env):
    client, _, _ = env
    setup_fire(client, "together", 60, 40)
    # a Together service can still bill Supply only or Installation only ...
    assert client.post(f"/api/billing/{SLUG}/ra", json={"service": "Fire", "series": "supply"}).status_code == 200
    # ... but only one draft per stage: Together would carry the supply stage again
    r = client.post(f"/api/billing/{SLUG}/ra", json={"service": "Fire", "series": "together"})
    assert r.status_code == 409 and "same stage" in r.json()["detail"]
    assert client.post(f"/api/billing/{SLUG}/ra", json={"service": "Fire", "series": "install"}).status_code == 200
    # the default can be changed after RAs exist (it only decides what is offered first)
    r = client.post(f"/api/billing/{SLUG}/setup", json={"service": "Fire", "mode": "separate", "supply_pct": 70, "install_pct": 30})
    assert r.status_code == 200, r.text
    st = client.get(f"/api/billing/{SLUG}/setup").json()["services"]
    assert next(x for x in st if x["service"] == "Fire")["mode"] == "separate"


# ----------------------------------------------------------- bill maths
def test_supply_ra_uses_stock_received_and_leaves_unlinked_manual(env):
    client, d, _ = prepared(env)
    ra = client.post(f"/api/billing/{SLUG}/ra", json={"service": "Fire", "series": "supply"}).json()["id"]
    b = client.get(f"/api/billing/{SLUG}/ra/{ra}").json()
    bf = line(b, "2")["supply"]
    assert bf["utd"] == 3 and bf["prev"] == 0 and bf["this"] == 3          # received = 1 used + 2 stock
    assert bf["amount"] == pytest.approx(5000 * 3 * 0.70)
    assert line(b, "3")["supply"]["this"] == 50 and line(b, "3")["supply"]["amount"] == pytest.approx(800 * 50 * .7)
    assert line(b, "5")["supply"]["this"] == 10                              # capped nowhere, == BOQ qty
    assert all(l["item_code"] != "1" for l in b["lines"])                  # diesel pump: no link -> not guessed
    assert b["unfilled"]["supply"] == 1
    assert b["this_total"] == pytest.approx(10500 + 28000 + 21000)
    assert b["ra"]["label"] == "Fire · Supply RA 1" and b["ra"]["status"] == "draft"


def test_install_ra_uses_site_progress_used_and_is_a_separate_series(env):
    client, d, _ = prepared(env)
    s = client.post(f"/api/billing/{SLUG}/ra", json={"service": "Fire", "series": "supply"}).json()
    i = client.post(f"/api/billing/{SLUG}/ra", json={"service": "Fire", "series": "install"}).json()
    b = client.get(f"/api/billing/{SLUG}/ra/{i['id']}").json()
    assert i["label"] == "Fire · Installation RA 1"                          # numbered per series
    assert line(b, "2")["install"]["this"] == pytest.approx(2) and line(b, "3")["install"]["this"] == pytest.approx(20)
    assert b["this_total"] == pytest.approx(5000 * 2 * .3 + 800 * 20 * .3)
    assert "supply" not in line(b, "2")


def test_typed_qty_needs_reason_context_and_cannot_pass_boq_qty(env):
    client, d, _ = prepared(env)
    ra = client.post(f"/api/billing/{SLUG}/ra", json={"service": "Fire", "series": "supply"}).json()["id"]
    r = client.post(f"/api/billing/{SLUG}/ra/{ra}/line", json={"item_code": "1", "stage": "supply", "qty": 5})
    assert r.status_code == 400 and "BOQ qty is 2" in r.json()["detail"]
    assert client.post(f"/api/billing/{SLUG}/ra/{ra}/line", json={"item_code": "1", "stage": "supply", "qty": 1, "reason": "gate pass 44"}).status_code == 200
    b = client.get(f"/api/billing/{SLUG}/ra/{ra}").json()
    dp = line(b, "1")["supply"]
    assert dp["this"] == 1 and dp["amount"] == pytest.approx(100000 * .7) and dp["reason"] == "gate pass 44"
    assert dp["source"] == "manual-entered" and b["unfilled"]["supply"] == 0
    client.post(f"/api/billing/{SLUG}/ra/{ra}/line", json={"item_code": "1", "stage": "supply", "qty": None})
    assert client.get(f"/api/billing/{SLUG}/ra/{ra}").json()["unfilled"]["supply"] == 1


def test_second_ra_is_cumulative_and_overbilled_line_goes_negative(env):
    client, d, tmp = prepared(env)
    ra1 = client.post(f"/api/billing/{SLUG}/ra", json={"service": "Fire", "series": "supply"}).json()["id"]
    assert client.post(f"/api/billing/{SLUG}/ra/{ra1}/lock").status_code == 200
    # new stock run: butterfly now shows 2 received (1 used + 1 stock) though 3 were billed; pipe 80 received
    make_run(tmp, "run2", "2026-10-01T10:00:00", [
        {"material": "BUTTERFLY VALVE 150MM", "unit": "Nos", "stock": 1.0, "total_consumed": 1.0},
        {"material": "GI PIPE 100MM", "unit": "Rmt", "stock": 40.0, "total_consumed": 40.0},
        {"material": "HYDRANT VALVE", "unit": "Nos", "stock": 10.0, "total_consumed": 0.0}])
    ra2 = client.post(f"/api/billing/{SLUG}/ra", json={"service": "Fire", "series": "supply"}).json()
    assert ra2["label"] == "Fire · Supply RA 2"
    b = client.get(f"/api/billing/{SLUG}/ra/{ra2['id']}").json()
    bf, gi = line(b, "2")["supply"], line(b, "3")["supply"]
    assert (bf["prev"], bf["utd"], bf["this"]) == (3, 2, -1)
    assert bf["amount"] == pytest.approx(-3500)
    assert (gi["prev"], gi["utd"], gi["this"]) == (50, 80, 30)
    assert all(l["item_code"] != "5" for l in b["lines"])                   # nothing new on this line -> not listed
    assert any("3 Nos was billed earlier" in w["msg"] for w in b["warnings"])
    assert b["this_total"] == pytest.approx(-3500 + 800 * 30 * .7)
    # group totals: previous bill value comes from the LOCKED snapshot
    g = {x["name"]: x for x in b["groups"]}
    assert g["Fire Pump & Accessories"]["stages"]["supply"]["prev"] == pytest.approx(10500 + 28000)
    assert g["Hydrant System"]["stages"]["supply"]["prev"] == pytest.approx(21000)


def test_locked_ra_is_a_snapshot_later_rate_or_progress_changes_do_not_move_it(env):
    client, d, tmp = prepared(env)
    ra = client.post(f"/api/billing/{SLUG}/ra", json={"service": "Fire", "series": "supply"}).json()["id"]
    before = client.post(f"/api/billing/{SLUG}/ra/{ra}/lock").json()["this_total"]
    rates = json.loads((d / "rates.json").read_text())
    rates["Fire"]["2"] = 9999
    (d / "rates.json").write_text(json.dumps(rates))
    progress(client, "2", 1.0)
    b = client.get(f"/api/billing/{SLUG}/ra/{ra}").json()
    assert b["this_total"] == pytest.approx(before) and b["ra"]["status"] == "locked"
    assert line(b, "2")["rate"] == 5000
    r = client.post(f"/api/billing/{SLUG}/ra/{ra}/line", json={"item_code": "2", "stage": "supply", "qty": 1})
    assert r.status_code == 409


def test_locked_ra_totals_feed_the_next_ra_even_after_a_rate_change(env):
    client, d, _ = prepared(env)
    ra1 = client.post(f"/api/billing/{SLUG}/ra", json={"service": "Fire", "series": "supply"}).json()["id"]
    client.post(f"/api/billing/{SLUG}/ra/{ra1}/lock")
    rates = json.loads((d / "rates.json").read_text()); rates["Fire"]["2"] = 9999
    (d / "rates.json").write_text(json.dumps(rates))
    ra2 = client.post(f"/api/billing/{SLUG}/ra", json={"service": "Fire", "series": "supply"}).json()["id"]
    b = client.get(f"/api/billing/{SLUG}/ra/{ra2}").json()
    assert b["totals"]["supply"]["prev"] == pytest.approx(59500)           # old rate, from snapshot


def test_one_draft_per_series_reopen_and_delete_rules(env):
    client, _, _ = prepared(env)
    ra1 = client.post(f"/api/billing/{SLUG}/ra", json={"service": "Fire", "series": "supply"}).json()["id"]
    assert client.post(f"/api/billing/{SLUG}/ra", json={"service": "Fire", "series": "supply"}).status_code == 409
    assert client.post(f"/api/billing/{SLUG}/ra/{ra1}/reopen").status_code == 409     # not locked
    client.post(f"/api/billing/{SLUG}/ra/{ra1}/lock")
    assert client.delete(f"/api/billing/{SLUG}/ra/{ra1}").status_code == 409          # locked
    ra2 = client.post(f"/api/billing/{SLUG}/ra", json={"service": "Fire", "series": "supply"}).json()["id"]
    assert client.post(f"/api/billing/{SLUG}/ra/{ra1}/reopen").status_code == 409     # draft exists
    client.post(f"/api/billing/{SLUG}/ra/{ra2}/lock")
    assert client.post(f"/api/billing/{SLUG}/ra/{ra1}/reopen").status_code == 409     # not the last
    assert client.post(f"/api/billing/{SLUG}/ra/{ra2}/reopen").status_code == 200
    assert client.get(f"/api/billing/{SLUG}/ra/{ra2}").json()["ra"]["status"] == "draft"
    assert client.delete(f"/api/billing/{SLUG}/ra/{ra2}").status_code == 200


def test_together_ra_carries_both_stages(env):
    client, _, _ = prepared(env)
    setup_fire(client, "together", 60, 40)
    ra = client.post(f"/api/billing/{SLUG}/ra", json={"service": "Fire", "series": "together"}).json()
    b = client.get(f"/api/billing/{SLUG}/ra/{ra['id']}").json()
    l = line(b, "2")
    assert l["supply"]["amount"] == pytest.approx(5000 * 3 * .6) and l["install"]["amount"] == pytest.approx(5000 * 2 * .4)
    assert ra["label"] == "Fire · Supply + Installation RA 1"


def test_line_percent_overrides_the_service_percent(env):
    client, _, _ = prepared(env)
    client.post(f"/api/billing/{SLUG}/setup", json={"service": "Fire", "mode": "separate", "supply_pct": 70, "install_pct": 30,
                                                  "line_pct": {"3": {"supply": 90}}})
    ra = client.post(f"/api/billing/{SLUG}/ra", json={"service": "Fire", "series": "supply"}).json()["id"]
    b = client.get(f"/api/billing/{SLUG}/ra/{ra}").json()
    assert line(b, "3")["supply"]["pct"] == 90 and line(b, "3")["supply"]["amount"] == pytest.approx(800 * 50 * .9)
    assert line(b, "2")["supply"]["pct"] == 70


# ---------------------------------------------------- supply safety rules
def test_shared_material_multi_link_and_missing_factor_stay_manual(env):
    client, d, tmp = env
    make_run(tmp, "run1", "2026-09-01T10:00:00", STOCK1 + [{"material": "PIPE B", "unit": "Mtr", "stock": 5.0, "total_consumed": 5.0}])
    link(client, "2", "BUTTERFLY VALVE 150MM")
    link(client, "5", "BUTTERFLY VALVE 150MM")            # same material on a second item -> shared
    link(client, "3", "PIPE B")                              # unit Mtr vs Rmt, no factor
    setup_fire(client)
    ra = client.post(f"/api/billing/{SLUG}/ra", json={"service": "Fire", "series": "supply"}).json()["id"]
    b = client.get(f"/api/billing/{SLUG}/ra/{ra}").json()
    assert b["lines"] == [] and b["unfilled"]["supply"] == 4
    link(client, "3", "PIPE B", factor=2.0)                  # now a known factor: 10 received / 2
    b = client.get(f"/api/billing/{SLUG}/ra/{ra}").json()
    assert line(b, "3")["supply"]["utd"] == 5


# --------------------------------------------------------------- groups
def test_groups_come_from_the_sheets_own_subtotal_rows(env):
    client, d, _ = env
    g = billing.suggest_groups(d, billing._lines(d, "Fire"))
    assert [(x["sr"], x["name"], x["from"], x["to"]) for x in g] == [
        ("A", "Fire Pump & Accessories", 2, 7), ("B", "Hydrant System", 8, 10)]


def test_saved_groups_replace_the_suggestion_and_wo_can_be_typed(env):
    client, d, _ = prepared(env)
    client.post(f"/api/billing/{SLUG}/groups/Fire", json={"groups": [
        {"sr": "1", "name": "Everything", "sheet": "Fire", "from": 2, "to": 20, "wo": 123456}]})
    ra = client.post(f"/api/billing/{SLUG}/ra", json={"service": "Fire", "series": "supply"}).json()["id"]
    b = client.get(f"/api/billing/{SLUG}/ra/{ra}").json()
    assert len(b["groups"]) == 1 and b["groups"][0]["wo_amount"] == 123456 and b["groups"][0]["wo_calc"] > 0
    client.post(f"/api/billing/{SLUG}/groups/Fire", json={"groups": []})
    assert len(client.get(f"/api/billing/{SLUG}/ra/{ra}").json()["groups"]) == 2


# ------------------------------------------------------------- overview
def test_overview_lists_ready_to_bill_and_blocked_services(env):
    client, d, _ = prepared(env)
    o = client.get(f"/api/billing/{SLUG}").json()
    fire = next(s for s in o["services"] if s["service"] == "Fire")
    assert fire["ready"]["supply"]["ready"] == pytest.approx(59500)
    assert fire["ready"]["install"]["ready"] == pytest.approx(3000 + 4800)
    hv = next(s for s in o["services"] if s["service"] == "HVAC")
    assert hv["set"] is False and hv["ready"] == {}
    assert o["ready_total"] == pytest.approx(59500 + 7800)
    assert o["contract_value"] == pytest.approx(2 * 100000 + 4 * 5000 + 100 * 800 + 10 * 3000 + 2 * 5e6 + 4 * 2e5)


# ---------------------------------------------------------------- export
import shutil
import subprocess


def _recalc(path, outdir):
    """Recalculate with LibreOffice so formula results are cached, then load values."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        pytest.skip("LibreOffice not installed")
    subprocess.run([soffice, "--headless", "--convert-to", "xlsx", "--outdir", str(outdir), str(path)],
                   check=True, capture_output=True, timeout=180)
    return openpyxl.load_workbook(Path(outdir) / Path(path).name, data_only=True)


def _export(client, tmp, name="out.xlsx"):
    r = client.get(f"/api/billing/{SLUG}/export")
    assert r.status_code == 200, r.text
    p = tmp / name
    p.write_bytes(r.content)
    return p


def _grand(ws):
    for r in range(8, ws.max_row + 1):
        if ws.cell(r, 2).value == "Grand Total":
            return [ws.cell(r, c).value for c in range(3, 7)]


def test_export_refuses_with_no_ra(env):
    client, _, _ = env
    assert client.get(f"/api/billing/{SLUG}/export").status_code == 409


def test_export_workbook_matches_the_api_and_keeps_original_columns(env):
    client, d, tmp = prepared(env)
    ra1 = client.post(f"/api/billing/{SLUG}/ra", json={"service": "Fire", "series": "supply"}).json()["id"]
    client.post(f"/api/billing/{SLUG}/ra/{ra1}/line", json={"item_code": "1", "stage": "supply", "qty": 1, "reason": "gate pass 44"})
    client.post(f"/api/billing/{SLUG}/ra/{ra1}/lock")
    make_run(tmp, "run2", "2026-10-01T10:00:00", [
        {"material": "BUTTERFLY VALVE 150MM", "unit": "Nos", "stock": 1.0, "total_consumed": 1.0},
        {"material": "GI PIPE 100MM", "unit": "Rmt", "stock": 40.0, "total_consumed": 40.0},
        {"material": "HYDRANT VALVE", "unit": "Nos", "stock": 10.0, "total_consumed": 0.0}])
    ra2 = client.post(f"/api/billing/{SLUG}/ra", json={"service": "Fire", "series": "supply"}).json()["id"]
    ri = client.post(f"/api/billing/{SLUG}/ra", json={"service": "Fire", "series": "install"}).json()["id"]
    client.post(f"/api/billing/{SLUG}/meta", json={"client": "Thoth Mall Pvt Ltd", "contractor": "Vedant Engitech", "wo_no": "WO/12"})
    p = _export(client, tmp)
    wb0 = openpyxl.load_workbook(d / "boq.xlsx")
    wbf = openpyxl.load_workbook(p)
    # original cells A..U never change
    for row0, rowf in zip(wb0["Fire"].iter_rows(max_col=21), wbf["Fire"].iter_rows(max_col=21)):
        assert [c.value for c in row0] == [c.value for c in rowf]
    assert wbf.sheetnames[0] == "RA Register" and "Invoice (format pending)" not in wbf.sheetnames
    assert {"Fire Supply RA-01", "Fire Supply RA-02", "Fire Install RA-01"} <= set(wbf.sheetnames)
    # every RA has its Tax Invoice sheet, right before its own annexure tab
    for tab, inv in (("Fire Supply RA-01", "Fire Supply RA-01 INVOICE"), ("Fire Supply RA-02", "Fire Supply RA-02 INVOICE"),
                     ("Fire Install RA-01", "Fire Install RA-01 INVOICE")):
        assert wbf.sheetnames.index(inv) == wbf.sheetnames.index(tab) - 1, inv
    rc = _recalc(p, tmp / "rc")
    for ra_id, tab, stage in ((ra1, "Fire Supply RA-01", "supply"), (ra2, "Fire Supply RA-02", "supply"), (ri, "Fire Install RA-01", "install")):
        api = client.get(f"/api/billing/{SLUG}/ra/{ra_id}").json()
        wo, prev, this, tot = _grand(rc[tab])
        assert this == pytest.approx(api["totals"][stage]["this"], abs=0.01), tab
        assert prev == pytest.approx(api["totals"][stage]["prev"], abs=0.01), tab
        assert tot == pytest.approx(prev + this, abs=0.01)
    # annexure text + group names
    a = rc["Fire Supply RA-02"]
    assert a["A1"].value == "Thoth Mall Pvt Ltd" and a["B2"].value == "WO/12" and a["B4"].value == "Vedant Engitech"
    assert [a.cell(r, 2).value for r in (8, 9)] == ["Fire Pump & Accessories", "Hydrant System"]
    assert a["D5"].value == "SUPPLY @ 70%"
    # register + service-sheet summary
    reg = rc["RA Register"]
    rows = {(reg.cell(r, 1).value, reg.cell(r, 2).value, reg.cell(r, 3).value): reg.cell(r, 5).value for r in range(5, 12)}
    assert rows[("Fire", "Supply", "RA-01")] == "Invoiced" and rows[("Fire", "Supply", "RA-02")] == "Draft"
    fire = rc["Fire"]
    row = 4   # Butterfly valve, internal ref 3
    assert fire["V4"].value == 70 and fire["W4"].value == 30
    assert fire["Z4"].value == pytest.approx(3 + (-1))                      # supply qty to date
    assert fire["AA4"].value == pytest.approx(2)                            # install qty to date
    assert fire["AE4"].value and "Over-billed" in fire["AE4"].value
    # locked RA-01 is written as values, draft as formulas
    assert not str(wbf["Fire"]["AG4"].value).startswith("=") and str(wbf["Fire"]["AI4"].value).startswith("=")


def test_export_refuses_to_overwrite_a_sheet_that_already_uses_the_summary_columns(env):
    client, d, tmp = prepared(env)
    client.post(f"/api/billing/{SLUG}/ra", json={"service": "Fire", "series": "supply"})
    wb = openpyxl.load_workbook(d / "boq.xlsx")
    wb["Fire"]["V5"] = 123
    wb.save(d / "boq.xlsx")
    r = client.get(f"/api/billing/{SLUG}/export")
    assert r.status_code == 409 and "column V" in r.json()["detail"] and "row 5" in r.json()["detail"]


# ----------------------------------------------- the real Thoth multi-sheet BOQ
REAL = os.environ.get("THOTH_BOQ", "/root/.claude/uploads/cc641dec-fdd2-5679-9e93-268c1cf0e51e/564f322c-Thoth_Malll_Multi_Sheet.xlsx")


@pytest.mark.skipif(not os.path.exists(REAL), reason="real Thoth BOQ not available here (set THOTH_BOQ)")
def test_real_thoth_boq_every_annexure_total_equals_the_api(tmp_path, monkeypatch):
    import random
    monkeypatch.setattr(sp, "PROJECTS", tmp_path / "p")
    monkeypatch.setattr(sp, "RUNS", tmp_path / "r")
    sp.PROJECTS.mkdir()
    sp.RUNS.mkdir()
    app = FastAPI()
    app.include_router(sp.router)
    app.include_router(billing.router)
    c = TestClient(app)
    with open(REAL, "rb") as fh:
        assert c.post(f"/api/siteprogress/{SLUG}/boq", files={"file": ("b.xlsx", fh)}).status_code == 200
    d = sp._dir(SLUG)
    (d / "structure.json").write_text(structure.hotel("Thoth", floors=["1", "2", "3"], room_labels=["a", "b"]).to_json())
    random.seed(3)
    made = []
    for svc in billing._services(d):
        lines = billing._lines(d, svc)
        assert lines and all(l["_ref_ok"] for l in lines)               # every billable line found its sheet row
        for l in random.sample(lines, min(40, len(lines))):
            c.post(f"/api/siteprogress/{SLUG}/progress/item", json={"service": svc, "item_code": l["item_code"], "frac": random.choice([.25, .5, 1.0])})
        assert c.post(f"/api/billing/{SLUG}/setup", json={"service": svc, "mode": "separate", "supply_pct": 70, "install_pct": 30}).status_code == 200
        made.append((svc, c.post(f"/api/billing/{SLUG}/ra", json={"service": svc, "series": "install"}).json()["id"]))
    c.post(f"/api/billing/{SLUG}/ra/{made[0][1]}/lock")
    nxt = c.post(f"/api/billing/{SLUG}/ra", json={"service": made[0][0], "series": "install"}).json()["id"]
    made.append((made[0][0], nxt))
    p = _export(c, tmp_path)
    rc = _recalc(p, tmp_path / "rc")
    for svc, ra_id in made:
        b = c.get(f"/api/billing/{SLUG}/ra/{ra_id}").json()
        assert b["groups"]
        assert b["lines"] or ra_id == made[-1][1]          # the 2nd Fire RA has nothing new, so no lines
        wo, prev, this, tot = _grand(rc[billing._tab_name(svc, "install", {"no": b["ra"]["no"]})])
        assert this == pytest.approx(b["totals"]["install"]["this"], abs=0.01)
        assert prev == pytest.approx(b["totals"]["install"]["prev"], abs=0.01)
        assert wo == pytest.approx(b["contract_value"], abs=1.0)
    assert rc["Fire"]["X185"].value == 3000000                           # Overhead lumpsum column left untouched


def test_lines_picker_lists_every_billable_line_once(env):
    client, _, _ = env
    r = client.get(f"/api/billing/{SLUG}/lines/Fire").json()["lines"]
    assert [l["item_code"] for l in r] == ["1", "2", "3", "5"] and r[0]["internal_ref"] == "2"
    assert client.get(f"/api/billing/{SLUG}/lines/Nope").status_code == 404


def test_installation_ra2_picks_up_only_what_was_not_billed_in_ra1(env):
    client, d, _ = prepared(env)                       # butterfly valve: BOQ 4, installed 50% = 2
    ra1 = client.post(f"/api/billing/{SLUG}/ra", json={"service": "Fire", "series": "install"}).json()["id"]
    assert line(client.get(f"/api/billing/{SLUG}/ra/{ra1}").json(), "2")["install"]["this"] == pytest.approx(2)
    client.post(f"/api/billing/{SLUG}/ra/{ra1}/lock")
    progress(client, "2", 0.75)                        # engineer moves it to 75% = 3
    ra2 = client.post(f"/api/billing/{SLUG}/ra", json={"service": "Fire", "series": "install"}).json()["id"]
    c = line(client.get(f"/api/billing/{SLUG}/ra/{ra2}").json(), "2")["install"]
    assert (c["prev"], c["utd"], c["this"]) == (2, pytest.approx(3), pytest.approx(1))
    assert c["amount"] == pytest.approx(5000 * 1 * .3)
    progress(client, "2", 1.0)                         # finished: 4 of 4, RA 2 now carries 2
    c = line(client.get(f"/api/billing/{SLUG}/ra/{ra2}").json(), "2")["install"]
    assert c["this"] == pytest.approx(2)
    client.post(f"/api/billing/{SLUG}/ra/{ra2}/lock")
    ra3 = client.post(f"/api/billing/{SLUG}/ra", json={"service": "Fire", "series": "install"}).json()["id"]
    assert all(l["item_code"] != "2" for l in client.get(f"/api/billing/{SLUG}/ra/{ra3}").json()["lines"])   # fully billed: nothing left
