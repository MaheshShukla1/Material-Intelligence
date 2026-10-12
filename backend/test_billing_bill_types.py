"""Supply only / Installation only / Together, chosen on EACH RA (the service's mode is only a default).
The real scenario: a service that is billed together one month and as separate invoices the next."""
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import openpyxl
import pytest

from backend import billing
from backend.test_billing import env, progress, SLUG, line, setup_fire, _recalc  # noqa: F401  (env is a fixture)
from backend.test_billing_challan import cs, ch, new_ra, bill, B, _find, _row_of

# Fire BOQ (see test_billing.make_workbook): line "2" Butterfly valve 150mm, 4 Nos @ 5000;
# line "3" GI pipe 100mm, 100 Rmt @ 800.


def lock(client, ra):
    r = client.post(f"{B}/ra/{ra}/lock")
    assert r.status_code == 200, r.text


def overview(client):
    return client.get(B).json()


def svc(client, name="Fire"):
    return next(x for x in overview(client)["services"] if x["service"] == name)


def rows(client):
    return {(r["series"], r["no"]): r for r in overview(client)["ras"]}


def test_the_new_ra_picker_offers_every_bill_type_whatever_the_default_is(env):
    client, _, _ = env
    cs(client, mode="together", sup=70, ins=30)
    can = svc(client)["can"]
    assert list(can) == ["supply", "install", "together"]                    # always in this order
    assert all(v["blocked"] is None and v["has_draft"] is False and v["next_no"] == 1 for v in can.values())
    assert list(svc(client)["ready"]) == ["together"]                         # the home figure still follows the default
    cs(client, mode="separate", sup=70, ins=30)
    assert list(svc(client)["ready"]) == ["supply", "install"]


def test_a_type_is_blocked_only_by_a_missing_percent_or_a_draft_on_the_same_stage(env):
    client, _, _ = env
    client.post(f"{B}/setup", json={"service": "Fire", "mode": "separate", "supply_pct": 70, "install_pct": None,
                                    "supply_source": "challan"})
    can = svc(client)["can"]
    assert can["supply"]["blocked"] is None
    assert "Installation %" in can["install"]["blocked"] and "Installation %" in can["together"]["blocked"]
    cs(client, sup=70, ins=30)
    new_ra(client, "supply")
    can = svc(client)["can"]
    assert can["supply"]["has_draft"] and can["together"]["has_draft"] and not can["install"]["has_draft"]
    assert can["together"]["draft_label"] == "Supply RA 1"
    r = client.post(f"{B}/ra", json={"service": "Fire", "series": "together"})
    assert r.status_code == 409 and "Supply draft is already open" in r.json()["detail"]
    assert client.post(f"{B}/ra", json={"service": "Fire", "series": "install"}).status_code == 200


def test_mixed_history_numbers_follow_the_stages_and_nothing_collides(env):
    client, _, _ = env
    cs(client, mode="together", sup=70, ins=30)
    ch(client, "614", "2026-07-04", [("2", 2)])
    r1 = new_ra(client, "together")                      # RA 1: both in one invoice
    lock(client, r1)
    ch(client, "620", "2026-08-04", [("2", 1)], ra=2)
    r2 = new_ra(client, "supply")                        # RA 2: supply invoice only
    r3 = new_ra(client, "install")                       # RA 2: installation invoice on its own
    assert (bill(client, r2)["ra"]["no"], bill(client, r3)["ra"]["no"]) == (2, 2)
    lock(client, r2)
    lock(client, r3)
    r4 = new_ra(client, "together")                      # RA 3 carries both stages again
    assert bill(client, r4)["ra"]["no"] == 3
    got = rows(client)
    assert set(got) == {("together", 1), ("supply", 2), ("install", 2), ("together", 3)}


def test_together_after_uneven_supply_and_install_history_takes_the_higher_number(env):
    client, _, _ = env
    cs(client, sup=70, ins=30)
    ch(client, "614", "2026-07-04", [("2", 1)])
    for _ in range(2):
        lock(client, new_ra(client, "install"))           # installation RA 1, RA 2 ...
    lock(client, new_ra(client, "supply"))                # ... supply only has RA 1
    r = new_ra(client, "together")
    assert bill(client, r)["ra"]["no"] == 3               # max(supply 2, install 3): never reuses a stage's number


def test_cumulative_qty_chains_across_bill_types(env):
    """Together RA 1 bills 2 valves supplied + installation; a later Supply-only RA must start from those 2, not from 0."""
    client, _, _ = env
    cs(client, mode="together", sup=70, ins=30)
    ch(client, "614", "2026-07-04", [("2", 2)])
    progress(client, "2", 0.25)                           # 1 of 4 installed
    r1 = new_ra(client, "together")
    b1 = bill(client, r1)
    assert line(b1, "2")["supply"]["this"] == 2 and line(b1, "2")["install"]["this"] == pytest.approx(1)
    lock(client, r1)
    ch(client, "620", "2026-08-04", [("2", 1)], ra=2)
    r2 = new_ra(client, "supply")
    b2 = bill(client, r2)
    s2 = line(b2, "2")["supply"]
    assert (s2["prev"], s2["utd"], s2["this"]) == (2, 3, 1)         # chained over the Together RA
    assert "install" not in line(b2, "2") and b2["ra"]["stages"] == ["supply"]
    progress(client, "2", 0.75)                           # 3 of 4 installed
    lock(client, r2)
    r3 = new_ra(client, "install")
    i3 = line(bill(client, r3), "2")["install"]
    assert (i3["prev"], i3["this"]) == (pytest.approx(1), pytest.approx(2))      # chained over the Together RA, not over the Supply RA


def test_up_to_date_and_reopen_follow_the_stage(env):
    client, _, _ = env
    cs(client, mode="together", sup=70, ins=30)
    ch(client, "614", "2026-07-04", [("2", 2)])
    r1 = new_ra(client, "together")
    lock(client, r1)
    ch(client, "620", "2026-08-04", [("2", 1)], ra=2)
    r2 = new_ra(client, "supply")
    lock(client, r2)
    # RA 1 (Together) cannot be reopened while the later Supply RA carries the same stage ...
    r = client.post(f"{B}/ra/{r1}/reopen")
    assert r.status_code == 409
    # ... the later one can, and then RA 1 stays locked
    assert client.post(f"{B}/ra/{r2}/reopen").status_code == 200
    got = rows(client)
    assert got[("supply", 2)]["status"] == "draft" and got[("together", 1)]["status"] == "locked"
    # "up to date" of a draft counts every locked RA that shares a stage with it
    assert got[("supply", 2)]["up_to_date"] == pytest.approx(got[("together", 1)]["this_total"] + got[("supply", 2)]["this_total"])


def test_installation_ra_after_a_together_ra_reopens_independently_of_the_supply_stage(env):
    client, _, _ = env
    cs(client, sup=70, ins=30)
    ch(client, "614", "2026-07-04", [("2", 2)])
    progress(client, "2", 0.25)
    r1 = new_ra(client, "together")
    lock(client, r1)
    r2 = new_ra(client, "install")
    lock(client, r2)
    r3 = new_ra(client, "supply")                        # supply chain is now ahead of nothing: only the Together RA is before it
    lock(client, r3)
    assert client.post(f"{B}/ra/{r2}/reopen").status_code == 200          # installation RA 2: no later RA on the install stage
    assert client.post(f"{B}/ra/{r1}/reopen").status_code == 409          # Together RA 1: later RAs exist on both stages


def test_change_of_default_after_ras_exist_does_not_touch_them(env):
    client, _, _ = env
    cs(client, mode="separate", sup=70, ins=30)
    lock(client, new_ra(client, "supply"))
    r = client.post(f"{B}/setup", json={"service": "Fire", "mode": "together", "supply_pct": 70, "install_pct": 30})
    assert r.status_code == 200, r.text
    assert list(svc(client)["ready"]) == ["together"] and ("supply", 1) in rows(client)
    # where the supply quantity comes from still cannot change once a Supply RA is locked
    r = client.post(f"{B}/setup", json={"service": "Fire", "mode": "together", "supply_pct": 70, "install_pct": 30, "supply_source": "stock"})
    assert r.status_code == 409


def test_open_challan_number_follows_a_locked_together_ra(env):
    client, _, _ = env
    cs(client, mode="together", sup=70, ins=30)
    ch(client, "614", "2026-07-04", [("2", 2)])
    lock(client, new_ra(client, "together"))
    j = client.get(f"{B}/challans/Fire").json()
    assert j["last_locked"] == 1 and j["open_no"] == 2
    r = ch(client, "620", "2026-08-04", [("2", 1)], ra=1)                  # RA 1 is submitted: frozen
    assert r.status_code == 409


def _book(client, ra, tmp, name="b.xlsx"):
    r = client.get(f"{B}/download", params={"ras": ra})
    assert r.status_code == 200, r.text
    p = tmp / name
    p.write_bytes(r.content)
    return openpyxl.load_workbook(p), r


def test_each_bill_type_downloads_as_its_own_invoice_with_the_right_sheets_and_wording(env):
    client, _, tmp = env
    cs(client, mode="together", sup=70, ins=30)
    ch(client, "614", "2026-07-04", [("2", 2)])
    progress(client, "2", 0.25)
    r1 = new_ra(client, "together")
    lock(client, r1)
    ch(client, "620", "2026-08-04", [("2", 1)], ra=2)
    progress(client, "2", 0.75)
    r2 = new_ra(client, "supply")
    r3 = new_ra(client, "install")
    wb1, resp1 = _book(client, r1, tmp, "t.xlsx")
    wb2, resp2 = _book(client, r2, tmp, "s.xlsx")
    wb3, resp3 = _book(client, r3, tmp, "i.xlsx")
    assert "Fire_Supply+Installation_RA-01.xlsx" in resp1.headers["content-disposition"]
    assert "Fire_Supply_RA-02_DRAFT" in resp2.headers["content-disposition"]
    assert "Fire_Installation_RA-02_DRAFT" in resp3.headers["content-disposition"]
    # supply-only carries the challan sheets, installation-only does not
    assert wb2.sheetnames == ["RA-02 INVOICE", "Annexure B-Summary BOQ", "Annexure B-BOQ", "CHALLAN ABTRACT", "CHALLAN"]
    assert wb3.sheetnames == ["RA-02 INVOICE", "Annexure B-Summary BOQ", "Annexure B-BOQ"]
    assert "Supply of" in wb2["RA-02 INVOICE"]["B14"].value and "Installation of" in wb3["RA-02 INVOICE"]["B14"].value
    assert "Supply and Installation of" in wb1["RA-01 INVOICE"]["B14"].value
    # the previous-bill quantity of the Supply-only RA 2 comes from the Together RA 1 (2), this bill is 1
    bq = _recalc(tmp / "s.xlsx", tmp / "rc_s")["Annexure B-BOQ"]
    r = _row_of(bq, "Butterfly valve 150mm")
    assert (bq.cell(r, 8).value, bq.cell(r, 9).value, bq.cell(r, 10).value) == (2, 1, 3)       # previous, this, total
    bi = _recalc(tmp / "i.xlsx", tmp / "rc_i")["Annexure B-BOQ"]
    r = _row_of(bi, "Butterfly valve 150mm")
    assert (bi.cell(r, 8).value, bi.cell(r, 9).value, bi.cell(r, 10).value) == (pytest.approx(1), pytest.approx(2), pytest.approx(3))


def test_download_dialog_data_lists_every_type_so_the_ticks_can_follow(env):
    client, _, _ = env
    cs(client, mode="together", sup=70, ins=30)
    ch(client, "614", "2026-07-04", [("2", 2)])
    lock(client, new_ra(client, "together"))
    new_ra(client, "supply")
    new_ra(client, "install")
    ras = overview(client)["ras"]
    assert sorted((r["series"], r["no"], r["status"]) for r in ras) == [
        ("install", 2, "draft"), ("supply", 2, "draft"), ("together", 1, "locked")]
    assert {r["series"]: r["stages"] for r in ras} == {"together": ["supply", "install"], "supply": ["supply"], "install": ["install"]}


def test_combined_export_handles_a_mixed_history(env):
    client, _, tmp = env
    cs(client, mode="together", sup=70, ins=30)
    ch(client, "614", "2026-07-04", [("2", 2)])
    progress(client, "2", 0.25)
    lock(client, new_ra(client, "together"))
    ch(client, "620", "2026-08-04", [("2", 1)], ra=2)
    progress(client, "2", 0.75)
    new_ra(client, "supply")
    new_ra(client, "install")
    r = client.get(f"{B}/export")
    assert r.status_code == 200, r.text
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    names = wb.sheetnames
    for n in ("Fire Supply RA-01", "Fire Install RA-01", "Fire Supply RA-02", "Fire Install RA-02"):
        assert n in names, names
    assert sum(1 for n in names if n.endswith("INVOICE")) == 3          # one Tax Invoice per RA (Together, Supply, Installation)
    reg = wb["RA Register"]
    got = [(reg.cell(r, 2).value, reg.cell(r, 3).value, reg.cell(r, 5).value) for r in range(5, 9)]
    assert got == [("Supply", "RA-01", "Invoiced"), ("Supply", "RA-02", "Draft"),          # the register lists the stages, a Together RA is in both
                   ("Installation", "RA-01", "Invoiced"), ("Installation", "RA-02", "Draft")]


def test_overview_tells_each_ra_what_each_stage_is_worth_so_a_zero_installation_can_be_flagged(env):
    """A Together RA created before any installation was logged: Supply has money, Installation is 0 -- the list says so
    (draft is live; the invoiced one is read from its frozen snapshot)."""
    client, _, _ = env
    cs(client, mode="together", sup=70, ins=30)
    ch(client, "614", "2026-07-04", [("2", 2)])
    r1 = new_ra(client, "together")
    st = rows(client)[("together", 1)]["stage_this"]
    assert st["supply"] > 0 and st["install"] == 0
    lock(client, r1)
    st = rows(client)[("together", 1)]["stage_this"]
    assert st["supply"] > 0 and st["install"] == 0 and rows(client)[("together", 1)]["status"] == "locked"
    # installation is then logged and billed as its own RA: that RA is worth money, the frozen one stays as it was
    progress(client, "2", 0.5)
    r2 = new_ra(client, "install")
    assert rows(client)[("install", 2)]["stage_this"]["install"] > 0
    assert rows(client)[("together", 1)]["stage_this"]["install"] == 0


def test_the_register_and_the_draft_banner_say_invoiced_not_locked(env, tmp_path):
    client, _, tmp = env
    cs(client, mode="together", sup=70, ins=30)
    ch(client, "614", "2026-07-04", [("2", 2)])
    r1 = new_ra(client, "together")
    lock(client, r1)
    ch(client, "620", "2026-08-04", [("2", 1)], ra=2)
    r2 = new_ra(client, "supply")
    wb, _ = _book(client, r2, tmp, "s.xlsx")
    assert "not invoiced yet" in str(wb["RA-02 INVOICE"]["A1"].value) + " ".join(str(c.value) for row in wb["RA-02 INVOICE"].iter_rows(max_row=8) for c in row if c.value)
    resp = client.get(f"{B}/download", params={"combined": "1"})
    reg = openpyxl.load_workbook(io.BytesIO(resp.content))["RA Register"]
    text = " ".join(str(c.value) for row in reg.iter_rows() for c in row if c.value is not None)
    assert "Invoiced" in text and "Locked" not in text
