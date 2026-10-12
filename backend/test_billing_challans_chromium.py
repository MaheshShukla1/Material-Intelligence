import sys, json, re, os
from playwright.sync_api import sync_playwright
"""Chromium test of the challan flow at desktop (1366) and phone (390) widths, against a SCRATCH copy of the
app (real Thoth BOQ). It drives the real UI: import the user's real HVAC invoice workbook, add a challan by
search, create the RA, download the invoice workbook and open it.
env: BASE (default http://localhost:8765), ROOT (scratch app root), INVOICE (HVAC invoice .xlsx), OUT (screenshots)."""
import urllib.request as u
BASE = os.environ.get("BASE", "http://localhost:8765")
ROOT = os.environ["ROOT"]
INVOICE = os.environ["INVOICE"]
OUT = os.environ.get("OUT", "/tmp/billing_challan_shots")
os.makedirs(OUT, exist_ok=True)
errs = []
def post(path, body):
    return u.urlopen(u.Request(BASE + path, data=json.dumps(body).encode(), headers={"Content-Type": "application/json"}))
def reset():
    f = os.path.join(ROOT, "data", "projects", "thoth-mall", "billing.json")
    if os.path.exists(f): os.remove(f)
    post("/api/billing/thoth-mall/setup", {"service": "HVAC", "mode": "separate", "supply_pct": 75, "install_pct": 25, "supply_source": "challan", "auto_draft": False})
    post("/api/billing/thoth-mall/setup", {"service": "Fire", "mode": "separate", "supply_pct": 70, "install_pct": 30, "auto_draft": False})
def noscroll(pg, label):
    assert 0 == pg.evaluate("document.documentElement.scrollWidth - document.documentElement.clientWidth"), label + ": page scrolls sideways"
with sync_playwright() as p:
    b = p.chromium.launch()
    for name, (w, h) in {"desk": (1366, 900), "phone": (390, 844)}.items():
        reset()
        ctx = b.new_context(viewport={"width": w, "height": h}, device_scale_factor=1, has_touch=(name == "phone"), is_mobile=(name == "phone"), accept_downloads=True)
        pg = ctx.new_page()
        pg.on("pageerror", lambda e, n=name: errs.append(("pageerror", n, str(e))))
        pg.on("console", lambda m, n=name: errs.append(("console", n, m.text)) if m.type == "error" else None)
        pg.goto(BASE); pg.wait_for_timeout(800)
        pg.click('.spnav button[data-v="billing"]'); pg.wait_for_selector(".bl-kpis")
        pg.screenshot(path=f"{OUT}/{name}_1_home.png", full_page=True); noscroll(pg, name + " home")
        pg.click('.bl-link[data-ch="HVAC"]'); pg.wait_for_selector(".bl-none")
        pg.screenshot(path=f"{OUT}/{name}_2_empty.png", full_page=True); noscroll(pg, name + " empty")
        # ---- import the real workbook through the UI
        pg.click("#bl-chimp"); pg.wait_for_selector("#bl-file")
        pg.set_input_files("#bl-file", INVOICE); pg.wait_for_selector(".bl-imp")
        assert "1,88,21,413" in pg.inner_text("#bl-modal"), "preview total"
        pg.screenshot(path=f"{OUT}/{name}_3_preview.png")
        assert pg.evaluate("(()=>{const m=document.querySelector('#bl-modal .card');return m.scrollWidth<=m.clientWidth+1})()"), name + ": import modal overflows sideways"
        pg.click("#bl-modal .mact .primary"); pg.wait_for_selector(".bl-ch")
        assert pg.locator(".bl-ch").count() == 12
        pg.screenshot(path=f"{OUT}/{name}_4_challans.png", full_page=True); noscroll(pg, name + " challans")
        pg.locator(".bl-ch .no").first.click(); pg.wait_for_timeout(150)
        pg.screenshot(path=f"{OUT}/{name}_5_expanded.png", full_page=True)
        # ---- add a challan by search
        pg.click("#bl-chadd"); pg.wait_for_selector("#bl-cno")
        pg.fill("#bl-cno", "700"); pg.fill("#bl-cfind", "horizontal endsuction pump"); pg.wait_for_selector("#bl-cfound .bl-hit")
        pg.screenshot(path=f"{OUT}/{name}_6_add_search.png")
        pg.locator("#bl-cfound .bl-hit").first.click(); pg.wait_for_selector("#bl-clist [data-q]")
        assert pg.evaluate("document.activeElement && document.activeElement.dataset.q === '0'"), "qty box not focused"
        pg.keyboard.type("1"); pg.screenshot(path=f"{OUT}/{name}_7_add_line.png")
        assert pg.evaluate("(()=>{const m=document.querySelector('#bl-modal .card');return m.scrollWidth<=m.clientWidth+1})()"), name + ": challan modal overflows sideways"
        pg.click("#bl-modal .mact .primary"); pg.wait_for_function("document.querySelectorAll('.bl-ch').length===13")
        # ---- re-tag the new challan to RA-02 -> the open RA keeps 12
        sel = pg.locator(".bl-ch", has_text="Challan 700").locator(".bl-retag")
        sel.select_option("2"); pg.wait_for_function("[...document.querySelectorAll('.bl-h3')].some(h=>/RA-02 · later/.test(h.textContent))")
        pg.screenshot(path=f"{OUT}/{name}_8_retag.png", full_page=True)
        # ---- RA from the challans
        pg.click("#bl-back"); pg.wait_for_selector(".bl-kpis")
        pg.click("#bl-new"); pg.wait_for_selector("#bl-modal .bl-opt")
        pg.click('#bl-modal [data-svc="HVAC"]'); pg.click('#bl-modal .bl-opt[data-series="supply"]')
        assert "challans" in pg.inner_text("#bl-modal").lower()
        pg.click("#bl-modal .mact .primary"); pg.wait_for_selector(".bl-lines")
        pg.wait_for_timeout(300)
        assert pg.locator(".bl-lr").count() == 45, pg.locator(".bl-lr").count()
        pg.screenshot(path=f"{OUT}/{name}_9_ra.png", full_page=True); noscroll(pg, name + " ra")
        pg.locator(".bl-lr .d").first.click(); pg.wait_for_timeout(150)
        pg.screenshot(path=f"{OUT}/{name}_10_ra_expand.png", full_page=True)
        pg.click("#bl-dl"); pg.wait_for_selector("#bl-modal .bl-dl")
        assert pg.inner_text("#bl-dl-go") == "Download .xlsx", pg.inner_text("#bl-dl-go")
        with pg.expect_download() as dl:
            pg.click("#bl-dl-go")
        path = os.path.join(OUT, name + "_invoice.xlsx"); dl.value.save_as(path)
        import openpyxl
        wb = openpyxl.load_workbook(path)
        assert wb.sheetnames == ["RA-01 INVOICE", "Annexure B-Summary BOQ", "Annexure B-BOQ", "CHALLAN ABTRACT", "CHALLAN"], wb.sheetnames
        assert re.search(r"HVAC_Supply_RA-01_DRAFT\.xlsx$", dl.value.suggested_filename), dl.value.suggested_filename
        # ---- setup screen has the new column
        pg.click("#bl-back"); pg.wait_for_selector(".bl-kpis"); pg.click("#bl-setup"); pg.wait_for_selector(".bl-sets")
        pg.screenshot(path=f"{OUT}/{name}_11_setup.png", full_page=True); noscroll(pg, name + " setup")
        pg.click('.spnav button[data-v="forecast"]'); pg.wait_for_timeout(300)
        assert pg.evaluate("document.getElementById('billing').hidden")
        ctx.close()
    b.close()
assert not errs, errs
print("billing challans chromium test: OK")
