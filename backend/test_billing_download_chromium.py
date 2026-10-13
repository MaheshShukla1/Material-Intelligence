import sys, json, re, os, io, zipfile
from playwright.sync_api import sync_playwright
import openpyxl
"""Chromium test of the Download dialog, the invoice-details form and the invoice strip at desktop (1366) and phone
(390) widths, against a SCRATCH copy of the app (real Thoth BOQ, the user's real HVAC challans). It drives the real UI
and opens every downloaded file.
env: BASE (default http://localhost:8765), ROOT (scratch app root), INVOICE (HVAC invoice .xlsx), OUT (screenshots)."""
import urllib.request as u
BASE = os.environ.get("BASE", "http://localhost:8765")
ROOT = os.environ["ROOT"]
INVOICE = os.environ["INVOICE"]
OUT = os.environ.get("OUT", "/tmp/billing_download_shots")
os.makedirs(OUT, exist_ok=True)
errs = []
P = "/api/billing/thoth-mall"
def post(path, body):
    return json.load(u.urlopen(u.Request(BASE + path, data=json.dumps(body).encode(), headers={"Content-Type": "application/json"})))
def get(path):
    return json.load(u.urlopen(BASE + path))
def importer():
    import uuid
    b = uuid.uuid4().hex
    raw = open(INVOICE, "rb").read()
    parts = [f'--{b}\r\nContent-Disposition: form-data; name="commit"\r\n\r\ntrue\r\n', f'--{b}\r\nContent-Disposition: form-data; name="lock_upto"\r\n\r\n0\r\n',
             f'--{b}\r\nContent-Disposition: form-data; name="file"; filename="i.xlsx"\r\nContent-Type: application/octet-stream\r\n\r\n']
    body = "".join(parts[:2]).encode() + parts[2].encode() + raw + f"\r\n--{b}--\r\n".encode()
    u.urlopen(u.Request(BASE + P + "/challans/HVAC/import", data=body, headers={"Content-Type": f"multipart/form-data; boundary={b}"}))
def reset():
    f = os.path.join(ROOT, "data", "projects", "thoth-mall", "billing.json")
    if os.path.exists(f): os.remove(f)
    post(P + "/setup", {"service": "HVAC", "mode": "separate", "supply_pct": 75, "install_pct": 25, "supply_source": "challan", "auto_draft": False})
    post(P + "/setup", {"service": "Fire", "mode": "separate", "supply_pct": 70, "install_pct": 30, "auto_draft": False})
    importer()
    r1 = post(P + "/ra", {"service": "HVAC", "series": "supply"})["id"]
    post(P + "/ra", {"service": "Fire", "series": "supply"})
    post(P + "/ra", {"service": "HVAC", "series": "install"})
    return r1
def noscroll(pg, label):
    assert 0 == pg.evaluate("document.documentElement.scrollWidth - document.documentElement.clientWidth"), label + ": page scrolls sideways"
def fits(pg, label):
    assert pg.evaluate("(()=>{const m=document.querySelector('#bl-modal .card');const r=m.getBoundingClientRect();return m.scrollWidth<=m.clientWidth+1 && r.left>=0 && r.right<=innerWidth+1 && r.top>=0 && r.bottom<=innerHeight+1})()"), label + ": the dialog does not fit the screen / overflows sideways"
def save(dl, name):
    path = os.path.join(OUT, name); dl.value.save_as(path); return path
with sync_playwright() as p:
    b = p.chromium.launch()
    for name, (w, h) in {"desk": (1366, 900), "phone": (390, 844)}.items():
        ra1 = reset()
        ctx = b.new_context(viewport={"width": w, "height": h}, device_scale_factor=1, has_touch=(name == "phone"), is_mobile=(name == "phone"), accept_downloads=True)
        pg = ctx.new_page()
        pg.on("pageerror", lambda e, n=name: errs.append(("pageerror", n, str(e))))
        pg.on("console", lambda m, n=name: errs.append(("console", n, m.text)) if m.type == "error" else None)
        pg.goto(BASE); pg.wait_for_timeout(800)
        pg.click('.spnav button[data-v="billing"]'); pg.wait_for_selector(".bl-kpis")
        pg.screenshot(path=f"{OUT}/{name}_1_home.png", full_page=True); noscroll(pg, name + " home")
        assert pg.locator("#bl-dl").count() == 1 and pg.locator("#bl-xl").count() == 0
        # ---- the dialog
        pg.click("#bl-dl"); pg.wait_for_selector("#bl-modal .bl-dl")
        pg.wait_for_timeout(200)
        assert pg.locator("#bl-modal [data-ra]:checked").count() == 0 and pg.locator('#bl-modal [data-dl="pb"]:checked').count() == 0 and pg.is_disabled("#bl-dl-go"), "from home nothing is ticked"
        pg.click('#bl-modal [data-q="latest"]'); pg.wait_for_timeout(100)
        pg.screenshot(path=f"{OUT}/{name}_2_dialog.png"); fits(pg, name + " dialog")
        n = pg.locator("#bl-modal .bl-fi").count()
        assert n == 3, n
        assert pg.inner_text("#bl-dl-go") == "Download 3 files (.zip)"
        txt = pg.inner_text("#bl-modal")
        assert "stock register" in txt and "DRAFT" in txt and "Blank on the Tax Invoice sheet" in txt
        if name == "phone":
            # the Download button stays visible without scrolling the dialog body
            assert pg.evaluate("(()=>{const r=document.getElementById('bl-dl-go').getBoundingClientRect();return r.bottom<=innerHeight && r.height>=44})()"), "phone: Download button off-screen or too small"
            pg.locator("#bl-modal .bl-mbody").evaluate("e=>e.scrollTo(0,e.scrollHeight)"); pg.wait_for_timeout(100)
            pg.screenshot(path=f"{OUT}/{name}_2b_dialog_bottom.png")
        # touch targets
        sel = "#bl-modal .bl-dr, #bl-modal .bl-dsh" + (", #bl-modal .bl-ck, #bl-modal .bl-seg button, #bl-modal .bl-qk .btn, #bl-modal .bl-wn .btn, #bl-modal .mact .btn" if name == "phone" else "")
        small = pg.evaluate("(s)=>[...document.querySelectorAll(s)].filter(e=>e.getBoundingClientRect().height<43).length", sel)
        assert small == 0, f"{small} tick rows under 43px"
        # ---- the real zip: HVAC (challans) has 5 sheets, Fire (stock) 3, install 3
        with pg.expect_download() as dl:
            pg.click("#bl-dl-go")
        path = save(dl, name + "_dl.zip")
        assert dl.value.suggested_filename.endswith(".zip"), dl.value.suggested_filename
        z = zipfile.ZipFile(path); names = sorted(z.namelist())
        assert [re.sub(r"^thoth-mall_", "", n) for n in names] == ["Fire_Supply_RA-01_DRAFT.xlsx", "HVAC_Installation_RA-01_DRAFT.xlsx", "HVAC_Supply_RA-01_DRAFT.xlsx"], names
        sh = {n: openpyxl.load_workbook(io.BytesIO(z.read(n))).sheetnames for n in names}
        assert sh[[n for n in names if "HVAC_Supply" in n][0]] == ["RA-01 INVOICE", "Annexure B-Summary BOQ", "Annexure B-BOQ", "CHALLAN ABTRACT", "CHALLAN"]
        assert sh[[n for n in names if "Fire_Supply" in n][0]] == ["RA-01 INVOICE", "Annexure B-Summary BOQ", "Annexure B-BOQ"]
        assert sh[[n for n in names if "HVAC_Installation" in n][0]] == ["RA-01 INVOICE", "Annexure B-Summary BOQ", "Annexure B-BOQ"]
        pg.wait_for_function("!document.getElementById('bl-modal')")
        # ---- fill the details through the UI
        pg.click("#bl-dl"); pg.wait_for_selector("#bl-modal .bl-dl")
        pg.click('#bl-modal [data-q="latest"]'); pg.wait_for_timeout(100)
        pg.click('#bl-modal [data-act="details"]'); pg.wait_for_selector('#bl-modal [data-m="client"]')
        pg.wait_for_timeout(150)
        pg.screenshot(path=f"{OUT}/{name}_3_details.png"); fits(pg, name + " details")
        vals = {"contractor": "Vedant Engitech And Infrastructure Pvt. Ltd.", "supplier_gstin": "24AADCV2799M1ZP", "supplier_pan": "AADCV2799M", "supplier_state": "GUJARAT",
                "supplier_state_code": "24", "supplier_category": "MEPF", "client": "Thoth Mall Pvt Ltd", "client_address": "Plot 5, SG Highway, Ahmedabad", "place_of_supply": "Gujarat",
                "client_gstin": "24ABCDE1234F1Z5", "client_state": "Gujarat", "client_state_code": "24", "project_name": "Thoth Mall", "wo_no": "WO/12", "sgst": "9", "cgst": "9"}
        for k, v in vals.items():
            pg.fill(f'#bl-modal [data-m="{k}"]', v)
        for i, svc in enumerate(["Electrical", "Fire", "HVAC", "Plumbing", "ELV"]):
            if pg.locator(f'#bl-modal [data-sac="{i}"]').count(): pg.fill(f'#bl-modal [data-sac="{i}"]', "995461")
        ids = pg.evaluate("[...document.querySelectorAll('#bl-modal [data-rno]')].map(e=>e.dataset.rno)")
        assert len(ids) == 3
        for j, rid in enumerate(ids):
            pg.fill(f'#bl-modal [data-rno="{rid}"]', f"VE/25-26/02{j}"); pg.fill(f'#bl-modal [data-rdt="{rid}"]', "2026-10-08")
        pg.click("#bl-isave"); pg.wait_for_selector("#bl-modal .bl-dl")
        pg.wait_for_timeout(200)
        assert pg.locator("#bl-modal .bl-wn.details").count() == 0, "details warning should be gone"
        pg.screenshot(path=f"{OUT}/{name}_4_after_details.png")
        # ---- ONE workbook, named by service
        pg.click('#bl-modal [data-mode="one"]'); pg.wait_for_timeout(100)
        assert pg.locator("#bl-modal .bl-fi").count() == 1
        pg.screenshot(path=f"{OUT}/{name}_5_one.png")
        with pg.expect_download() as dl:
            pg.click("#bl-dl-go")
        path = save(dl, name + "_one.xlsx")
        assert dl.value.suggested_filename.endswith(".xlsx") and "RA_invoices" in dl.value.suggested_filename
        wb = openpyxl.load_workbook(path)
        assert "HVAC RA01S INVOICE" in wb.sheetnames and "HVAC RA01S CH ABSTRACT" in wb.sheetnames and "HVAC RA01S CHALLAN" in wb.sheetnames, wb.sheetnames
        assert "Fire RA01 INVOICE" in wb.sheetnames and not any(s.startswith("Fire") and "CHALLAN" in s for s in wb.sheetnames), wb.sheetnames
        cv = wb["HVAC RA01S INVOICE"]
        assert cv["B2"].value == "Thoth Mall Pvt Ltd" and cv["B5"].value == "24ABCDE1234F1Z5" and cv["E7"].value == "HVAC" and cv["D14"].value == 995461
        assert cv["A41"].value == "For Vedant Engitech And Infrastructure Pvt. Ltd." and cv["B36"].value == "24AADCV2799M1ZP" and cv["E3"].value.startswith("VE/25-26/02")
        pg.wait_for_function("!document.getElementById('bl-modal')")
        # ---- the All services Excel alone
        pg.click("#bl-dl"); pg.wait_for_selector("#bl-modal .bl-dl")
        pg.uncheck('#bl-modal [data-dl="invoice"]'); pg.check('#bl-modal [data-dl="combined"]'); pg.wait_for_timeout(100)
        pg.screenshot(path=f"{OUT}/{name}_6_combined.png")
        with pg.expect_download() as dl:
            pg.click("#bl-dl-go")
        path = save(dl, name + "_all.xlsx")
        wb = openpyxl.load_workbook(path)
        assert "HVAC CHALLAN ABSTRACT" in wb.sheetnames and "HVAC CHALLAN" in wb.sheetnames and "Invoice (format pending)" not in wb.sheetnames, wb.sheetnames
        assert any(s.endswith("INVOICE") for s in wb.sheetnames)
        pg.wait_for_function("!document.getElementById('bl-modal')")
        # ---- RA screen: invoice strip saves, Download is this RA only
        pg.click(f'.bl-rar[data-ra="{ra1}"]'); pg.wait_for_selector("#bl-ino")
        pg.screenshot(path=f"{OUT}/{name}_7_ra.png", full_page=True); noscroll(pg, name + " ra")
        pg.fill("#bl-ino", "VE/25-26/099"); pg.fill("#bl-ida", "2026-10-09"); pg.press("#bl-ida", "Tab")
        # straight after typing: Download waits for the invoice no. to be saved
        pg.click("#bl-dl"); pg.wait_for_selector("#bl-modal .bl-dl")
        assert pg.locator("#bl-modal .bl-fi").count() == 1
        with pg.expect_download() as dl:
            pg.click("#bl-dl-go")
        path = save(dl, name + "_ra.xlsx")
        assert re.search(r"HVAC_Supply_RA-01_DRAFT\.xlsx$", dl.value.suggested_filename), dl.value.suggested_filename
        cv = openpyxl.load_workbook(path)["RA-01 INVOICE"]
        assert get(P + "/ra/" + ra1)["ra"]["inv_no"] == "VE/25-26/099"
        assert cv["E3"].value == "VE/25-26/099" and cv["E4"].value.date().isoformat() == "2026-10-09", (name, cv["E3"].value, cv["E4"].value)
        pg.wait_for_function("!document.getElementById('bl-modal')")
        # ---- setup card
        pg.click("#bl-back"); pg.wait_for_selector(".bl-kpis"); pg.click("#bl-setup"); pg.wait_for_selector(".bl-metac")
        pg.screenshot(path=f"{OUT}/{name}_8_setup.png", full_page=True)
        ctx.close()
    b.close()
assert not errs, errs
print("billing download chromium test: OK")
