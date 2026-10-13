"""Billing -- quantity ABOVE the BOQ is accepted line by line, billed as Extra, highlighted, never billed twice.
Scenario from the report (Thoth HVAC, Supply RA 1): challans total 2 Nos for two lines whose BOQ qty is 1. It used to show red CHECK
cards on the Billing home and bill only 1; typing 2 on the RA said "More than that cannot be billed"; Upload filled sheet refused it.
Same harness as test_billing_challan.py (Fire: valve BOQ 4 @ 5000, pump BOQ 2 @ 100000, GI pipe BOQ 100 @ 800; supply 70 / install 30)."""
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl
import pytest

from backend.test_billing import env, prepared, make_run, link, progress, line, STOCK1, SLUG   # noqa: F401
from backend.test_billing_challan import cs, ch, new_ra, bill, B
from backend.test_billing_invoice_cover import invoice_of

SP = f"/api/siteprogress/{SLUG}"


def accept(client, ra, items, accept=True):
    r = client.post(f"{B}/ra/{ra}/accept-excess", json={"items": items, "accept": accept})
    assert r.status_code == 200, r.text
    return r.json()


def over_valve(client):
    """Valve BOQ is 4; 6 delivered by challan: 2 above the BOQ."""
    cs(client)
    ch(client, "700", "2026-07-04", [("2", 3)])
    ch(client, "701", "2026-07-06", [("2", 3)])
    return new_ra(client)


def boq_ws(wb):
    return next(wb[n] for n in wb.sheetnames if "BOQ" in n and "Summary" not in n and "EXTRA" not in n.upper())


def this_cell(ws, desc):
    hdr = next(r for r in range(1, 20) if "THIS BILL QTY" in [ws.cell(r, c).value for c in range(1, 25)])
    col = next(c for c in range(1, 25) if ws.cell(hdr, c).value == "THIS BILL QTY")
    row = next(r for r in range(8, ws.max_row + 1) if ws.cell(r, 2).value == desc)
    return ws.cell(row, col)


# ----------------------------------------------------------------- not accepted: as before, plus a precise list
def test_not_accepted_the_bill_stays_at_the_boq_and_the_excess_is_listed_with_its_value(env):
    client, _, _ = env
    ra = over_valve(client)
    b = bill(client, ra)
    c = line(b, "2")["supply"]
    assert (c["utd"], c["this"]) == (4, 4) and "extra" not in c and "accepted" not in c
    assert len(b["excess"]) == 1
    x = b["excess"][0]
    assert (x["item_code"], x["stage"], x["boq_qty"], x["have"], x["extra"], x["source"]) == ("2", "supply", 4, 6, 2, "challan")
    assert x["amount"] == pytest.approx(5000 * 2 * .70)                   # real rate x extra x supply %
    assert b["extra_total"] == 0 and b["extra_lines"] == []
    assert b["this_total"] == pytest.approx(5000 * 4 * .70)               # unchanged: nothing is billed above the BOQ by itself


def test_billing_home_has_no_over_boq_card_only_a_count_on_that_ra(env):
    client, _, _ = env
    ra = over_valve(client)
    ov = client.get(B).json()
    assert [n for n in ov["notes"] if n["kind"] == "over"] == []
    row = next(r for r in ov["ras"] if r["id"] == ra)
    assert row["over_n"] == 1 and row["extra_n"] == 0


# ----------------------------------------------------------------- accept
def test_accept_bills_the_extra_at_the_same_rate_and_marks_it(env):
    client, d, _ = env
    ra = over_valve(client)
    r = accept(client, ra, [{"item_code": "2", "stage": "supply"}])
    assert r["extra_total"] == pytest.approx(5000 * 2 * .70) and r["excess"] == 0
    b = bill(client, ra)
    c = line(b, "2")["supply"]
    assert (c["utd"], c["this"], c["extra"], c["accepted"]) == (6, 6, 2, True)
    assert c["amount"] == pytest.approx(5000 * 6 * .70) and c["extra_amount"] == pytest.approx(5000 * 2 * .70)
    assert b["excess"] == [] and b["extra_total"] == pytest.approx(7000)
    assert [(x["item_code"], x["extra"], x["accepted"]) for x in b["extra_lines"]] == [("2", 2, True)]
    assert b["this_total"] == pytest.approx(5000 * 6 * .70)
    ov = client.get(B).json()
    assert next(r for r in ov["ras"] if r["id"] == ra)["extra_n"] == 1


def test_only_the_ticked_line_is_accepted(env):
    client, _, _ = env
    cs(client)
    ch(client, "700", "2026-07-04", [("2", 6), ("1", 3)])                # valve 6 of 4, pump 3 of 2
    ra = new_ra(client)
    assert {x["item_code"] for x in bill(client, ra)["excess"]} == {"1", "2"}
    accept(client, ra, [{"item_code": "1", "stage": "supply"}])
    b = bill(client, ra)
    assert line(b, "1")["supply"]["this"] == 3 and line(b, "2")["supply"]["this"] == 4
    assert [x["item_code"] for x in b["excess"]] == ["2"]


def test_back_to_boq_takes_it_back(env):
    client, _, _ = env
    ra = over_valve(client)
    accept(client, ra, [{"item_code": "2", "stage": "supply"}])
    accept(client, ra, [{"item_code": "2", "stage": "supply"}], accept=False)
    b = bill(client, ra)
    assert line(b, "2")["supply"]["this"] == 4 and len(b["excess"]) == 1 and b["extra_lines"] == []


def test_a_locked_ra_cannot_be_changed_and_an_unknown_line_is_refused(env):
    client, _, _ = env
    ra = over_valve(client)
    assert client.post(f"{B}/ra/{ra}/accept-excess", json={"items": [{"item_code": "999", "stage": "supply"}]}).status_code == 400
    assert client.post(f"{B}/ra/{ra}/accept-excess", json={"items": []}).status_code == 400
    client.post(f"{B}/ra/{ra}/lock")
    assert client.post(f"{B}/ra/{ra}/accept-excess", json={"items": [{"item_code": "2", "stage": "supply"}]}).status_code == 409


# ----------------------------------------------------------------- the next RA: never twice, never negative
def test_next_ra_starts_after_the_extra_so_it_is_never_billed_twice(env):
    client, _, _ = env
    ra1 = over_valve(client)
    accept(client, ra1, [{"item_code": "2", "stage": "supply"}])
    client.post(f"{B}/ra/{ra1}/lock")
    ch(client, "702", "2026-08-02", [("1", 1)])                          # something else arrives, so RA 2 has a reason to exist
    ra2 = new_ra(client)
    b = bill(client, ra2)
    assert all(x["item_code"] != "2" for x in b["lines"])                # valve: 6 already billed, 6 delivered: nothing more
    assert b["excess"] == [] and not any((w.get("msg") or "").startswith("2") for w in b["warnings"])
    assert not [w for w in b["warnings"] if "was billed earlier" in w["msg"]]            # no negative line
    ch(client, "703", "2026-08-09", [("2", 1)], ra=2)                     # one more valve: 7 delivered, 6 billed -> 1 above
    b = bill(client, ra2)
    assert [x["extra"] for x in b["excess"]] == [1]                       # asks again: the earlier accept does not cover new extra
    assert all(x["item_code"] != "2" for x in b["lines"])                 # still nothing new to bill on the valve: 6 billed, 6 allowed


# ----------------------------------------------------------------- typed quantity on the RA screen
def test_typing_above_the_boq_needs_the_accept_flag(env):
    client, _, _ = env
    cs(client)
    ra = new_ra(client)
    r = client.post(f"{B}/ra/{ra}/line", json={"item_code": "2", "stage": "supply", "qty": 6})
    assert r.status_code == 400 and "More than that cannot be billed" in r.json()["detail"]
    r = client.post(f"{B}/ra/{ra}/line", json={"item_code": "2", "stage": "supply", "qty": 6, "accept_excess": True})
    assert r.status_code == 200
    c = line(bill(client, ra), "2")["supply"]
    assert (c["this"], c["extra"], c["accepted"]) == (6, 2, True)


def test_back_to_boq_also_drops_a_typed_quantity_that_stood_only_because_of_the_accept(env):
    client, _, _ = env
    cs(client)
    ra = new_ra(client)
    client.post(f"{B}/ra/{ra}/line", json={"item_code": "2", "stage": "supply", "qty": 6, "accept_excess": True})
    accept(client, ra, [{"item_code": "2", "stage": "supply"}], accept=False)
    b = bill(client, ra)
    assert "2" not in {x["item_code"] for x in b["lines"]}                # no typed quantity left, no challan: the line is not in the draft
    assert b["extra_lines"] == [] and b["this_total"] == 0


# ----------------------------------------------------------------- installation (Site Progress above the BOQ)
def test_installation_above_the_boq_from_site_progress_is_accepted_the_same_way(env):
    client, d, _ = env
    cs(client)
    r = client.post(f"{SP}/progress/item-qty", json={"service": "Fire", "item_code": "3", "qty": 130})
    assert r.status_code == 200, r.text
    ra = new_ra(client, "install")
    b = bill(client, ra)
    assert line(b, "3")["install"]["this"] == 100                         # BOQ 100: capped
    x = b["excess"][0]
    assert (x["stage"], x["source"], x["extra"], x["have"]) == ("install", "progress", 30, 130)
    assert x["amount"] == pytest.approx(800 * 30 * .30)
    accept(client, ra, [{"item_code": "3", "stage": "install"}])
    b = bill(client, ra)
    c = line(b, "3")["install"]
    assert (c["this"], c["extra"]) == (130, 30) and c["extra_amount"] == pytest.approx(800 * 30 * .30) and b["excess"] == []


# ----------------------------------------------------------------- the Excel
def test_excel_highlights_the_accepted_extra_and_adds_one_sheet(env):
    client, d, tmp = env
    ra = over_valve(client)
    accept(client, ra, [{"item_code": "2", "stage": "supply"}])
    p = invoice_of(client, ra, tmp, "x.xlsx")
    wb = openpyxl.load_workbook(p)
    assert "Extra over BOQ" in wb.sheetnames and wb.sheetnames[-1] == "Extra over BOQ"
    cell = this_cell(boq_ws(wb), "Butterfly valve 150mm")
    assert cell.fill.fgColor.rgb.endswith("FFE699")
    assert cell.comment and "Extra over BOQ: 2" in cell.comment.text and "BOQ quantity of 4" in cell.comment.text
    assert "MIN(" not in str(cell.value) and "CHALLAN ABTRACT" in str(cell.value)            # the cap is lifted for this line only
    es = wb["Extra over BOQ"]
    rows = [[c.value for c in r] for r in es.iter_rows(min_row=8)]
    assert rows[0][1] == "Butterfly valve 150mm" and rows[0][3:6] == [4, 6, 2]
    assert rows[0][-1] == pytest.approx(7000) and rows[1][1] == "Total extra" and rows[1][-1] == pytest.approx(7000)


def test_excel_without_an_accept_is_exactly_as_before(env):
    client, d, tmp = env
    ra = over_valve(client)
    wb = openpyxl.load_workbook(invoice_of(client, ra, tmp, "y.xlsx"))
    assert "Extra over BOQ" not in wb.sheetnames
    cell = this_cell(boq_ws(wb), "Butterfly valve 150mm")
    assert "MIN(D" in str(cell.value) and cell.comment is None
    assert not (cell.fill.fgColor.rgb or "").endswith("FFE699")


def test_excel_numbers_match_the_screen_when_recalculated(env):
    """The workbook's formulas, recalculated by LibreOffice, give the bill the screen shows."""
    from backend.test_billing import _recalc
    client, d, tmp = env
    ra = over_valve(client)
    accept(client, ra, [{"item_code": "2", "stage": "supply"}])
    p = invoice_of(client, ra, tmp, "z.xlsx")
    out = tmp / "re"
    out.mkdir()
    wb = _recalc(p, out)
    assert this_cell(boq_ws(wb), "Butterfly valve 150mm").value == pytest.approx(6)


def test_excel_of_the_next_ra_does_not_go_negative_on_a_line_that_billed_extra_before(env):
    """RA 1 billed 6 of a BOQ 4. In RA 2's workbook the formula MIN(D, challan total) - previous would read 4 - 6 = -2."""
    from backend.test_billing import _recalc
    client, d, tmp = env
    ra1 = over_valve(client)
    accept(client, ra1, [{"item_code": "2", "stage": "supply"}])
    client.post(f"{B}/ra/{ra1}/lock")
    ch(client, "702", "2026-08-02", [("1", 1)])
    ra2 = new_ra(client)
    p = invoice_of(client, ra2, tmp, "n.xlsx")
    wb = openpyxl.load_workbook(p)
    assert "MAX(D" in str(this_cell(boq_ws(wb), "Butterfly valve 150mm").value)
    out = tmp / "re2"
    out.mkdir()
    wb = _recalc(p, out)
    assert this_cell(boq_ws(wb), "Butterfly valve 150mm").value == pytest.approx(0)
    assert "Extra over BOQ" not in wb.sheetnames                           # nothing is extra in THIS bill


def test_dump_the_real_payloads_for_the_ui_test(env):
    """Only when MI_DUMP_FIXTURE=<path> is set: writes the REAL API payloads of this scenario, so the jsdom test
    (frontend/test_billing_extra_ui.mjs) runs on what the backend really sends, not on a hand-made shape."""
    import json, os
    out = os.environ.get("MI_DUMP_FIXTURE")
    if not out:
        pytest.skip("fixture dump not requested")
    client, d, tmp = env
    cs(client)
    ch(client, "700", "2026-07-04", [("2", 3), ("1", 3)])
    ch(client, "701", "2026-07-06", [("2", 3)])                                            # valve 6 of 4, pump 3 of 2
    ra = new_ra(client)
    fx = {"raid": ra, "before": bill(client, ra), "overview_before": client.get(B).json()}
    accept(client, ra, [{"item_code": "2", "stage": "supply"}])
    fx["after"], fx["overview_after"] = bill(client, ra), client.get(B).json()
    accept(client, ra, [{"item_code": "2", "stage": "supply"}], accept=False)
    ra2 = new_ra(client, "install")
    p = invoice_of(client, ra2, tmp, "u.xlsx")
    wb = openpyxl.load_workbook(p)
    ws = boq_ws(wb)
    this_cell(ws, "GI pipe 100mm").value = 130
    wb.save(p)
    with open(p, "rb") as fh:
        fx["upload_preview"] = client.post(f"{B}/import-sheet/Fire", files={"file": ("f.xlsx", fh)}, data={"commit": "false", "ra_id": ra2}).json()
    fx["raid2"] = ra2
    json.dump(fx, open(out, "w"))
