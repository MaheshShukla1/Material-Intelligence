"""Billing -- the Tax Invoice cover sheet (the user's invoice format), invoice details, the
per-RA invoice no./date, challan sheets named by service, and the one Download entry point
(zip / single workbook / sheet ticks). Same harness as test_billing_challan.py."""
import io
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl
import pytest

from backend import billing, billing_cover
from backend.test_billing import env, progress, SLUG, _recalc  # noqa: F401  (env is a fixture)
from backend.test_billing_challan import cs, ch, new_ra, bill, B, _find, _row_of

FULL = {"client": "Thoth Mall Pvt Ltd", "contractor": "Vedant Engitech And Infrastructure Pvt. Ltd.", "wo_no": "WO/12",
        "client_address": "Plot 5, SG Highway, Ahmedabad", "place_of_supply": "Gujarat", "client_gstin": "24ABCDE1234F1Z5",
        "client_state": "Gujarat", "client_state_code": "24", "project_name": "Thoth Mall",
        "supplier_gstin": "24AADCV2799M1ZP", "supplier_state": "GUJARAT", "supplier_state_code": "24",
        "supplier_pan": "AADCV2799M", "supplier_category": "MEPF", "sgst": 9, "cgst": 9,
        "sac": {"Fire": "995461"}}


def meta(c, **kw):
    r = c.post(f"{B}/meta", json=kw)
    assert r.status_code == 200, r.text
    return r.json()["meta"]


def full(c, **over):
    return meta(c, **{**FULL, **over})


def invoice_of(client, ra, tmp, name="i.xlsx"):
    r = client.get(f"{B}/invoice/{ra}")
    assert r.status_code == 200, r.text
    p = tmp / name
    p.write_bytes(r.content)
    return p


# ------------------------------------------------------------------ words / money
def test_amount_in_words_is_indian_and_matches_the_users_own_example():
    assert billing_cover.words(11215801) == "One Crore Twelve Lakh Fifteen Thousand Eight Hundred One Only."
    assert billing_cover.words(100000) == "One Lakh Only."
    assert billing_cover.words(0) == "Zero Only."
    assert billing_cover.words(119) == "One Hundred Nineteen Only."


def test_money_rounds_to_the_rupee_like_the_sheets_own_round_off():
    m = billing_cover.money(100000.50, 9, 9)
    assert m["gst"] == pytest.approx(18000.09)
    assert m["total"] == 118001.0 and m["round"] == pytest.approx(118001 - 118000.59)
    assert billing_cover.money(-1000.5, 0, 0)["total"] == -1001.0           # half away from zero, like Excel ROUND


# --------------------------------------------------------------- invoice details
def test_meta_saves_invoice_details_validates_gst_and_blank_means_not_set(env):
    client, _, _ = env
    m = full(client)
    assert m["supplier_gstin"] == "24AADCV2799M1ZP" and m["sgst"] == 9 and m["sac"] == {"Fire": "995461"}
    m = meta(client, client_gstin=" 24abcde1234f1z5 ", supplier_pan="aadcv2799m")
    assert (m["client_gstin"], m["supplier_pan"]) == ("24ABCDE1234F1Z5", "AADCV2799M")     # upper-cased
    assert client.post(f"{B}/meta", json={"sgst": 120}).status_code == 400
    assert client.post(f"{B}/meta", json={"cgst": -1}).status_code == 400
    assert client.post(f"{B}/meta", json={"sgst": "abc"}).status_code == 400
    assert meta(client, sgst="")["sgst"] is None and meta(client)["cgst"] == 9               # blank clears, others stay
    assert meta(client, sac={"Fire": ""})["sac"]["Fire"] == ""
    # the old three fields still behave as before
    assert meta(client, client="X")["client"] == "X"


def test_ra_invoice_no_and_date_save_validate_and_work_on_a_locked_ra(env):
    client, _, _ = env
    cs(client)
    ch(client, "614", "2026-07-04", [("1", 2)])
    ra = new_ra(client)
    r = client.post(f"{B}/ra/{ra}/invoice", json={"no": " VE/25-26/014 ", "date": "05.07.2026"})
    assert r.status_code == 200 and r.json()["inv_no"] == "VE/25-26/014" and r.json()["inv_date"] == "2026-07-05"
    assert client.post(f"{B}/ra/{ra}/invoice", json={"date": "not a date"}).status_code == 400
    assert client.post(f"{B}/ra/ra999/invoice", json={"no": "1"}).status_code == 404
    client.post(f"{B}/ra/{ra}/lock")
    r = client.post(f"{B}/ra/{ra}/invoice", json={"no": "VE/25-26/015"})       # the invoice is raised after locking
    assert r.status_code == 200 and r.json()["inv_no"] == "VE/25-26/015" and r.json()["inv_date"] == "2026-07-05"
    assert bill(client, ra)["ra"]["inv_no"] == "VE/25-26/015"
    ov = client.get(B).json()
    row = next(x for x in ov["ras"] if x["id"] == ra)
    assert "Invoice no." not in row["inv_missing"] and "Invoice date" not in row["inv_missing"]
    assert "Client GSTIN" in row["inv_missing"] and "SAC code of Fire" in row["inv_missing"] and "GST % (SGST / CGST)" in row["inv_missing"]
    full(client)
    assert next(x for x in client.get(B).json()["ras"] if x["id"] == ra)["inv_missing"] == []


# ------------------------------------------------------------------ the cover sheet
def test_draft_cover_is_the_users_format_filled_with_live_formulas(env):
    client, _, tmp = env
    cs(client)
    full(client)
    ch(client, "614", "2026-07-04", [("1", 2), ("2", 1)])
    ra = new_ra(client)
    client.post(f"{B}/ra/{ra}/invoice", json={"no": "VE/25-26/014", "date": "2026-07-05"})
    b = bill(client, ra)
    p = invoice_of(client, ra, tmp)
    wb = openpyxl.load_workbook(p)
    ws = wb["RA-01 INVOICE"]
    assert wb.sheetnames[0] == "RA-01 INVOICE" and wb.sheetnames[1] == "Annexure B-Summary BOQ"     # just before the annexures
    assert ws["E14"].value.startswith("='Annexure B-Summary BOQ'!E")                                    # follows the annexure
    assert ws["E31"].value == "=ROUND(E30,0)-E30" and ws["E24"].value == "=+E21*D24"
    assert "DRAFT" in ws["A1"].value
    assert ws["E4"].number_format == "dd-mm-yyyy"                                                       # date shows like the user's own sheet
    rc = _recalc(p, tmp / "rc")["RA-01 INVOICE"]
    # the user's own labels are untouched
    assert (rc["D1"].value, rc["A2"].value, rc["D10"].value, rc["E10"].value) == ("TAX INVOICE", "Company Name :", "SAC CODE", "AMOUNT")
    assert (rc["B2"].value, rc["B3"].value, rc["B4"].value, rc["B5"].value, rc["B6"].value) == (
        "Thoth Mall Pvt Ltd", "Plot 5, SG Highway, Ahmedabad", "Gujarat", "24ABCDE1234F1Z5", "Thoth Mall")
    assert (rc["B7"].value, rc["B8"].value, rc["B9"].value) == ("N", "Gujarat", "24")
    assert (rc["E3"].value, rc["E5"].value, rc["E6"].value, rc["E7"].value) == ("VE/25-26/014", "WO/12", 1, "Fire")
    assert rc["E4"].value.date().isoformat() == "2026-07-05"
    assert rc["B14"].value == "Supply of Fire Work" and rc["D14"].value == 995461 and rc["B15"].value == "As per Annexure"
    assert "Supply Amount" in rc["A22"].value
    amt = b["this_total"]
    assert rc["E14"].value == pytest.approx(amt) and rc["E21"].value == pytest.approx(amt)
    assert (rc["C24"].value, rc["C25"].value, rc["D24"].value, rc["D25"].value) == ("SGST", "CGST", 0.09, 0.09)
    assert rc["E24"].value == pytest.approx(amt * .09) and rc["E25"].value == pytest.approx(amt * .09)
    assert rc["E28"].value == pytest.approx(amt * .18) and rc["E30"].value == pytest.approx(amt * 1.18)
    total = round(amt * 1.18)
    assert rc["E32"].value == pytest.approx(total) and rc["E31"].value == pytest.approx(total - amt * 1.18)
    assert rc["A33"].value == "Amount in Words :  " + billing_cover.words(total)
    # our own company at the bottom, clearly
    assert [rc.cell(r, 2).value for r in range(36, 41)] == ["24AADCV2799M1ZP", "GUJARAT", "24", "AADCV2799M", "MEPF"]
    assert rc["A41"].value == "For Vedant Engitech And Infrastructure Pvt. Ltd."


def test_nothing_is_invented_for_details_that_are_not_filled_in(env):
    client, _, tmp = env
    cs(client)
    ch(client, "614", "2026-07-04", [("1", 2)])
    ra = new_ra(client)
    p = invoice_of(client, ra, tmp)
    rc = _recalc(p, tmp / "rc")["RA-01 INVOICE"]
    for ref in ("B2", "B3", "B4", "B5", "B6", "B8", "B9", "E3", "E4", "E5", "D14", "B36", "B37", "B38", "B39", "B40", "D24", "D25", "E24", "E25"):
        assert rc[ref].value in (None, ""), ref
    assert rc["A41"].value in (None, "") and rc["E28"].value in (0, None)       # no GST % set: GST is not guessed
    assert rc["E21"].value > 0 and rc["B14"].value == "Supply of Fire Work"      # the bill itself is still there


def test_locked_cover_is_values_only_and_the_gst_rates_freeze_with_the_bill(env):
    client, _, tmp = env
    cs(client)
    full(client)
    ch(client, "614", "2026-07-04", [("1", 2), ("2", 1)])
    ra = new_ra(client)
    client.post(f"{B}/ra/{ra}/lock")
    meta(client, sgst=2.5, cgst=2.5)                         # a later rate change must not move the submitted invoice
    p = invoice_of(client, ra, tmp)
    ws = openpyxl.load_workbook(p)["RA-01 INVOICE"]
    assert not any(isinstance(c.value, str) and c.value.startswith("=") for row in ws.iter_rows() for c in row)
    assert (ws["D24"].value, ws["D25"].value) == (0.09, 0.09) and "DRAFT" not in str(ws["A1"].value or "")
    amt = bill(client, ra)["this_total"]
    assert ws["E21"].value == pytest.approx(amt) and ws["E28"].value == pytest.approx(amt * .18)
    assert ws["E32"].value == round(amt * 1.18)
    # a reopened RA goes back to the live rates
    client.post(f"{B}/ra/{ra}/reopen")
    ws2 = openpyxl.load_workbook(invoice_of(client, ra, tmp, "re.xlsx"))["RA-01 INVOICE"]
    assert (ws2["D24"].value, ws2["D25"].value) == (0.025, 0.025)


def test_together_ra_cover_bills_supply_and_installation_in_one_line(env):
    client, _, tmp = env
    cs(client, 75, 25, svc="HVAC", mode="together")
    full(client, sac={"HVAC": "995461"}, inv_desc={"HVAC": "Supply and Installation of HVAC Work (custom text)"})
    ch(client, "9", "2026-07-04", [("1", 1)], svc="HVAC")
    client.post(f"/api/siteprogress/{SLUG}/progress/item", json={"service": "HVAC", "item_code": "1", "frac": 0.5})
    ra = new_ra(client, "together", "HVAC")
    p = invoice_of(client, ra, tmp)
    rc = _recalc(p, tmp / "rc")["RA-01 INVOICE"]
    assert rc["E14"].value == pytest.approx(3750000 + 1250000)
    assert rc["B14"].value == "Supply and Installation of HVAC Work (custom text)"
    assert "Supply and Installation Amount" in rc["A22"].value and rc["E7"].value == "HVAC"


def test_installation_ra_cover_says_installation(env):
    client, _, tmp = env
    progress(client, "2", 0.5)
    client.post(f"{B}/setup", json={"service": "Fire", "mode": "separate", "supply_pct": 70, "install_pct": 30})
    ra = new_ra(client, "install")
    rc = openpyxl.load_workbook(invoice_of(client, ra, tmp))["RA-01 INVOICE"]
    assert rc["B14"].value == "Installation of Fire Work" and "Installation Amount" in rc["A22"].value


# ---------------------------------------------------------------- Download endpoint
def _two_services(client):
    """Fire on challans (RA 1 draft + RA 1 install draft), HVAC on the stock register."""
    cs(client)
    client.post(f"{B}/setup", json={"service": "HVAC", "mode": "together", "supply_pct": 75, "install_pct": 25,
                                    "supply_source": "stock"})
    full(client, sac={"Fire": "995461", "HVAC": "995461"})
    ch(client, "614", "2026-07-04", [("1", 2), ("2", 1)])
    r1 = new_ra(client)
    r2 = new_ra(client, "install")
    r3 = new_ra(client, "together", "HVAC")
    return r1, r2, r3


def dl(client, **q):
    r = client.get(f"{B}/download", params=q)
    assert r.status_code == 200, r.text
    return r


def test_download_one_ra_is_a_plain_workbook_many_is_a_zip(env):
    client, _, tmp = env
    r1, r2, r3 = _two_services(client)
    r = dl(client, ras=r1)
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats") and ".xlsx" in r.headers["content-disposition"]
    assert openpyxl.load_workbook(io.BytesIO(r.content)).sheetnames[0] == "RA-01 INVOICE"
    r = dl(client, ras=f"{r1},{r2},{r3}")
    assert r.headers["content-type"] == "application/zip" and ".zip" in r.headers["content-disposition"]
    z = zipfile.ZipFile(io.BytesIO(r.content))
    names = z.namelist()
    assert len(names) == 3 and all(n.endswith(".xlsx") and ("DRAFT" in n) for n in names)
    assert any("Fire_Supply_RA-01" in n for n in names) and any("Fire_Installation_RA-01" in n for n in names)
    assert any("HVAC_Supply+Installation_RA-01" in n for n in names)
    for n in names:
        openpyxl.load_workbook(io.BytesIO(z.read(n)))                       # every file opens


def test_sheet_ticks_keep_only_the_chosen_sheets_and_no_formula_is_left_pointing_at_a_dropped_sheet(env):
    client, _, tmp = env
    r1, _, _ = _two_services(client)
    r = dl(client, ras=r1, sheets="boq")
    p = tmp / "boq.xlsx"
    p.write_bytes(r.content)
    wb = openpyxl.load_workbook(p)
    assert wb.sheetnames == ["Annexure B-BOQ"]
    bad = [c.coordinate for row in wb.active.iter_rows() for c in row
           if isinstance(c.value, str) and c.value.startswith("=") and "!" in c.value]
    assert bad == []
    rc = _recalc(p, tmp / "rc")["Annexure B-BOQ"]
    rp = _row_of(rc, "Diesel engine pump set")
    assert rc.cell(rp, 9).value == 2 and rc.cell(rp, 12).value == pytest.approx(140000)       # chain value kept as a number
    # invoice + summary only: the cover still follows the summary, abstract/challan are gone
    r = dl(client, ras=r1, sheets="invoice,summary")
    p2 = tmp / "is.xlsx"
    p2.write_bytes(r.content)
    assert openpyxl.load_workbook(p2).sheetnames == ["RA-01 INVOICE", "Annexure B-Summary BOQ"]
    rc2 = _recalc(p2, tmp / "rc2")
    g = _find(rc2["Annexure B-Summary BOQ"], "Grand Total")
    assert rc2["RA-01 INVOICE"]["E14"].value == pytest.approx(rc2["Annexure B-Summary BOQ"].cell(g, 5).value)
    # the summary alone has no BOQ sheet to point at: values, not broken references
    p3 = tmp / "s.xlsx"
    p3.write_bytes(dl(client, ras=r1, sheets="summary").content)
    rc3 = _recalc(p3, tmp / "rc3")["Annexure B-Summary BOQ"]
    assert rc3.cell(g, 5).value == pytest.approx(rc2["Annexure B-Summary BOQ"].cell(g, 5).value)
    assert dl(client, ras=r1, sheets="challan").status_code == 200


def test_a_stock_register_service_gets_no_challan_sheets_and_a_challan_service_does(env):
    client, _, _ = env
    r1, _, r3 = _two_services(client)
    assert openpyxl.load_workbook(io.BytesIO(dl(client, ras=r1).content)).sheetnames == [
        "RA-01 INVOICE", "Annexure B-Summary BOQ", "Annexure B-BOQ", "CHALLAN ABTRACT", "CHALLAN"]
    assert openpyxl.load_workbook(io.BytesIO(dl(client, ras=r3).content)).sheetnames == [
        "RA-01 INVOICE", "Annexure B-Summary BOQ", "Annexure B-BOQ"]


def test_one_workbook_mode_names_every_sheet_by_service_and_ra(env):
    client, _, tmp = env
    r1, r2, r3 = _two_services(client)
    r = dl(client, ras=f"{r1},{r2},{r3}", mode="one")
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    p = tmp / "one.xlsx"
    p.write_bytes(r.content)
    wb = openpyxl.load_workbook(p)
    names = wb.sheetnames
    assert len(names) == len(set(n.lower() for n in names)) and all(len(n) <= 31 for n in names)
    # Fire has a Supply RA 1 and an Installation RA 1: the tags keep them apart; HVAC is one RA, no tag
    assert "Fire RA01S INVOICE" in names and "Fire RA01I INVOICE" in names and "HVAC RA01 INVOICE" in names
    assert "Fire RA01S CH ABSTRACT" in names and "Fire RA01S CHALLAN" in names
    assert not any(n.startswith("HVAC") and ("CHALLAN" in n or "CH ABS" in n) for n in names)      # stock service: none
    assert not any("Fire RA01I" in n and "CHALLAN" in n for n in names)                              # install RA: none
    rc = _recalc(p, tmp / "rc")
    # the formulas of each RA still reach their own sheets (nothing broken, nothing crossed)
    s = rc["Fire RA01S Summary"]
    g = _find(s, "Grand Total")
    assert rc["Fire RA01S INVOICE"]["E14"].value == pytest.approx(s.cell(g, 5).value) and s.cell(g, 5).value > 0
    b = rc["Fire RA01S BOQ"]
    assert b.cell(_row_of(b, "Diesel engine pump set"), 12).value == pytest.approx(140000)
    for ws in rc.worksheets:
        assert not any(isinstance(c.value, str) and c.value.startswith(("#", "Err:")) for row in ws.iter_rows() for c in row), ws.title


def test_combined_export_has_each_ras_invoice_and_challan_sheets_named_by_service(env):
    client, _, tmp = env
    r1, r2, r3 = _two_services(client)
    r = client.get(f"{B}/export")
    assert r.status_code == 200
    p = tmp / "all.xlsx"
    p.write_bytes(r.content)
    wb = openpyxl.load_workbook(p)
    names = wb.sheetnames
    assert "Invoice (format pending)" not in names
    for inv, tab in (("Fire Supply RA-01 INVOICE", "Fire Supply RA-01"), ("Fire Install RA-01 INVOICE", "Fire Install RA-01"),
                     ("HVAC S+I RA-01 INVOICE", "HVAC Supply RA-01")):
        assert names.index(inv) == names.index(tab) - 1, inv
    assert "Fire CHALLAN ABSTRACT" in names and "Fire CHALLAN" in names
    assert not any(n.startswith("HVAC") and "CHALLAN" in n for n in names)                  # HVAC is on the stock register
    assert not any(n.startswith("~") for n in names)
    rc = _recalc(p, tmp / "rc")
    fa, fc = rc["Fire CHALLAN ABSTRACT"], rc["Fire CHALLAN"]
    rv = _row_of(fa, "Butterfly valve 150mm")
    assert fa.cell(8, 5).value == "RA 01" and fa.cell(rv, 5).value == 1 and fc.cell(rv, 6).value == 1
    amt = bill(client, r1)["this_total"]
    assert rc["Fire Supply RA-01 INVOICE"]["E14"].value == pytest.approx(amt)
    # the cover of the draft follows the annexure tab of the combined workbook
    assert "'Fire Supply RA-01'!E" in openpyxl.load_workbook(p)["Fire Supply RA-01 INVOICE"]["E14"].value
    for ws in rc.worksheets:
        assert not any(isinstance(c.value, str) and c.value.startswith(("#", "Err:")) for row in ws.iter_rows() for c in row), ws.title


def test_combined_export_challan_sheets_are_cumulative_over_ra_1_and_2(env):
    client, _, tmp = env
    cs(client)
    ch(client, "614", "2026-07-04", [("1", 2), ("2", 1)])
    ra1 = new_ra(client)
    client.post(f"{B}/ra/{ra1}/lock")
    ch(client, "616", "2026-08-13", [("2", 2), ("3", 40)])
    new_ra(client)
    p = tmp / "all2.xlsx"
    p.write_bytes(client.get(f"{B}/export").content)
    rc = _recalc(p, tmp / "rc")
    fa = rc["Fire CHALLAN ABSTRACT"]
    rv = _row_of(fa, "Butterfly valve 150mm")
    assert (fa.cell(8, 5).value, fa.cell(8, 6).value) == ("RA 01", "RA 02")
    assert (fa.cell(rv, 5).value, fa.cell(rv, 6).value, fa.cell(rv, 7).value) == (1, 2, 3)


def test_download_validation(env):
    client, _, _ = env
    r1, _, _ = _two_services(client)
    assert client.get(f"{B}/download").status_code == 400                                   # nothing ticked
    assert client.get(f"{B}/download", params={"ras": "ra999"}).status_code == 404
    assert client.get(f"{B}/download", params={"ras": r1, "mode": "zigzag"}).status_code == 400
    assert client.get(f"{B}/download", params={"ras": r1, "sheets": "nonsense"}).status_code == 400
    r = client.get(f"{B}/download", params={"combined": 1})                                  # the old All services Excel alone
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    r = client.get(f"{B}/download", params={"ras": r1, "combined": 1})
    assert r.headers["content-type"] == "application/zip" and len(zipfile.ZipFile(io.BytesIO(r.content)).namelist()) == 2


def test_service_names_with_odd_characters_still_make_legal_sheet_names(env):
    ra = {"service": "Fire & Gas: 'Main' [Block/A]?*", "no": 3, "series": "supply"}
    n = billing._sheet_names(ra, plain=False)
    for v in n.values():
        assert len(v) <= 31 and not any(ch_ in v for ch_ in ":\\/?*[]'"), v
    assert len(set(v.lower() for v in n.values())) == 5


# ------------------------------------------------- look of the CHALLAN / CHALLAN ABTRACT tabs
def _two_ra_challan_book(client, tmp):
    """The user's real case: challans over two RAs, RA 1 locked, RA 2 open; RA 2's workbook."""
    cs(client)
    full(client)
    ch(client, "614", "2026-07-04", [("1", 2), ("2", 1)])
    ra1 = new_ra(client)
    client.post(f"{B}/ra/{ra1}/lock")
    ch(client, "616", "2026-08-13", [("2", 2), ("3", 40)])
    ra2 = new_ra(client)
    return openpyxl.load_workbook(invoice_of(client, ra2, tmp, "two.xlsx"))


def _boxed(c):
    b = c.border
    return all(getattr(b, k).style == "thin" for k in ("left", "right", "top", "bottom"))


def test_challan_tab_header_block_is_boxed_bold_and_centred_like_the_users_own_sheet(env):
    client, _, tmp = env
    wb = _two_ra_challan_book(client, tmp)
    w = wb["CHALLAN"]
    # F, G = the two challans; H, I, J = Total / Previous / This bill
    for col in "FGHIJ":
        for r in range(3, 9):
            c = w[f"{col}{r}"]
            assert _boxed(c), f"{col}{r} has no box"
            assert c.font.b and c.alignment.horizontal == "center", f"{col}{r} not bold/centred"
    assert (w["F3"].value, w["F5"].value, w["F6"].value, w["F7"].value, w["F8"].value) == ("Challan Date", "Challan No", 614, "RA-01", "Qty")
    assert (w["G6"].value, w["G7"].value) == (616, "RA-02")
    assert (w["H3"].value, w["I3"].value, w["J3"].value) == ("Total QTY", "Previous Bill QTY", "This Bill QTY")
    assert w["J7"].value == "RA-02"
    # his row heights and centred labels
    assert [w.row_dimensions[r].height for r in (3, 4, 5, 6)] == [32, 22, 22, 26]
    for r in (2, 3, 4, 6):
        assert w[f"A{r}"].font.b and w[f"A{r}"].alignment.horizontal == "center"
    # the work-details block has one frame (A1..E7), like the original
    assert w["A1"].border.left.style == "thin" and w["A1"].border.top.style == "thin"
    side = lambda c, k: getattr(getattr(c.border, k), "style", None) if getattr(c.border, k) else None
    assert side(w["E4"], "right") == "thin" and side(w["C7"], "bottom") == "thin" and side(w["C4"], "bottom") is None
    # header row 8 matches his sheet: no peach fill; quantities are centred, the three bill columns bold
    assert w["B8"].fill.fill_type is None
    r = _row_of(w, "Butterfly valve 150mm")
    assert w[f"F{r}"].alignment.horizontal == "center" and w[f"H{r}"].font.b and not w[f"F{r}"].font.b


def test_abstract_tab_header_says_certified_qty_once_over_the_ra_columns_and_total_once(env):
    client, _, tmp = env
    wb = _two_ra_challan_book(client, tmp)
    w = wb["CHALLAN ABTRACT"]
    merged = {str(m) for m in w.merged_cells.ranges}
    assert {"E7:F7", "G7:G8"} <= merged, merged                       # CERTIFIED QTY over RA 01..02, TOTAL QTY over both header rows
    assert (w["E7"].value, w["G7"].value, w["E8"].value, w["F8"].value) == ("CERTIFIED QTY", "TOTAL QTY", "RA 01", "RA 02")
    assert w["G8"].value is None                                       # not said twice
    for ref in ("E7", "G7", "E8", "F8", "B8"):                 # (the other cells of a merged block share its edge)
        assert _boxed(w[ref]), ref
    assert w["B8"].fill.fill_type is None
    r = _row_of(w, "Butterfly valve 150mm")
    assert w[f"E{r}"].alignment.horizontal == "center" and w[f"G{r}"].font.b


def test_abstract_with_one_ra_has_no_one_cell_merge(env):
    client, _, tmp = env
    cs(client)
    ch(client, "614", "2026-07-04", [("1", 2)])
    ra = new_ra(client)
    w = openpyxl.load_workbook(invoice_of(client, ra, tmp))["CHALLAN ABTRACT"]
    assert {str(m) for m in w.merged_cells.ranges} == {"A1:B1", "F7:F8"}


def test_boq_tab_header_is_left_as_it_was_and_the_draft_banner_has_a_font(env):
    client, _, tmp = env
    wb = _two_ra_challan_book(client, tmp)
    assert wb["Annexure B-BOQ"]["B8"].fill.fgColor.rgb.endswith("FCE4D6")          # B-BOQ keeps its own look
    for n in ("CHALLAN", "CHALLAN ABTRACT", "Annexure B-BOQ"):
        assert wb[n]["B5"].font.name == "Calibri" and wb[n]["B5"].font.b


def test_a_challan_number_typed_as_an_ra_tag_shows_twice_exactly_as_typed(env):
    """The user's file: challan number saved as 'RA-01' -> row 6 (the number) and row 7 (the RA tag) both read RA-01.
    The sheet prints what was entered; it does not rewrite a challan number."""
    client, _, tmp = env
    cs(client)
    ch(client, "RA-01", "2026-10-09", [("1", 2)])
    ra = new_ra(client)
    w = openpyxl.load_workbook(invoice_of(client, ra, tmp))["CHALLAN"]
    assert (w["F5"].value, w["F6"].value, w["F7"].value) == ("Challan No", "RA-01", "RA-01")
