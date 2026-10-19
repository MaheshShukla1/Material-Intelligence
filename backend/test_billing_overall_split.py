"""Overall annexure cut into Supply and Installation.

Reported scenario: one Work Order whose "In draft" holds a Supply draft AND an Installation draft (and an invoiced Supply RA).
The annexure only said "In draft ₹ x"; nobody could tell how much of it was Supply and how much Installation.
The server now sends the same numbers cut by stage (additive keys: rows[].stages, units[].stages, totals.stages,
ras[].by_stage) and the Excel gets a second sheet. Nothing is invented: a stage with no % set, or a group whose Work Order
value is typed in the BOQ sheet, has no stage share (None / a dash)."""
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import openpyxl
import pytest

from backend.test_billing import env, progress, SLUG  # noqa: F401  (env is a fixture)
from backend.test_billing_challan import cs, ch, new_ra, bill, B


def overall(client, base="Fire"):
    r = client.get(f"{B}/overall/{base}")
    assert r.status_code == 200, r.text
    return r.json()


def reported_scenario(client):
    """Supply RA 1 invoiced, Supply RA 2 draft, Installation RA 1 draft: all three in the same Work Order."""
    cs(client)                                                           # Fire: Supply 70% / Installation 30%, supply from challans
    assert ch(client, "614", "2026-07-04", [("1", 1)]).status_code == 200
    s1 = new_ra(client)
    assert client.post(f"{B}/ra/{s1}/lock").status_code == 200
    assert ch(client, "615", "2026-07-05", [("1", 1), ("2", 2)]).status_code == 200
    s2 = new_ra(client)
    progress(client, "1", 0.5)                                           # installation comes from Site Progress
    progress(client, "2", 0.5)
    i1 = new_ra(client, series="install")
    return s1, s2, i1


def test_the_draft_is_cut_into_supply_and_installation_and_adds_back_to_the_whole(env):
    client, _, _ = env
    s1, s2, i1 = reported_scenario(client)
    o = overall(client)
    u = o["units"][0]
    sup, ins = u["stages"]["supply"], u["stages"]["install"]
    # the numbers are exactly what each RA says on its own screen
    assert sup["invoiced"] == pytest.approx(bill(client, s1)["this_total"])
    assert sup["draft"] == pytest.approx(bill(client, s2)["this_total"]) and sup["draft"] > 0
    assert ins["draft"] == pytest.approx(bill(client, i1)["this_total"]) and ins["draft"] > 0
    assert ins["invoiced"] == 0
    # nothing lost, nothing counted twice: the two stages add up to what the annexure already showed
    assert sup["invoiced"] + ins["invoiced"] == pytest.approx(u["invoiced"])
    assert sup["draft"] + ins["draft"] == pytest.approx(u["draft"])
    assert o["totals"]["stages"]["supply"]["draft"] == pytest.approx(sup["draft"])
    assert o["totals"]["stages"]["install"]["draft"] == pytest.approx(ins["draft"])
    # the same holds for every group row
    for x in u["rows"]:
        assert sum(x["stages"][k]["invoiced"] for k in ("supply", "install")) == pytest.approx(x["invoiced"])
        assert sum(x["stages"][k]["draft"] for k in ("supply", "install")) == pytest.approx(x["draft"])
    assert sum(x["draft"] for x in u["rows"]) == pytest.approx(u["draft"])
    # the old keys are untouched
    assert u["draft"] == pytest.approx(bill(client, s2)["this_total"] + bill(client, i1)["this_total"])
    assert o["totals"]["todate"] == pytest.approx(u["invoiced"] + u["draft"])


def test_each_stage_has_its_own_share_of_the_work_order(env):
    client, _, _ = env
    reported_scenario(client)
    u = overall(client)["units"][0]
    assert u["stages"]["supply"]["wo"] == pytest.approx(u["wo_value"] * .70)
    assert u["stages"]["install"]["wo"] == pytest.approx(u["wo_value"] * .30)
    assert u["stages"]["supply"]["pct"] == 70 and u["stages"]["install"]["pct"] == 30
    for x in u["rows"]:
        assert x["stages"]["supply"]["wo"] == pytest.approx(x["wo"] * .70)
        assert x["stages"]["install"]["wo"] == pytest.approx(x["wo"] * .30)


def test_every_open_draft_is_listed_with_the_stage_it_carries(env):
    client, _, _ = env
    s1, s2, i1 = reported_scenario(client)
    ras = {r["id"]: r for r in overall(client)["units"][0]["ras"]}
    assert set(ras) == {s1, s2, i1}
    assert set(ras[s2]["by_stage"]) == {"supply"} and ras[s2]["by_stage"]["supply"] == pytest.approx(ras[s2]["this_total"])
    assert set(ras[i1]["by_stage"]) == {"install"} and ras[i1]["by_stage"]["install"] == pytest.approx(ras[i1]["this_total"])
    assert ras[s1]["status"] == "locked" and set(ras[s1]["by_stage"]) == {"supply"}
    assert [r["label"] for r in ras.values() if r["status"] == "draft"].count("Installation RA 1") == 1


def test_a_together_bill_is_read_out_in_its_two_parts_and_counted_once(env):
    client, _, _ = env
    cs(client, mode="together")
    assert ch(client, "614", "2026-07-04", [("1", 1), ("2", 2)]).status_code == 200
    progress(client, "1", 0.5)
    progress(client, "2", 0.25)
    ra = new_ra(client, series="together")
    b = bill(client, ra)
    u = overall(client)["units"][0]
    r = next(x for x in u["ras"] if x["id"] == ra)
    assert set(r["by_stage"]) == {"supply", "install"}
    assert r["by_stage"]["supply"] == pytest.approx(b["totals"]["supply"]["this"])
    assert r["by_stage"]["install"] == pytest.approx(b["totals"]["install"]["this"])
    assert r["by_stage"]["supply"] + r["by_stage"]["install"] == pytest.approx(r["this_total"])
    assert r["this_total"] == pytest.approx(u["draft"])
    assert u["stages"]["supply"]["draft"] + u["stages"]["install"]["draft"] == pytest.approx(u["draft"])
    assert u["stages"]["supply"]["draft"] > 0 and u["stages"]["install"]["draft"] > 0


def test_a_stage_with_no_percentage_is_not_guessed(env):
    """HVAC-style: only Supply is set up. Installation has no share and no number, the Supply side is still real."""
    client, _, _ = env
    r = client.post(f"{B}/setup", json={"service": "Fire", "mode": "separate", "supply_pct": 70, "supply_source": "challan"})
    assert r.status_code == 200, r.text
    assert ch(client, "614", "2026-07-04", [("1", 1)]).status_code == 200
    new_ra(client)
    u = overall(client)["units"][0]
    assert u["stages"]["install"]["set"] is False and u["stages"]["install"]["pct"] is None
    assert u["stages"]["install"]["wo"] is None and all(x["stages"]["install"]["wo"] is None for x in u["rows"])
    assert u["stages"]["supply"]["set"] is True and u["stages"]["supply"]["wo"] == pytest.approx(u["wo_value"] * .70)
    assert u["stages"]["supply"]["draft"] == pytest.approx(u["draft"]) and u["draft"] > 0
    assert overall(client)["totals"]["stages"]["install"]["wo"] is None


def test_a_group_value_typed_in_the_sheet_has_no_stage_share_but_keeps_its_real_amounts(env):
    client, d, _ = env
    reported_scenario(client)
    calc = overall(client)["units"][0]["wo_value"]
    client.post(f"{B}/groups/Fire", json={"groups": [{"sr": "1", "name": "Everything", "sheet": "Fire", "from": 2, "to": 20, "wo": 123456}]})
    u = overall(client)["units"][0]
    (row,) = u["rows"]
    assert row["wo"] == 123456 and row["stages"]["supply"]["wo"] is None and row["stages"]["install"]["wo"] is None
    assert row["stages"]["supply"]["draft"] > 0 and row["stages"]["install"]["draft"] > 0          # billed amounts are real
    assert u["stages"]["supply"]["wo"] is None
    # typed value that equals rate x qty is not a different number, so the share is known again
    client.post(f"{B}/groups/Fire", json={"groups": [{"sr": "1", "name": "Everything", "sheet": "Fire", "from": 2, "to": 20, "wo": calc}]})
    (row2,) = overall(client)["units"][0]["rows"]
    assert row2["stages"]["supply"]["wo"] == pytest.approx(calc * .70)


def sheet2(client, drafts=1):
    r = client.get(f"{B}/overall/Fire/xlsx?drafts={drafts}")
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    return wb


def test_excel_keeps_sheet_1_and_adds_a_supply_vs_installation_sheet(env):
    client, _, _ = env
    reported_scenario(client)
    for drafts, c1, c2 in ((1, 7, 13), (0, 6, 11)):
        wb = sheet2(client, drafts)
        assert wb.sheetnames == ["Overall annexure", "Supply vs Installation"]
        assert wb["Overall annexure"].max_column == c1                               # sheet 1: same columns as before
        ws = wb["Supply vs Installation"]
        assert ws.max_column == c2
        u = overall(client)["units"][0]
        top = [c.value for c in ws[8]]
        assert top[3] == "Supply · 70%" and top[3 + (c2 - 3) // 2] == "Installation · 30%"
        head = [c.value for c in ws[9]]
        assert head[:3] == ["Sr.No.", "Description", "Work Order amount"]
        assert head[3:3 + (c2 - 3) // 2] == (["Value", "Invoiced", "In draft", "Billed to date", "Done %"] if drafts else ["Value", "Invoiced", "Billed to date", "Done %"])
        first = [c.value for c in ws[10]]
        x = u["rows"][0]
        assert first[1] == x["name"] and first[2] == pytest.approx(x["wo"])
        assert first[3] == pytest.approx(x["stages"]["supply"]["wo"]) and first[4] == pytest.approx(x["stages"]["supply"]["invoiced"])
        w = (c2 - 3) // 2
        assert first[3 + w] == pytest.approx(x["stages"]["install"]["wo"])
        if drafts:
            assert first[5] == pytest.approx(x["stages"]["supply"]["draft"]) and first[3 + w + 2] == pytest.approx(x["stages"]["install"]["draft"])
            assert first[6] == "=E10+F10" and first[3 + w + 3] == f"={ws.cell(10, 5 + w).column_letter}10+{ws.cell(10, 6 + w).column_letter}10"
        else:
            assert first[5] == "=E10"
        texts = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
        assert any("Subtotal" in t for t in texts) and any("Fire, all Work Orders" in t for t in texts)


def test_excel_shows_a_dash_not_a_number_for_a_stage_that_is_not_set(env):
    client, _, _ = env
    client.post(f"{B}/setup", json={"service": "Fire", "mode": "separate", "supply_pct": 70, "supply_source": "challan"})
    ch(client, "614", "2026-07-04", [("1", 1)])
    new_ra(client)
    ws = sheet2(client)["Supply vs Installation"]
    assert [c.value for c in ws[8]][8] == "Installation"                              # no % to print
    for row in ws.iter_rows(min_row=10, max_row=ws.max_row):
        if row[1].value and row[1].value != "":
            assert all(c.value == "—" for c in row[8:13]), [c.value for c in row]
    sup_first = [c.value for c in ws[10]]
    assert isinstance(sup_first[3], (int, float)) and sup_first[3] > 0
