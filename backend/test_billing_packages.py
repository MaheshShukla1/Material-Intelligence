"""Work Orders (packages): one service billed under several Work Orders, each with its OWN RA numbers.
Real scenario (Thoth Electrical): "Excl. point wiring" and "Only point wiring" both have an RA 1; the service also
has one overall annexure that adds the Work Orders up. Fire is split here the way Electrical is in the real project:
group A "Fire Pump & Accessories" (rows 2-6) = Work Order 'Pumps', group B "Hydrant System" (rows 7-9) = the remainder."""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import io
import openpyxl
import pytest

from backend import billing
from backend.test_billing import env, progress, SLUG  # noqa: F401  (env is a fixture)
from backend.test_billing_challan import cs, ch, new_ra, bill, B

PUMPS, HYD = "Fire › Pumps", "Fire › Hydrants"


def groups(client):
    return client.get(f"{B}/packages/Fire").json()["groups"]


def split(client, expect=200, **kw):
    g = groups(client)
    items = [{"name": "Pumps", "groups": [{"sheet": g[0]["sheet"], "from": g[0]["from"], "to": g[0]["to"]}]},
             {"name": "Hydrants", "rest": True}]
    r = client.post(f"{B}/packages/Fire", json={"items": kw.get("items", items)})
    assert r.status_code == expect, r.text
    return r.json()


def both_setup(client, sup=70, ins=30):
    for u in (PUMPS, HYD):
        r = client.post(f"{B}/setup", json={"service": u, "mode": "separate", "supply_pct": sup, "install_pct": ins,
                                            "supply_source": "challan"})
        assert r.status_code == 200, r.text


def unit(client, name):
    return next(s for s in client.get(B).json()["services"] if s["service"] == name)


def test_a_service_is_one_unit_until_it_is_split_then_one_unit_per_work_order(env):
    client, _, _ = env
    assert [s["service"] for s in client.get(B).json()["services"]] == ["Fire", "HVAC"]
    out = split(client)
    assert out["units"] == [PUMPS, HYD]
    svcs = client.get(B).json()["services"]
    assert [s["service"] for s in svcs] == [PUMPS, HYD, "HVAC"]
    assert (unit(client, PUMPS)["base"], unit(client, PUMPS)["package"]) == ("Fire", "Pumps")
    # contract value is split by the lines, and the two add up to the whole service
    pumps_v, hyd_v = unit(client, PUMPS)["contract_value"], unit(client, HYD)["contract_value"]
    assert pumps_v == 2 * 100000 + 4 * 5000 + 100 * 800 and hyd_v == 10 * 3000
    assert [b["units"] for b in client.get(B).json()["bases"] if b["base"] == "Fire"] == [[PUMPS, HYD]]


def test_each_work_order_has_its_own_ra_1_and_its_own_lines_only(env):
    client, _, _ = env
    split(client)
    both_setup(client)
    ch(client, "614", "2026-07-04", [("1", 2)], svc=PUMPS)
    ch(client, "615", "2026-07-05", [("5", 3)], svc=HYD)
    a, b = new_ra(client, svc=PUMPS), new_ra(client, svc=HYD)
    ba, bb = bill(client, a), bill(client, b)
    assert ba["ra"]["no"] == 1 and bb["ra"]["no"] == 1                      # RA 1 for each Work Order
    assert {x["item_code"] for x in ba["lines"]} == {"1"} and {x["item_code"] for x in bb["lines"]} == {"5"}
    assert ba["ra"]["label"] == "Fire › Pumps · Supply RA 1"
    assert ba["this_total"] == pytest.approx(100000 * 2 * .70) and bb["this_total"] == pytest.approx(3000 * 3 * .70)
    # a line of the other Work Order cannot be billed in this one
    assert client.post(f"{B}/ra/{a}/line", json={"item_code": "5", "stage": "supply", "qty": 1}).status_code == 400
    # locking one does not move the other's numbering
    assert client.post(f"{B}/ra/{a}/lock").status_code == 200
    again = new_ra(client, "supply", svc=PUMPS)
    assert bill(client, again)["ra"]["no"] == 2
    assert bill(client, b)["ra"]["no"] == 1 and bill(client, b)["ra"]["status"] == "draft"


def test_challans_belong_to_their_work_order_and_an_lines_outside_it_are_refused(env):
    client, _, _ = env
    split(client)
    both_setup(client)
    assert ch(client, "614", "2026-07-04", [("5", 2)], svc=PUMPS).status_code == 400       # hydrant valve in the Pumps WO
    assert ch(client, "614", "2026-07-04", [("1", 2)], svc=PUMPS).status_code == 200
    assert ch(client, "614", "2026-07-04", [("5", 2)], svc=HYD).status_code == 200         # same challan no. in the other WO is fine
    assert len(client.get(f"{B}/challans/{PUMPS}").json()["challans"]) == 1


def test_installation_comes_from_site_progress_for_the_right_work_order(env):
    client, _, _ = env
    split(client)
    both_setup(client)
    progress(client, "1", 0.5)                                            # a pump line
    progress(client, "5", 1.0)                                            # a hydrant line
    made = client.post(f"{B}/auto").json()["created"]
    assert {m["service"] for m in made} == {PUMPS, HYD}
    labels = {m["service"]: m["label"] for m in made}
    assert labels[PUMPS].startswith("Fire › Pumps · Installation RA 1") and labels[HYD].startswith("Fire › Hydrants · Installation RA 1")
    amt = {m["service"]: m["install"] for m in made}
    assert amt[PUMPS] == pytest.approx(100000 * 1 * .30) and amt[HYD] == pytest.approx(3000 * 10 * .30)


def test_first_split_moves_an_invoiced_ra_and_its_challans_into_the_work_order_it_belongs_to(env):
    client, _, _ = env
    cs(client)                                                           # plain 'Fire', challan source
    ch(client, "614", "2026-07-04", [("1", 2), ("2", 1)])
    ra = new_ra(client)
    assert client.post(f"{B}/ra/{ra}/lock").status_code == 200
    ch(client, "700", "2026-07-20", [("2", 1), ("5", 4)])                # a challan with lines of BOTH Work Orders, billed later
    out = split(client)
    assert out["removed_drafts"] == 0
    ras = client.get(B).json()["ras"]
    assert [(r["service"], r["no"], r["status"]) for r in ras] == [(PUMPS, 1, "locked")]
    snap = bill(client, ra)
    assert snap["ra"]["service"] == PUMPS and snap["ra"]["status"] == "locked"
    assert {x["item_code"] for x in snap["lines"]} == {"1", "2"}
    assert snap["contract_value"] == unit(client, PUMPS)["contract_value"]
    assert [g["name"] for g in snap["groups"]] == ["Fire Pump & Accessories"]          # the other WO's group is not on this bill
    # the locked bill is exactly what it was (70% of 2 pumps + 1 valve)
    assert snap["this_total"] == pytest.approx(100000 * 2 * .70 + 5000 * 1 * .70)
    # setup carried over to both Work Orders; challan 700 is split by lines
    assert unit(client, PUMPS)["supply_pct"] == 70 and unit(client, HYD)["supply_pct"] == 70
    pch = {c["no"]: c for c in client.get(f"{B}/challans/{PUMPS}").json()["challans"]}
    hch = {c["no"]: c for c in client.get(f"{B}/challans/{HYD}").json()["challans"]}
    assert set(pch) == {"614", "700"} and set(hch) == {"700"}
    assert [l["item_code"] for l in pch["700"]["lines"]] == ["2"] and [l["item_code"] for l in hch["700"]["lines"]] == ["5"]
    # Hydrants has no RA yet: it starts at RA 1, Pumps' next one is RA 2
    assert bill(client, new_ra(client, svc=HYD))["ra"]["no"] == 1
    assert bill(client, new_ra(client, svc=PUMPS))["ra"]["no"] == 2


def test_an_invoiced_ra_with_lines_of_two_work_orders_is_never_cut_in_two(env):
    client, _, _ = env
    cs(client)
    ch(client, "614", "2026-07-04", [("1", 2), ("5", 1)])
    ra = new_ra(client)
    client.post(f"{B}/ra/{ra}/lock")
    r = client.post(f"{B}/packages/Fire", json={"items": [
        {"name": "Pumps", "groups": [{"sheet": groups(client)[0]["sheet"], "from": groups(client)[0]["from"], "to": groups(client)[0]["to"]}]},
        {"name": "Hydrants", "rest": True}]})
    assert r.status_code == 409 and "more than one Work Order" in r.json()["detail"]
    assert [s["service"] for s in client.get(B).json()["services"]] == ["Fire", "HVAC"]     # nothing changed
    assert bill(client, ra)["ra"]["status"] == "locked"


def test_an_untouched_draft_is_dropped_on_split(env):
    client, _, _ = env
    cs(client)
    ch(client, "614", "2026-07-04", [("1", 2)])
    new_ra(client)
    out = split(client)
    assert out["removed_drafts"] == 1 and client.get(B).json()["ras"] == []     # it is live: it opens again by itself


def test_a_draft_with_typed_quantities_blocks_the_split_instead_of_being_thrown_away(env):
    client, _, _ = env
    cs(client)
    d = new_ra(client)
    assert client.post(f"{B}/ra/{d}/line", json={"item_code": "1", "stage": "supply", "qty": 1, "reason": "x"}).status_code == 200
    g = groups(client)
    r = client.post(f"{B}/packages/Fire", json={"items": [
        {"name": "Pumps", "groups": [{"sheet": g[0]["sheet"], "from": g[0]["from"], "to": g[0]["to"]}]},
        {"name": "Hydrants", "rest": True}]})
    assert r.status_code == 409 and "open draft" in r.json()["detail"]
    assert [x["id"] for x in client.get(B).json()["ras"]] == [d]


def test_work_orders_can_be_merged_back_only_while_nothing_is_billed_in_them(env):
    client, _, _ = env
    split(client)
    both_setup(client)
    assert client.post(f"{B}/packages/Fire", json={"items": []}).status_code == 200
    assert [s["service"] for s in client.get(B).json()["services"]] == ["Fire", "HVAC"]
    assert unit(client, "Fire")["supply_pct"] == 70                                  # the first Work Order's term comes back


def test_definition_is_validated(env):
    client, _, _ = env
    g = groups(client)
    grp = lambda i: {"sheet": g[i]["sheet"], "from": g[i]["from"], "to": g[i]["to"]}
    bad = lambda items: client.post(f"{B}/packages/Fire", json={"items": items})
    assert bad([{"name": "A", "groups": [grp(0)]}, {"name": "B", "groups": [grp(1)]}]).status_code == 400      # no remainder
    assert bad([{"name": "A", "groups": [grp(0)]}, {"name": "a", "rest": True}]).status_code == 400             # same name
    assert bad([{"name": "", "groups": [grp(0)]}, {"name": "B", "rest": True}]).status_code == 400
    assert bad([{"name": "A/B", "groups": [grp(0)]}, {"name": "B", "rest": True}]).status_code == 400
    assert bad([{"name": "A", "groups": []}, {"name": "B", "rest": True}]).status_code == 400                  # nothing ticked
    assert bad([{"name": "A", "groups": [grp(0), grp(1)]}, {"name": "B", "rest": True}]).status_code == 400     # remainder empty
    assert bad([{"name": "A", "groups": [grp(0)]}, {"name": "B", "groups": [grp(0)]}, {"name": "C", "rest": True}]).status_code == 400
    assert bad([{"name": "A", "groups": [grp(0)]}, {"name": "B", "rest": True}]).status_code == 200


def test_a_work_order_that_bills_keeps_its_lines_and_can_be_renamed(env):
    client, _, _ = env
    split(client)
    both_setup(client)
    ch(client, "614", "2026-07-04", [("1", 2)], svc=PUMPS)
    ra = new_ra(client, svc=PUMPS)
    client.post(f"{B}/ra/{ra}/lock")
    g = groups(client)
    # moving the pump group out of Pumps is refused
    r = client.post(f"{B}/packages/Fire", json={"items": [
        {"name": "Pumps", "was": "Pumps", "groups": [{"sheet": g[1]["sheet"], "from": g[1]["from"], "to": g[1]["to"]}]},
        {"name": "Hydrants", "was": "Hydrants", "rest": True}]})
    assert r.status_code == 409
    # renaming is fine and takes the RA, challans and setup with it
    r = client.post(f"{B}/packages/Fire", json={"items": [
        {"name": "Pump sets", "was": "Pumps", "groups": [{"sheet": g[0]["sheet"], "from": g[0]["from"], "to": g[0]["to"]}]},
        {"name": "Hydrants", "was": "Hydrants", "rest": True}]})
    assert r.status_code == 200, r.text
    ras = client.get(B).json()["ras"]
    assert [(x["service"], x["no"]) for x in ras] == [("Fire › Pump sets", 1)]
    assert len(client.get(f"{B}/challans/Fire › Pump sets").json()["challans"]) == 1
    assert unit(client, "Fire › Pump sets")["supply_pct"] == 70
    # merging back is refused while a Work Order has RAs
    assert client.post(f"{B}/packages/Fire", json={"items": []}).status_code == 409


def test_overall_annexure_adds_up_both_work_orders_and_separates_invoiced_from_draft(env):
    client, _, _ = env
    split(client)
    both_setup(client)
    ch(client, "614", "2026-07-04", [("1", 2)], svc=PUMPS)
    ch(client, "615", "2026-07-05", [("5", 3)], svc=HYD)
    a, b = new_ra(client, svc=PUMPS), new_ra(client, svc=HYD)
    client.post(f"{B}/ra/{a}/lock")                                       # Pumps invoiced, Hydrants still a draft
    o = client.get(f"{B}/overall/Fire").json()
    assert o["split"] is True and [u["unit"] for u in o["units"]] == [PUMPS, HYD]
    p, h = o["units"]
    assert p["invoiced"] == pytest.approx(100000 * 2 * .70) and p["draft"] == 0
    assert h["invoiced"] == 0 and h["draft"] == pytest.approx(3000 * 3 * .70)
    assert p["latest"]["status"] == "locked" and h["latest"]["status"] == "draft"
    assert p["wo_value"] == 2 * 100000 + 4 * 5000 + 100 * 800 and h["wo_value"] == 30000
    assert o["totals"]["wo_value"] == p["wo_value"] + h["wo_value"]
    assert o["totals"]["invoiced"] == pytest.approx(p["invoiced"]) and o["totals"]["draft"] == pytest.approx(h["draft"])
    assert o["totals"]["todate"] == pytest.approx(p["invoiced"] + h["draft"])
    row = next(x for x in p["rows"] if x["name"] == "Fire Pump & Accessories")
    assert row["wo"] == p["wo_value"] and row["billed"] and row["pct"] == pytest.approx(row["todate"] / row["wo"] * 100)
    # a Work Order's rows only list its own groups
    assert [x["name"] for x in h["rows"]] == ["Hydrant System"]
    # the same total through the RA list
    ovr = client.get(B).json()
    assert ovr["billed"] == pytest.approx(p["invoiced"])


def test_overall_annexure_for_a_service_that_is_not_split_is_just_that_service(env):
    client, _, _ = env
    cs(client)
    ch(client, "614", "2026-07-04", [("1", 2)])
    client.post(f"{B}/ra/{new_ra(client)}/lock")
    o = client.get(f"{B}/overall/Fire").json()
    assert o["split"] is False and [u["unit"] for u in o["units"]] == ["Fire"]
    assert client.get(f"{B}/overall/Nope").status_code == 404


def test_overall_annexure_excel_has_every_work_order_and_a_grand_total(env):
    client, _, _ = env
    split(client)
    both_setup(client)
    ch(client, "614", "2026-07-04", [("1", 2)], svc=PUMPS)
    ch(client, "615", "2026-07-05", [("5", 3)], svc=HYD)
    a, b = new_ra(client, svc=PUMPS), new_ra(client, svc=HYD)
    client.post(f"{B}/ra/{a}/lock")
    for drafts, cols in ((1, 7), (0, 6)):
        r = client.get(f"{B}/overall/Fire/xlsx?drafts={drafts}")
        assert r.status_code == 200 and "overall_annexure" in r.headers["content-disposition"]
        ws = openpyxl.load_workbook(io.BytesIO(r.content))["Overall annexure"]
        texts = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
        assert any("Work Order: Pumps" in t for t in texts) and any("Work Order: Hydrants" in t for t in texts)
        assert any("Fire, all Work Orders" in t for t in texts)
        assert ws.max_column == cols
    assert "_invoiced" in client.get(f"{B}/overall/Fire/xlsx?drafts=0").headers["content-disposition"]


def test_wo_number_sac_and_scope_text_are_per_work_order_and_reach_the_invoice(env):
    client, _, _ = env
    split(client)
    both_setup(client)
    ch(client, "614", "2026-07-04", [("1", 2)], svc=PUMPS)
    ch(client, "615", "2026-07-05", [("5", 3)], svc=HYD)
    a, b = new_ra(client, svc=PUMPS), new_ra(client, svc=HYD)
    client.post(f"{B}/meta", json={"wo_no": "PROJECT-WO", "wo_of": {PUMPS: "WO/77"}, "sac": {PUMPS: "995461", HYD: "995462"},
                                   "work": {PUMPS: "Pump works only"}})
    def sheet(ra, title_has):
        r = client.get(f"{B}/invoice/{ra}")
        assert r.status_code == 200, r.text
        return openpyxl.load_workbook(io.BytesIO(r.content))
    wb = sheet(a, "")
    cells = {str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None}
    assert "WO/77" in cells and "Pump works only" in cells and "PROJECT-WO" not in cells
    wb2 = sheet(b, "")
    cells2 = {str(c.value) for ws in wb2 for row in ws.iter_rows() for c in row if c.value is not None}
    assert "PROJECT-WO" in cells2 and "WO/77" not in cells2                # no own number -> the project's
    missing = {r["id"]: r["inv_missing"] for r in client.get(B).json()["ras"]}
    assert "W/O number" not in missing[a] and "SAC code of Fire › Pumps" not in missing[a]


def test_downloads_and_the_all_services_excel_work_with_two_work_orders(env):
    client, _, _ = env
    split(client)
    both_setup(client)
    ch(client, "614", "2026-07-04", [("1", 2)], svc=PUMPS)
    ch(client, "615", "2026-07-05", [("5", 3)], svc=HYD)
    a, b = new_ra(client, svc=PUMPS), new_ra(client, svc=HYD)
    client.post(f"{B}/ra/{a}/lock")
    r = client.get(f"{B}/download?ras={a},{b}&mode=files")
    assert r.status_code == 200 and r.headers["content-type"] == "application/zip"
    import zipfile
    names = zipfile.ZipFile(io.BytesIO(r.content)).namelist()
    assert len(names) == 2 and len(set(names)) == 2 and all("Fire-" in n for n in names), names
    r = client.get(f"{B}/download?ras={a},{b}&mode=one")
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert len(wb.sheetnames) == len(set(n.lower() for n in wb.sheetnames))
    r = client.get(f"{B}/download?combined=1")
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert len(wb.sheetnames) == len(set(n.lower() for n in wb.sheetnames))
    reg = wb["RA Register"]
    svc = {reg.cell(i, 1).value for i in range(5, 12)}
    assert PUMPS in svc and HYD in svc
