import sys, json, os, io
from playwright.sync_api import sync_playwright
import urllib.request as u
import numpy as np, pandas as pd
"""Chromium test (desktop 1366 + phone 390), real backend and the real Thoth BOQ. Electrical's "Only point wiring" Work Order is billed from
challans. Its stock register has dated INs of ordinary material and the real wire deliveries of the site's WIRE CHALLAN sheet (28 Mar, 2 Jul,
16 Jul). Nothing is typed twice: the engineer searches the inbox, adds a delivery to RA-01 with a real challan number, confirms the wire
materials once, gives the metres to the point lines, and the supply of RA-01 is the real 4,020 + 300 + 130 points = Rs 64,28,813.
Phone: nothing scrolls sideways, controls are finger-sized.
env: BASE (default http://localhost:8765), ROOT (scratch app root), OUT (screenshots)."""
BASE = os.environ.get("BASE", "http://localhost:8765")
ROOT = os.environ["ROOT"]
OUT = os.environ.get("OUT", "/tmp/billing_stockinbox_shots")
os.makedirs(OUT, exist_ok=True)
P = "/api/billing/thoth-mall"
PW = "Electrical › Only point wiring"
errs = []
def post(path, body, method="POST"):
    r = u.Request(BASE + path, data=json.dumps(body).encode(), headers={"Content-Type": "application/json"}, method=method)
    return json.load(u.urlopen(r))
def get(path):
    return json.load(u.urlopen(BASE + path))
def noscroll(pg, label):
    assert 0 == pg.evaluate("document.documentElement.scrollWidth - document.documentElement.clientWidth"), label + ": page scrolls sideways"
def fingers(pg, sel, label, min_h=40):
    hs = pg.eval_on_selector_all(sel, "els => els.filter(e => e.offsetParent).map(e => Math.round(e.getBoundingClientRect().height))")
    assert hs, label + ": nothing matched " + sel
    for h in hs:
        assert h >= min_h, f"{label}: control only {h}px tall ({sel})"
WIRE = {"2026-03-28": {"1.5": dict(RED=300, YELLOW=600, BLUE=300, BLACK=900, GREEN=1500), "2.5": dict(RED=300, YELLOW=300, BLUE=300, BLACK=600)},
        "2026-07-02": {"1.5": dict(RED=3300, YELLOW=3000, BLUE=2700, BLACK=9600, GREEN=13800), "2.5": dict(RED=1500, YELLOW=1500, BLUE=1200, BLACK=3900)},
        "2026-07-16": {"1.5": dict(RED=3900, YELLOW=3600, BLUE=2400, BLACK=10500, GREEN=14700), "2.5": dict(RED=1500, YELLOW=1500, BLUE=1800, BLACK=4800)}}
def reset():
    f = os.path.join(ROOT, "data", "projects", "thoth-mall", "billing.json")
    if os.path.exists(f): os.remove(f)
    rows = [("Electrical", "SWITCH 6A MODULE BOX", "Box", "2026-07-04", 5, 0.0),
            ("Electrical", "PVC CLAMP 20MM", "Nos", "2026-07-14", 500, 0.0),
            ("Electrical", "GI BOX 2 MODULE", "Nos", "2026-07-09", 200, 0.0),
            ("Plumbing", "CPVC PIPE 25MM", "Mtr", "2026-07-10", 200, 0.0)]
    for dte, bysz in WIRE.items():
        for size, cols in bysz.items():
            for col, m in cols.items():
                rows.append(("Electrical", f"{size}SQMM FRLS WIRE {col}", "Mtr", dte, m, 0.0))
    df = pd.DataFrame(rows, columns=["service", "material", "unit", "date", "qty_in", "opening"])
    df["date"] = pd.to_datetime(df["date"]); df["qty_out"] = 0.0; df["balance"] = np.nan
    rd = os.path.join(ROOT, "data", "runs", "r_stockinbox"); os.makedirs(rd, exist_ok=True)
    df.to_parquet(os.path.join(rd, "daily.parquet"))
    json.dump({"run_id": "r_stockinbox", "project_slug": "thoth-mall", "created": "2099-01-01T00:00:00"}, open(os.path.join(rd, "meta.json"), "w"))
    lf = os.path.join(ROOT, "data", "projects", "thoth-mall", "links.json")
    links = json.load(open(lf)) if os.path.exists(lf) else {}
    links.setdefault("Electrical", {})["1.1.6"] = [{"material": "SWITCH 6A MODULE BOX", "factor": 0.1}]
    links["Electrical"]["1.1.13"] = [{"material": "GI BOX 2 MODULE", "factor": None}]
    json.dump(links, open(lf, "w"))
    g = get(P + "/packages/Electrical")["groups"]
    grp = lambda x: {"sheet": x["sheet"], "from": x["from"], "to": x["to"]}
    pw = next(x for x in g if x["name"].startswith("POINT WIRING"))
    post(P + "/packages/Electrical", {"items": [{"name": "Only point wiring", "groups": [grp(pw)]}, {"name": "Other Electrical", "rest": True}]})
    post(P + "/setup", {"service": PW, "mode": "separate", "supply_pct": 75, "install_pct": 15, "supply_source": "challan"})
    post(P + "/setup", {"service": "Electrical › Other Electrical", "mode": "separate", "supply_pct": 75, "install_pct": 15, "supply_source": "stock"})
with sync_playwright() as p:
    b = p.chromium.launch()
    for name, (w, h) in {"desk": (1366, 900), "phone": (390, 844)}.items():
        reset()
        ctx = b.new_context(viewport={"width": w, "height": h}, has_touch=(name == "phone"), is_mobile=(name == "phone"))
        ctx.add_init_script("window.__BL_LIVE_MS = 3600000")
        pg = ctx.new_page()
        pg.on("pageerror", lambda e, n=name: errs.append(("pageerror", n, str(e))))
        pg.on("console", lambda m, n=name: errs.append(("console", n, m.text)) if m.type == "error" else None)
        pg.goto(BASE); pg.wait_for_timeout(800)
        pg.click('.spnav button[data-v="billing"]'); pg.wait_for_selector(".bl-kpis")
        pg.click(f'[data-ch="{PW}"]'); pg.wait_for_selector("#bl-ib .bl-ibr")
        pg.screenshot(path=f"{OUT}/{name}_1_inbox.png", full_page=True); noscroll(pg, name + " inbox")
        n0 = pg.locator("#bl-iblist .bl-ibr").count()
        assert n0 == 3, n0                                    # switch, gi box, pvc clamp (plumbing + wire are not rows)
        assert "CPVC" not in pg.inner_text("#bl-iblist") and "WIRE" not in pg.inner_text("#bl-iblist")
        if name == "phone":
            fingers(pg, "#bl-ibq", "search", 44); fingers(pg, ".bl-ibr .act .btn", "row buttons", 44); fingers(pg, ".bl-ibt", "tabs", 44)
        # ---- search
        cnt = lambda: pg.locator("#bl-iblist .bl-ibr").count()
        pg.fill("#bl-ibq", "gi box"); assert cnt() == 1                       # by stock material
        pg.fill("#bl-ibq", "two way"); assert cnt() == 1                      # by the BOQ line text (1.1.6 only)
        pg.fill("#bl-ibq", "1.1.13"); assert cnt() == 1                       # by BOQ line code
        pg.fill("#bl-ibq", "switch"); assert cnt() == 2                       # 1.1.13's BOQ text says "switch" too, so both match
        pg.fill("#bl-ibq", "nothing like this"); assert cnt() == 0 and "Nothing matches" in pg.inner_text("#bl-iblist")
        pg.fill("#bl-ibq", "")
        # ---- the add menu offers both ways
        pg.click("#bl-chadd"); pg.wait_for_selector("#bl-chmenu:not([hidden])")
        pg.screenshot(path=f"{OUT}/{name}_2_menu.png")
        assert "From stock register" in pg.inner_text("#bl-chmenu") and "Type a challan" in pg.inner_text("#bl-chmenu")
        pg.click("#bl-chfrom")
        # ---- add the switch delivery with the real challan number
        row = pg.locator("#bl-iblist .bl-ibr", has_text="SWITCH 6A MODULE BOX")
        row.locator("[data-ibno]").fill("1188")
        row.locator("[data-ibadd]").click()
        pg.wait_for_function("document.getElementById('bl-toast') && /added to RA-01/.test(document.getElementById('bl-toast').textContent)")
        pg.wait_for_selector(".bl-ch")
        ch = get(P + f"/challans/{u.quote(PW)}")["challans"]
        assert len(ch) == 1 and ch[0]["no"] == "1188" and ch[0]["no_missing"] is False and ch[0]["lines"][0]["qty"] == 50, ch
        assert pg.locator("#bl-iblist .bl-ibr").count() == 2, "left the inbox"
        pg.screenshot(path=f"{OUT}/{name}_3_added.png", full_page=True); noscroll(pg, name + " challans")
        # ---- the unlinked material asks to be linked; the engineer says it is not billable
        pg.locator("#bl-iblist .bl-ibr", has_text="PVC CLAMP").locator("[data-ibmark]").click()
        pg.wait_for_function("document.querySelectorAll('#bl-iblist .bl-ibr').length === 1")
        # ---- wire, from the register
        pg.click("#bl-ibwire"); pg.wait_for_selector(".bl-wmt")
        pg.screenshot(path=f"{OUT}/{name}_4_wire_confirm.png", full_page=True); noscroll(pg, name + " wire confirm")
        assert pg.locator(".bl-wmt tbody tr").count() == 9
        if name == "phone": fingers(pg, ".bl-segb", "source switch", 44)
        pg.click("#bl-wmconf")
        pg.wait_for_function("document.querySelectorAll('.bl-wtab thead [data-wch]').length === 3")
        sums = pg.eval_on_selector_all(".bl-wtab:not(.bl-wmt) tr.sum td.n b", "els => els.map(e => e.textContent.trim())")
        assert sums == ["20,100", "9,900"], sums
        pg.screenshot(path=f"{OUT}/{name}_5_wire_from_stock.png", full_page=True); noscroll(pg, name + " wire from stock")
        for code, size, metres in (("1.1.5", "1.5", 20100), ("1.1.3", "2.5", 6000), ("1.1.4", "2.5", 3900)):
            pg.fill("#bl-wfind", code); pg.wait_for_selector(f'#bl-wfound .bl-hit[data-c="{code}"]')
            pg.click(f'#bl-wfound .bl-hit[data-c="{code}"]')
            blk = pg.locator(".bl-wb").last
            blk.locator("select[data-sz]").select_option(size)
            blk.locator("[data-m]").fill(str(metres))
            pg.fill("#bl-wfind", "")
        assert pg.inner_text("#bl-wtot").strip() == "₹ 64,28,813", pg.inner_text("#bl-wtot")
        pg.click("#bl-wsave"); pg.wait_for_function("document.getElementById('bl-toast') && document.getElementById('bl-toast').textContent.includes('Metres saved')")
        st = get(f"{P}/wire/{u.quote(PW)}")
        assert {l["item_code"]: l["points"] for l in st["lines"]} == {"1.1.3": 300, "1.1.4": 130, "1.1.5": 4020}
        assert st["source"] == "stock" and all(c["src"] for c in st["challans"])
        # ---- the RA: the switch delivery (50 Nos) and the wire points are both in the supply
        ra = post(P + "/ra", {"service": PW, "series": "supply"})["id"]
        bill = get(P + f"/ra/{ra}")
        by = {l["item_code"]: l for l in bill["lines"]}
        assert by["1.1.6"]["supply"]["this"] == 50 and by["1.1.5"]["supply"]["this"] == 4020, (by["1.1.6"]["supply"], by["1.1.5"]["supply"])
        # ---- the invoice workbook still carries both CHALLAN sheets
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(u.urlopen(BASE + P + f"/invoice/{ra}").read()))
        assert "WIRE CHALLAN" in wb.sheetnames and "CHALLAN" in wb.sheetnames, wb.sheetnames
        ctx.close()
    b.close()
assert not errs, errs[:5]
print("billing stock inbox chromium: all passed")
