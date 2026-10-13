"""Billing -- Supply + Installation TOGETHER, for us (internal view, not a bill).
Scenario from the report (Fire): Supply RA 1 is already invoiced, Installation RA 1 is an open draft. A colleague wanted both in one
place, but the client still gets two separate invoices, so a real "Together" RA (one invoice, a new RA number) is the wrong tool.
GET /{slug}/combined?service=... gives one workbook with both stages side by side, reading the latest Supply RA and the latest
Installation RA. It creates no RA, uses no invoice number and locks nothing. Same harness as test_billing_extra.py
(Fire: valve BOQ 4 @ 5000, pump BOQ 2 @ 100000, GI pipe BOQ 100 @ 800; supply 70 / install 30)."""
import io
import sys
import zipfile
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl
import pytest

from backend.test_billing import env, prepared, make_run, link, progress, line, STOCK1, SLUG, _recalc   # noqa: F401
from backend.test_billing_challan import cs, ch, new_ra, bill, B
from backend.test_billing_invoice_cover import invoice_of, FULL, meta, full
from backend.test_billing_extra import accept

URL = f"{B}/combined"


def typed(client, ra, code, stage, qty, **kw):
    r = client.post(f"{B}/ra/{ra}/line", json={"item_code": code, "stage": stage, "qty": qty, **kw})
    assert r.status_code == 200, r.text


def combined(client, svc="Fire"):
    r = client.get(URL, params={"service": svc})
    assert r.status_code == 200, r.text
    return openpyxl.load_workbook(io.BytesIO(r.content)), r


def fire_both(client):
    """Supply RA 1 invoiced (valve 6 of 4 accepted as extra, pump 2), Installation RA 1 an open draft (typed)."""
    full(client)
    cs(client)
    ch(client, "700", "2026-07-04", [("2", 6), ("1", 2)])
    sup = new_ra(client)
    accept(client, sup, [{"item_code": "2", "stage": "supply"}])
    client.post(f"{B}/ra/{sup}/invoice", json={"no": "FIRE/S/01", "date": "2026-07-10"})
    assert client.post(f"{B}/ra/{sup}/lock").status_code == 200
    ins = new_ra(client, "install")
    typed(client, ins, "2", "install", 3)
    typed(client, ins, "1", "install", 1)
    return sup, ins


def boq(wb):
    return wb["Annexure B-BOQ"]


def row_of(ws, desc):
    return next(r for r in range(9, ws.max_row + 1) if ws.cell(r, 2).value == desc)


def descs(ws):
    return [ws.cell(r, 2).value for r in range(9, ws.max_row + 1) if ws.cell(r, 3).value]


# ----------------------------------------------------------------- availability and refusals
def test_nothing_to_combine_without_an_ra_and_unknown_service_is_404(env):
    client, _, _ = env
    cs(client)
    r = client.get(URL, params={"service": "Fire"})
    assert r.status_code == 409 and "no Supply or Installation RA" in r.json()["detail"]
    assert client.get(URL, params={"service": "Nope"}).status_code == 404


def test_overview_says_which_ra_each_stage_comes_from(env):
    client, _, _ = env
    cs(client)
    assert client.get(B).json()["services"][0]["both_view"] == {"supply": None, "install": None}
    sup, ins = fire_both(client)
    bv = next(s for s in client.get(B).json()["services"] if s["service"] == "Fire")["both_view"]
    assert (bv["supply"]["id"], bv["supply"]["status"], bv["supply"]["no"]) == (sup, "locked", 1)
    assert (bv["install"]["id"], bv["install"]["status"], bv["install"]["no"]) == (ins, "draft", 1)


# ----------------------------------------------------------------- the file: same numbers as the two separate workbooks
def test_every_line_equals_the_separate_supply_and_installation_workbooks(env, tmp_path):
    client, _, _ = env
    sup, ins = fire_both(client)
    wb, resp = combined(client)
    assert wb.sheetnames[:3] == ["Internal Summary", "Annexure B-Summary BOQ", "Annexure B-BOQ"]
    ws = boq(wb)
    assert ws.cell(7, 7).value == "RA-01 · SUPPLY" and ws.cell(7, 14).value.startswith("RA-01 · INSTALLATION")
    s_wb = openpyxl.load_workbook(invoice_of(client, sup, tmp_path, "s.xlsx"), data_only=True)
    i_wb = _recalc(invoice_of(client, ins, tmp_path, "i.xlsx"), tmp_path / "rc")
    s_ws, i_ws = s_wb["Annexure B-BOQ"], i_wb["Annexure B-BOQ"]
    checked = 0
    for d in descs(ws):
        r, rs, ri = row_of(ws, d), row_of(s_ws, d), row_of(i_ws, d)
        for k in range(7):                                       # G..M of the supply block, N..T of the installation block
            for blk, sep in ((7, s_ws.cell(rs, 7 + k).value), (14, i_ws.cell(ri, 7 + k).value)):
                got = ws.cell(r, blk + k).value
                assert (got or 0) == pytest.approx(sep or 0), (d, blk, k)
        checked += 1
    assert checked >= 3
    # the total rows too
    for r in range(9, ws.max_row + 1):
        if str(ws.cell(r, 2).value or "").startswith("TOTAL FOR"):
            rs = next(x for x in range(9, s_ws.max_row + 1) if s_ws.cell(x, 2).value == ws.cell(r, 2).value)
            for k in (4, 5, 6):
                assert (ws.cell(r, 7 + k).value or 0) == pytest.approx(s_ws.cell(rs, 7 + k).value or 0)
                assert (ws.cell(r, 14 + k).value or 0) == pytest.approx(i_ws.cell(rs, 7 + k).value or 0)


def test_internal_summary_totals_and_gst_equal_the_two_invoices(env, tmp_path):
    client, _, _ = env
    sup, ins = fire_both(client)
    wb, _ = combined(client)
    ci = wb["Internal Summary"]
    assert "INTERNAL COPY, not an invoice" in ci["A2"].value
    rows = {ci.cell(r, 1).value: [ci.cell(r, c).value for c in (2, 3, 4)] for r in range(10, 18)}
    s_tot = bill(client, sup)["this_total"]
    i_tot = bill(client, ins)["this_total"]
    assert s_tot > 0 and i_tot > 0
    assert rows["Value of services"] == [pytest.approx(s_tot), pytest.approx(i_tot), pytest.approx(s_tot + i_tot)]
    assert rows["SGST 9%"][0] == pytest.approx(s_tot * .09) and rows["CGST 9%"][1] == pytest.approx(i_tot * .09)
    # each invoice rounds on its own; Both = sum of the two invoice totals
    from backend import billing_cover
    t1 = billing_cover.money(s_tot, 9, 9)["total"]
    t2 = billing_cover.money(i_tot, 9, 9)["total"]
    assert rows["Invoice total"] == [pytest.approx(t1), pytest.approx(t2), pytest.approx(t1 + t2)]
    assert rows["Which RA"][0] == "Supply RA-01 · Invoiced" and rows["Which RA"][1] == "Installation RA-01 · Draft"
    assert rows["Invoice no."][0] == "FIRE/S/01" and rows["Invoice no."][1] == "—"
    words = next(ci.cell(r, 2).value for r in range(17, 24) if ci.cell(r, 1).value == "Amount in words")
    assert words == billing_cover.words(t1 + t2)
    # the separate Supply invoice cover says the same amount
    s_inv = openpyxl.load_workbook(invoice_of(client, sup, tmp_path, "s2.xlsx"), data_only=True)["RA-01 INVOICE"]
    nums = [c.value for row in s_inv.iter_rows() for c in row if isinstance(c.value, (int, float))]
    assert pytest.approx(t1) in nums


def test_annexure_summary_has_supply_installation_and_total_per_group(env):
    client, _, _ = env
    sup, ins = fire_both(client)
    wb, _ = combined(client)
    w = wb["Annexure B-Summary BOQ"]
    assert [w.cell(5, c).value for c in (4, 7, 10)] == ["SUPPLY @ 70%", "INSTALLATION @ 30%", "TOTAL"]
    gr = next(r for r in range(8, w.max_row + 1) if w.cell(r, 2).value == "Grand Total")
    s_tot, i_tot = bill(client, sup)["this_total"], bill(client, ins)["this_total"]
    assert w.cell(gr, 5).value == pytest.approx(s_tot) and w.cell(gr, 8).value == pytest.approx(i_tot)
    assert w.cell(gr, 11).value == pytest.approx(s_tot + i_tot)
    assert w.cell(gr, 12).value == pytest.approx(w.cell(gr, 6).value + w.cell(gr, 9).value)


# ----------------------------------------------------------------- Extra over BOQ appears by itself
def test_accepted_extra_gets_its_own_sheet_and_yellow_cells(env):
    client, _, _ = env
    fire_both(client)
    wb, _ = combined(client)
    assert "Extra over BOQ" in wb.sheetnames
    x = wb["Extra over BOQ"]
    assert "INTERNAL COPY" in x["A5"].value and "Supply RA-01 and Installation RA-01" in x["B3"].value
    data = [[c.value for c in row] for row in x.iter_rows(min_row=8, max_row=x.max_row)]
    row = next(r for r in data if r[1] and "alve" in r[1])
    assert row[3] == "Supply" and row[4] == 4 and row[5] == 6 and row[6] == 2           # boq, billed, extra
    assert row[8] == pytest.approx(5000 * 2 * .70)
    ws = boq(wb)
    r = row_of(ws, row[1])
    assert ws.cell(r, 9).fill.fgColor.rgb.endswith("FFE699") and ws.cell(r, 9).comment is not None
    assert "above the BOQ quantity" in ws.cell(r, 9).comment.text


def test_no_extra_no_extra_sheet(env):
    client, _, _ = env
    full(client)
    cs(client)
    ch(client, "700", "2026-07-04", [("2", 3)])
    new_ra(client)
    wb, _ = combined(client)
    assert "Extra over BOQ" not in wb.sheetnames


# ----------------------------------------------------------------- one stage only: the other block is empty, never invented
def test_only_supply_exists_installation_block_is_empty(env):
    client, _, _ = env
    full(client)
    cs(client)
    ch(client, "700", "2026-07-04", [("2", 3), ("1", 1)])
    new_ra(client)
    wb, _ = combined(client)
    ws = boq(wb)
    assert ws.cell(7, 14).value == "INSTALLATION · NO RA YET"
    for r in range(9, ws.max_row + 1):
        assert all(ws.cell(r, c).value is None for c in range(14, 21)), r
    ci = wb["Internal Summary"]
    vals = {ci.cell(r, 1).value: [ci.cell(r, c).value for c in (2, 3, 4)] for r in range(10, 18)}
    assert vals["Which RA"][1] == "Installation: no RA yet"
    assert vals["Value of services"][1] == "—" and vals["Invoice total"][1] == "—"
    assert vals["Value of services"][2] == pytest.approx(vals["Value of services"][0])
    w = wb["Annexure B-Summary BOQ"]
    gr = next(r for r in range(8, w.max_row + 1) if w.cell(r, 2).value == "Grand Total")
    assert [w.cell(gr, c).value for c in (7, 8, 9)] == [None, None, None]


def test_only_installation_exists_supply_block_is_empty(env):
    client, _, _ = env
    full(client)
    cs(client)
    ins = new_ra(client, "install")
    typed(client, ins, "2", "install", 2)
    wb, _ = combined(client)
    ws = boq(wb)
    assert ws.cell(7, 7).value == "SUPPLY · NO RA YET"
    assert all(ws.cell(r, c).value is None for r in range(9, ws.max_row + 1) for c in range(7, 14))
    assert next(r for r in boq(wb).iter_rows(min_row=9) if r[1].value and "alve" in str(r[1].value))[15].value == 2


# ----------------------------------------------------------------- latest RA of each stage; previous comes from the earlier ones
def test_the_latest_ra_of_each_stage_is_used_and_previous_qty_comes_from_the_earlier_one(env):
    client, _, _ = env
    sup, ins = fire_both(client)
    client.post(f"{B}/ra/{ins}/lock")
    ch(client, "701", "2026-08-02", [("1", 1)], ra=2)                        # a second Supply RA has a reason to exist
    sup2 = new_ra(client)
    ins2 = new_ra(client, "install")
    typed(client, ins2, "2", "install", 1)
    wb, _ = combined(client)
    ws = boq(wb)
    assert ws.cell(7, 7).value.startswith("RA-02 · SUPPLY") and ws.cell(7, 14).value.startswith("RA-02 · INSTALLATION")
    r = row_of(ws, next(d for d in descs(ws) if "alve" in d))
    assert ws.cell(r, 15).value == 3 and ws.cell(r, 16).value == 1 and ws.cell(r, 17).value == 4        # prev 3 (RA1), this 1, total 4
    ci = wb["Internal Summary"]
    assert ci["B10"].value == "Supply RA-02 · Draft" and ci["C10"].value == "Installation RA-02 · Draft"


# ----------------------------------------------------------------- it is a view: nothing is created, locked or numbered
def test_downloading_it_changes_nothing(env):
    client, d, _ = env
    sup, ins = fire_both(client)
    before = client.get(B).json()
    state = [(r["id"], r["status"], r["no"], r["inv_no"], r["this_total"]) for r in before["ras"]]
    for _ in range(2):
        combined(client)
    after = client.get(B).json()
    assert [(r["id"], r["status"], r["no"], r["inv_no"], r["this_total"]) for r in after["ras"]] == state
    assert len(after["ras"]) == 2
    assert next(r for r in after["ras"] if r["id"] == ins)["inv_no"] == ""            # no invoice number was used


def test_every_sheet_says_it_is_internal_and_the_file_name_says_which_ras(env):
    client, _, _ = env
    fire_both(client)
    wb, resp = combined(client)
    for n in ("Internal Summary", "Annexure B-BOQ", "Extra over BOQ"):
        assert "INTERNAL COPY" in " ".join(str(c.value) for row in wb[n].iter_rows(min_row=1, max_row=6) for c in row if c.value)
    assert "INTERNAL COPY" in wb["Annexure B-Summary BOQ"]["B5"].value
    cd = resp.headers["content-disposition"]
    assert "Supply+Installation_TOGETHER_Supply-RA01_Install-RA01_INTERNAL.xlsx" in cd


# ----------------------------------------------------------------- GST not set: no guessed total
def test_gst_not_set_shows_a_dash_and_says_so(env):
    client, _, _ = env
    cs(client)
    ch(client, "700", "2026-07-04", [("2", 3)])
    new_ra(client)
    ins = new_ra(client, "install")
    typed(client, ins, "2", "install", 1)
    wb, _ = combined(client)
    ci = wb["Internal Summary"]
    vals = {ci.cell(r, 1).value: [ci.cell(r, c).value for c in (2, 3, 4)] for r in range(10, 18)}
    assert vals["Value of services"][2] > 0
    assert vals["Invoice total"] == ["—", "—", "—"] and vals["SGST"] == ["—", "—", "—"]
    assert any("not set in Billing setup" in str(ci.cell(r, 2).value) for r in range(17, 24))


# ----------------------------------------------------------------- a real Together RA (one invoice): GST once
def test_one_together_ra_shows_gst_once(env):
    client, _, _ = env
    full(client)
    cs(client, mode="together")
    ch(client, "700", "2026-07-04", [("2", 3)])
    ra = new_ra(client, "together")
    typed(client, ra, "2", "install", 2)
    wb, _ = combined(client)
    ci = wb["Internal Summary"]
    vals = {ci.cell(r, 1).value: [ci.cell(r, c).value for c in (2, 3, 4)] for r in range(10, 18)}
    tot = bill(client, ra)["this_total"]
    assert vals["Value of services"][2] == pytest.approx(tot)
    assert vals["Invoice total"][:2] == ["—", "—"] and isinstance(vals["Invoice total"][2], (int, float))
    assert any("one invoice" in str(ci.cell(r, 2).value) for r in range(17, 24))


# ----------------------------------------------------------------- inside the Download
def test_download_adds_it_beside_the_ticked_ras(env):
    client, _, _ = env
    sup, ins = fire_both(client)
    only = client.get(f"{B}/download", params={"invoice": 0, "both": ["Fire"]})
    assert only.status_code == 200 and only.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "TOGETHER" in only.headers["content-disposition"]
    z = client.get(f"{B}/download", params={"ras": f"{sup},{ins}", "both": ["Fire"]})
    assert z.headers["content-type"] == "application/zip"
    names = zipfile.ZipFile(io.BytesIO(z.content)).namelist()
    assert len(names) == 3 and sum("TOGETHER" in n for n in names) == 1


def test_download_one_workbook_mode_still_adds_it_as_its_own_file(env):
    client, _, _ = env
    sup, ins = fire_both(client)
    z = client.get(f"{B}/download", params={"ras": f"{sup},{ins}", "mode": "one", "both": ["Fire"]})
    assert z.status_code == 200 and z.headers["content-type"] == "application/zip"
    names = zipfile.ZipFile(io.BytesIO(z.content)).namelist()
    assert len(names) == 2 and sum("RA_invoices" in n for n in names) == 1 and sum("TOGETHER" in n for n in names) == 1


def test_download_both_for_a_service_without_ra_is_refused_clearly(env):
    client, _, _ = env
    cs(client)
    r = client.get(f"{B}/download", params={"invoice": 0, "both": ["Fire"]})
    assert r.status_code == 409 and "no Supply or Installation RA" in r.json()["detail"]


def test_the_separate_workbooks_are_unchanged_by_the_combined_code(env, tmp_path):
    client, _, _ = env
    sup, ins = fire_both(client)
    a = invoice_of(client, sup, tmp_path, "a.xlsx").read_bytes()
    combined(client)
    b = invoice_of(client, sup, tmp_path, "b.xlsx").read_bytes()
    wa = openpyxl.load_workbook(io.BytesIO(a))
    wb_ = openpyxl.load_workbook(io.BytesIO(b))
    assert wa.sheetnames == wb_.sheetnames
    for n in wa.sheetnames:
        assert [[c.value for c in r] for r in wa[n].iter_rows()] == [[c.value for c in r] for r in wb_[n].iter_rows()]


def test_dump(env):
    """Not a check: MI_DUMP_FIXTURE=<path> writes the Fire sample to look at it."""
    import os
    out = os.environ.get("MI_DUMP_FIXTURE")
    if not out:
        pytest.skip("only with MI_DUMP_FIXTURE")
    client, _, _ = env
    sup, ins = fire_both(client)
    Path(out).write_bytes(client.get(URL, params={"service": "Fire"}).content)
    if os.environ.get("MI_DUMP_JSON"):
        import json
        Path(os.environ["MI_DUMP_JSON"]).write_text(json.dumps({"slug": SLUG, "overview": client.get(B).json(), "sup": sup, "ins": ins}))
