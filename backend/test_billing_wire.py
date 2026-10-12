"""Wire metres -> points (point wiring is bought as wire, billed as points).
Real scenario: Thoth Mall Electrical "Only point wiring", RA 1. The site's own Excel (sheet WIRE CHALLAN) takes the
metres of the phase wires Red + Yellow + Blue (neutral Black and earth Green run beside them and are not counted),
divides what is given to a BOQ line by the average point length the BOQ states for it, and bills those points:
20,100 m / 5 = 4,020 points (1.1.5), 6,000 m / 20 = 300 (1.1.3), 3,900 m / 30 = 130 (1.1.4) = Rs 64,28,813 at 75%.
The descriptions, quantities and rates below are the real ones from the Thoth BOQ."""
import os
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from backend import siteprogress as sp, billing, structure

SLUG = "test-wire"
B = f"/api/billing/{SLUG}"
HEAD = ["Internal Ref No", "Client Ref No", "BOQ", "BOQ Unit", "Order Quantity", "Design Quantity", "Rate",
        "Discount %", "HSN Code", "Tax", "Progress Method", "Billing Term", "Summary Item"]
D = {
    "1.1.1": "One light point controlled by one single way 6 A switch.with average point length of 15 Rmtr.",
    "1.1.2": "Loop light point controlled by same switch same as 13.1.1.with average point length of 5 Rmtr. ",
    "1.1.3": "One light point controlled by SP MCB. Cost inculdes rate of MS conduit & Wire with average point length upto & including 20 Rmtr. (MCB Rates does not inclueded) ",
    "1.1.4": "One light point controlled by SP MCB. Cost inculdes rate of MS conduit & Wire with average point length of 30 Rmtr. (MCB Rates does not inclueded) ",
    "1.1.5": "Loop light point controlled by same mcb same as 13.1.3.1&2\nCost inculdes rate of MS conduit & Wire with average point length upto & including 5 Rmtr.",
    "1.1.6": "One light point controlled by two, two way 6 A switch",
    "1.1.13": "Exhaust fan outlet point with 6A 5 pin socket controlled by one 5amp single way switch located at different location.(Rate includes cost of wiring & piping with average point length of 15 Rmtr0.)",
    "1.1.18": "32 A 5 pin x 415 Volt (3 Phase) Metal clad socket outlet point with 32 Amp. FPMCB but without cost of wiring and piping..",
    "1.1.21": "SITC computer power supply point with 2 No.s of 6 amp 5 pin plugs with 1 No.. of 16 Amp. Plug with 1 No. 16 A  Switch with indicator & wiring of computer point with average point length of 10 Rmtr)",
    "1.2.1": "Supplying and Providing mains with 4 No.s. of 6 sq mm Cu. PVC insulated cable",
}
ROWS = [("1.1.3", 493, 6025), ("1.1.4", 211, 8895), ("1.1.5", 4277, 1395), ("1.1.6", 55, 5850), ("1.1.18", 11, 11650),
        ("1.1.21", 28, 4750), ("1.2.1", 121, 620)]
RYB = [("RED", "YELLOW", "BLUE"), ("BLACK", "GREEN")]


def wb_path(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Electrical"
    ws.append(HEAD)
    ws.append(["1", "1", "Internal Wiring", None, None, None, None, None, None, None, None, None, None])
    for i, (ref, qty, rate) in enumerate(ROWS, start=2):
        ws.append([str(i), ref, D[ref], "Nos" if not ref.startswith("1.2") else "Rmt", qty, qty, rate,
                   None, None, "GST @ 18%", None, None, None])
    wb.save(path)


@pytest.fixture
def env(tmp_path, monkeypatch):
    monkeypatch.setattr(sp, "PROJECTS", tmp_path / "projects")
    monkeypatch.setattr(sp, "RUNS", tmp_path / "runs")
    sp.PROJECTS.mkdir(parents=True)
    sp.RUNS.mkdir(parents=True)
    for c in (sp._read_forecast_parquet_cached, sp._forecast_pool_cached, sp._full_run_rows_cached, sp._linkage_match_cached):
        c.cache_clear()
    app = FastAPI()
    app.include_router(sp.router)
    app.include_router(billing.router)
    client = TestClient(app)
    wb_path(tmp_path / "boq.xlsx")
    with open(tmp_path / "boq.xlsx", "rb") as fh:
        assert client.post(f"/api/siteprogress/{SLUG}/boq", files={"file": ("boq.xlsx", fh)}).status_code == 200
    d = sp._dir(SLUG)
    (d / "structure.json").write_text(structure.hotel("T", floors=["1"], room_labels=["a"]).to_json())
    r = client.post(f"{B}/setup", json={"service": "Electrical", "mode": "separate", "supply_pct": 75, "install_pct": 15,
                                        "supply_source": "challan"})
    assert r.status_code == 200, r.text
    return client, d


def code_of(client, ref):
    ls = client.get(f"{B}/lines/Electrical").json()["lines"]
    return next(l["item_code"] for l in ls if l.get("internal_ref") == ref or l["desc"] == D[ref])


def code(client, ref):
    ls = client.get(f"{B}/wire/Electrical").json()
    all_ = {e["ref"]: e["item_code"] for e in ls["eligible"]}
    all_.update({e["ref"]: e["item_code"] for e in ls["free_lines"]})
    return all_[ref]


def wch(client, no, date, sizes, **kw):
    return client.post(f"{B}/wire/Electrical/challans", json={"no": no, "date": date, "sizes": sizes, **kw})


def sz(size, **cols):
    return {"size": size, "colours": {k.upper(): v for k, v in cols.items()}}


def real_wire(client):
    """The three real wire challans of the site's WIRE CHALLAN sheet (1207, 615, 617), metres per colour."""
    for no, date, a15, a25 in (
            ("1207", "2026-03-28", dict(red=300, yellow=600, blue=300, black=900, green=1500), dict(red=300, yellow=300, blue=300, black=600)),
            ("615", "2026-07-02", dict(red=3300, yellow=3000, blue=2700, black=9600, green=13800), dict(red=1500, yellow=1500, blue=1200, black=3900)),
            ("617", "2026-07-16", dict(red=3900, yellow=3600, blue=2400, black=10500, green=14700), dict(red=1500, yellow=1500, blue=1800, black=4800))):
        r = wch(client, no, date, [sz("1.5", **a15), sz("2.5", **a25)])
        assert r.status_code == 200, r.text


def alloc(client, rows):
    return client.put(f"{B}/wire/Electrical/alloc", json={"rows": rows})


def real_alloc(client):
    return alloc(client, [{"item_code": code(client, "1.1.5"), "size": "1.5", "metres": 20100},
                          {"item_code": code(client, "1.1.3"), "size": "2.5", "metres": 6000},
                          {"item_code": code(client, "1.1.4"), "size": "2.5", "metres": 3900}])


def new_ra(client):
    r = client.post(f"{B}/ra", json={"service": "Electrical", "series": "supply"})
    assert r.status_code == 200, r.text
    return r.json()["id"]


# ---------------------------------------------------------------- the BOQ's own point length
@pytest.mark.parametrize("ref,per,upto", [("1.1.1", 15, False), ("1.1.2", 5, False), ("1.1.3", 20, True), ("1.1.4", 30, False),
                                          ("1.1.5", 5, True), ("1.1.13", 15, False), ("1.1.21", 10, False)])
def test_point_length_is_read_from_the_boq_text_exactly(ref, per, upto):
    r = billing._point_len(D[ref])
    assert r and r["per"] == per and r["upto"] is upto


@pytest.mark.parametrize("ref", ["1.1.6", "1.1.18", "1.2.1"])
def test_a_line_whose_text_states_no_point_length_gets_none(ref):
    assert billing._point_len(D[ref]) is None


# ------------------------------------------------------------------ the real RA 1
def test_real_ra1_wire_becomes_4020_300_130_points_and_the_supply_bill(env):
    client, _ = env
    real_wire(client)
    st = client.get(f"{B}/wire/Electrical").json()
    # only Red + Yellow + Blue are counted; Black and Green run beside them
    assert {x["size"]: (x["counted"], x["given"], x["left"]) for x in st["size_status"]} == {"1.5": (20100, 0, 20100), "2.5": (9900, 0, 9900)}
    assert real_alloc(client).status_code == 200
    st = client.get(f"{B}/wire/Electrical").json()
    pts = {l["ref"]: l["points"] for l in st["lines"]}
    assert pts == {"1.1.5": 4020, "1.1.3": 300, "1.1.4": 130}
    assert [x["left"] for x in st["size_status"]] == [0, 0]
    assert st["totals"]["supply_this"] == pytest.approx(6428812.5)
    b = client.get(f"{B}/ra/{new_ra(client)}").json()
    assert b["this_total"] == pytest.approx(6428812.5) and int(b["this_total"] + .5) == 6428813
    by = {l["item_code"]: l for l in b["lines"]}
    assert by[code(client, "1.1.5")]["supply"]["this"] == 4020 and by[code(client, "1.1.5")]["supply"]["source"] == "challan"
    assert "Wire RA-01" in by[code(client, "1.1.5")]["supply"]["note"]


def test_the_points_show_as_one_wire_challan_per_ra_and_cannot_be_edited_as_a_normal_challan(env):
    client, _ = env
    real_wire(client)
    real_alloc(client)
    chs = client.get(f"{B}/challans/Electrical").json()
    w = [c for c in chs["challans"] if c["virtual"]]
    assert len(w) == 1 and w[0]["no"] == "Wire RA-01" and w[0]["ra"] == 1 and w[0]["date"] == "2026-07-16"
    assert sum(l["qty"] for l in w[0]["lines"]) == 4450
    assert chs["wire"]["available"] is True
    assert client.post(f"{B}/challans/Electrical/{w[0]['id']}", json={"no": "x", "date": "2026-01-01", "lines": []}).status_code == 409
    assert client.delete(f"{B}/challans/Electrical/{w[0]['id']}").status_code == 409


# ------------------------------------------------------------------ guard rails
def test_more_metres_than_were_received_is_refused(env):
    client, _ = env
    real_wire(client)
    r = alloc(client, [{"item_code": code(client, "1.1.5"), "size": "1.5", "metres": 20101}])
    assert r.status_code == 409 and "20101" in r.text


def test_a_line_can_never_be_claimed_above_its_boq_points(env):
    client, _ = env
    assert wch(client, "9", "2026-07-01", [sz("2.5", red=4000, yellow=4000, blue=4000)]).status_code == 200
    # 1.1.4: BOQ 211 points x 30 m = 6,330 m at most
    assert alloc(client, [{"item_code": code(client, "1.1.4"), "size": "2.5", "metres": 6330}]).status_code == 200      # exactly 211 points
    r = alloc(client, [{"item_code": code(client, "1.1.4"), "size": "2.5", "metres": 6390}])
    assert r.status_code == 409 and "211" in r.text and "6330" in r.text


def test_only_whole_points_count_and_the_rest_waits(env):
    client, _ = env
    wch(client, "9", "2026-07-01", [sz("1.5", red=100)])
    assert alloc(client, [{"item_code": code(client, "1.1.5"), "size": "1.5", "metres": 7}]).status_code == 200
    l = client.get(f"{B}/wire/Electrical").json()["lines"][0]
    assert (l["metres"], l["points"]) == (7, 1)          # 7 m / 5 m = 1 point; 2 m wait for the next metres


def test_a_line_with_no_length_in_the_boq_needs_a_typed_length(env):
    client, _ = env
    wch(client, "9", "2026-07-01", [sz("1.5", red=100)])
    c6 = code(client, "1.1.6")
    assert c6 in {x["item_code"] for x in client.get(f"{B}/wire/Electrical").json()["free_lines"]}
    r = alloc(client, [{"item_code": c6, "size": "1.5", "metres": 60}])
    assert r.status_code == 400 and "average point length" in r.text
    assert alloc(client, [{"item_code": c6, "size": "1.5", "metres": 60, "per": 12}]).status_code == 200
    l = client.get(f"{B}/wire/Electrical").json()["lines"][0]
    assert (l["points"], l["typed"]) == (5, True)


def test_a_line_takes_wire_or_challans_never_both(env):
    client, _ = env
    wch(client, "9", "2026-07-01", [sz("1.5", red=100)])
    c5 = code(client, "1.1.5")
    assert alloc(client, [{"item_code": c5, "size": "1.5", "metres": 50}]).status_code == 200
    r = client.post(f"{B}/challans/Electrical", json={"no": "A1", "date": "2026-07-02", "lines": [{"item_code": c5, "qty": 3}]})
    assert r.status_code == 409 and "wire" in r.text
    c3 = code(client, "1.1.3")
    assert client.post(f"{B}/challans/Electrical", json={"no": "A2", "date": "2026-07-02", "lines": [{"item_code": c3, "qty": 3}]}).status_code == 200
    r = alloc(client, [{"item_code": c3, "size": "1.5", "metres": 40}])
    assert r.status_code == 409 and "by hand" in r.text


def test_wire_challan_number_must_be_unique_and_needs_metres(env):
    client, _ = env
    assert wch(client, "9", "2026-07-01", [sz("1.5", red=100)]).status_code == 200
    assert wch(client, "9", "2026-07-02", [sz("1.5", red=5)]).status_code == 409
    assert wch(client, "10", "2026-07-02", [sz("1.5", red=0)]).status_code == 400
    assert wch(client, "11", "2026-07-02", [{"size": "1.5", "colours": {"BROWN": 5}}]).status_code == 400    # not a colour of this code


# ------------------------------------------------------------ colour code (industry: legacy vs IEC)
def test_the_colour_code_decides_which_wires_are_phases(env):
    client, _ = env
    r = client.put(f"{B}/wire-scheme", json={"scheme": "iec"})
    assert r.status_code == 200
    assert wch(client, "9", "2026-07-01", [sz("1.5", brown=100, black=200, grey=300, blue=50, green_yellow=400)]).status_code == 200
    st = client.get(f"{B}/wire/Electrical").json()
    assert st["scheme"]["id"] == "iec" and st["size_status"][0]["counted"] == 600      # brown + black + grey; blue is neutral
    assert [c["key"] for c in st["scheme"]["colours"] if c["counted"]] == ["BROWN", "BLACK", "GREY"]
    # the code cannot change once wire is entered (the colours would no longer mean the same)
    assert client.put(f"{B}/wire-scheme", json={"scheme": "ryb"}).status_code == 409
    assert client.put(f"{B}/wire-scheme", json={"scheme": "nonsense"}).status_code == 400


def test_default_colour_code_is_red_yellow_blue(env):
    client, _ = env
    st = client.get(f"{B}/wire/Electrical").json()
    assert st["scheme"]["id"] == "ryb" and [c["key"] for c in st["scheme"]["colours"] if c["counted"]] == ["RED", "YELLOW", "BLUE"]


# --------------------------------------------------------------- RAs, freezing, carry forward
def test_second_ra_carries_the_remainder_forward_and_the_first_is_frozen(env):
    client, _ = env
    real_wire(client)
    real_alloc(client)
    ra1 = new_ra(client)
    assert client.post(f"{B}/ra/{ra1}/lock").status_code == 200
    st = client.get(f"{B}/wire/Electrical").json()
    assert st["open_no"] == 2 and all(c["frozen"] for c in st["challans"])
    r1207 = st["challans"][0]
    assert client.post(f"{B}/wire/Electrical/challans/{r1207['id']}", json={"no": "1207", "date": "2026-03-28", "sizes": [sz("1.5", red=1)]}).status_code == 409
    assert client.delete(f"{B}/wire/Electrical/challans/{r1207['id']}").status_code == 409
    # RA 2: a new wire challan, and 12 m more on 1.1.5 -> 2 more points (cumulative 20112 / 5 = 4022)
    assert wch(client, "700", "2026-08-01", [sz("1.5", red=12)]).status_code == 200
    c5 = code(client, "1.1.5")
    assert alloc(client, [{"item_code": c5, "size": "1.5", "metres": 12}]).status_code == 200
    l = next(x for x in client.get(f"{B}/wire/Electrical").json()["lines"] if x["ref"] == "1.1.5")
    assert (l["points"], l["points_this"]) == (4022, 2)
    assert [(r["size"], r["earlier"], r["this"]) for r in l["rows"]] == [("1.5", 20100, 12)]
    b = client.get(f"{B}/ra/{new_ra(client)}").json()
    assert {x["item_code"]: x["supply"]["this"] for x in b["lines"]}[c5] == 2
    # the invoiced RA 1 did not move
    assert client.get(f"{B}/ra/{ra1}").json()["this_total"] == pytest.approx(6428812.5)


def test_frozen_metres_survive_a_resave_of_the_open_ra(env):
    client, _ = env
    real_wire(client)
    real_alloc(client)
    client.post(f"{B}/ra/{new_ra(client)}/lock")
    assert alloc(client, []).status_code == 200            # the open RA has nothing; RA 1 keeps its points
    st = client.get(f"{B}/wire/Electrical").json()
    assert {l["ref"]: l["points"] for l in st["lines"]} == {"1.1.5": 4020, "1.1.3": 300, "1.1.4": 130}


def test_wire_cannot_be_removed_from_under_metres_already_given(env):
    client, _ = env
    real_wire(client)
    real_alloc(client)
    st = client.get(f"{B}/wire/Electrical").json()
    c615 = next(c for c in st["challans"] if c["no"] == "615")
    assert client.delete(f"{B}/wire/Electrical/challans/{c615['id']}").status_code == 409


def test_a_service_with_wire_cannot_be_split_into_work_orders_later(env):
    client, _ = env
    real_wire(client)
    g = client.get(f"{B}/packages/Electrical").json()["groups"]
    r = client.post(f"{B}/packages/Electrical", json={"items": [
        {"name": "A", "groups": [{"sheet": g[0]["sheet"], "from": g[0]["from"], "to": g[0]["to"]}]}, {"name": "B", "rest": True}]})
    assert r.status_code in (400, 409)


def test_wire_is_offered_only_where_the_boq_states_a_point_length(env):
    client, _ = env
    assert client.get(f"{B}/challans/Electrical").json()["wire"]["available"] is True


def test_the_invoice_workbook_carries_the_wire_challan_and_its_points(env):
    client, _ = env
    real_wire(client)
    real_alloc(client)
    ra = new_ra(client)
    r = client.get(f"{B}/invoice/{ra}")
    assert r.status_code == 200
    import io
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    names = wb.sheetnames
    assert "WIRE CHALLAN" in names and "CHALLAN" in names
    ws = wb["WIRE CHALLAN"]
    cells = {c.value for row in ws.iter_rows() for c in row if c.value is not None}
    assert {"WIRE 1.5", "RED", "YELLOW", "BLUE", "BLACK", "GREEN", 20100, 4020, 300, 130, "1.1.5", "1.1.3", "1.1.4"} <= cells
    chs = wb["CHALLAN"]
    assert any(c.value == "Wire RA-01" for row in chs.iter_rows(max_row=8) for c in row)


# ---------------------------------------------------------------- switching the wire source with metres already given
def test_switching_source_with_metres_given_is_refused_in_plain_words_and_changes_nothing(env):
    client, _ = env
    real_wire(client)
    assert real_alloc(client).status_code == 200
    before = client.get(f"{B}/wire/Electrical").json()
    r = client.put(f"{B}/wire/Electrical/source", json={"source": "stock"})
    assert r.status_code == 409, r.text
    msg = r.json()["detail"]
    assert msg.startswith("RA-01 already has 30000 m of wire given to point lines")
    assert "Wire → points" in msg and "typed wire challans" in msg
    after = client.get(f"{B}/wire/Electrical").json()
    assert after["source"] == before["source"] == "typed"
    assert [x["given"] for x in after["size_status"]] == [x["given"] for x in before["size_status"]]


def test_switching_source_works_once_the_metres_given_are_removed(env):
    client, _ = env
    real_wire(client)
    assert real_alloc(client).status_code == 200
    assert client.put(f"{B}/wire/Electrical/source", json={"source": "stock"}).status_code == 409
    assert alloc(client, []).status_code == 200
    r = client.put(f"{B}/wire/Electrical/source", json={"source": "stock"})
    assert r.status_code == 200, r.text
    assert client.get(f"{B}/wire/Electrical").json()["source"] == "stock"


def test_switching_source_with_nothing_given_is_unchanged(env):
    client, _ = env
    real_wire(client)
    assert client.put(f"{B}/wire/Electrical/source", json={"source": "stock"}).status_code == 200
    assert client.put(f"{B}/wire/Electrical/source", json={"source": "typed"}).status_code == 200
