import sys, json, os, io
from playwright.sync_api import sync_playwright
import urllib.request as u
"""Chromium test (desktop 1366 + phone 390), real backend and the real Thoth BOQ. Electrical's "Only point wiring" Work Order is billed from
challans; its wire comes as challans (metres per colour), the engineer gives the metres to BOQ point lines in the browser, and the points become
the supply of RA 1: 4,020 + 300 + 130 points = Rs 64,28,813 (the invoiced RA 1 of the site's own workbook). The invoice workbook carries the
WIRE CHALLAN sheet. Phone: nothing scrolls sideways, controls are finger-sized.
env: BASE (default http://localhost:8765), ROOT (scratch app root), OUT (screenshots)."""
BASE = os.environ.get("BASE", "http://localhost:8765")
ROOT = os.environ["ROOT"]
OUT = os.environ.get("OUT", "/tmp/billing_wire_shots")
os.makedirs(OUT, exist_ok=True)
P = "/api/billing/thoth-mall"
PW = "Electrical › Only point wiring"
errs = []
def post(path, body, method="POST"):
    r = u.Request(BASE + path, data=json.dumps(body).encode(), headers={"Content-Type": "application/json"}, method=method)
    return json.load(u.urlopen(r))
def get(path):
    return json.load(u.urlopen(BASE + path))
def noscroll(pg, label):
    assert 0 == pg.evaluate("document.documentElement.scrollWidth - document.documentElement.clientWidth"), label + ": page scrolls sideways"
def fingers(pg, sel, label, min_h=40):
    hs = pg.eval_on_selector_all(sel, "els => els.filter(e => e.offsetParent).map(e => Math.round(e.getBoundingClientRect().height))")
    assert hs, label + ": nothing matched " + sel
    for h in hs:
        assert h >= min_h, f"{label}: control only {h}px tall ({sel})"
def reset():
    f = os.path.join(ROOT, "data", "projects", "thoth-mall", "billing.json")
    if os.path.exists(f): os.remove(f)
    g = get(P + "/packages/Electrical")["groups"]
    grp = lambda x: {"sheet": x["sheet"], "from": x["from"], "to": x["to"]}
    pw = next(x for x in g if x["name"].startswith("POINT WIRING"))
    post(P + "/packages/Electrical", {"items": [{"name": "Only point wiring", "groups": [grp(pw)]}, {"name": "Other Electrical", "rest": True}]})
    post(P + "/setup", {"service": PW, "mode": "separate", "supply_pct": 75, "install_pct": 15, "supply_source": "challan"})
    post(P + "/setup", {"service": "Electrical › Other Electrical", "mode": "separate", "supply_pct": 75, "install_pct": 15, "supply_source": "stock"})
sz = lambda s, **k: {"size": s, "colours": {a.upper(): v for a, v in k.items()}}
with sync_playwright() as p:
    b = p.chromium.launch()
    for name, (w, h) in {"desk": (1366, 900), "phone": (390, 844)}.items():
        reset()
        W = f"{P}/wire/{u.quote(PW)}"
        # two of the three real wire challans by API; the third is typed in the browser below
        post(W + "/challans", {"no": "1207", "date": "2026-03-28", "sizes": [sz("1.5", red=300, yellow=600, blue=300, black=900, green=1500), sz("2.5", red=300, yellow=300, blue=300, black=600)]})
        post(W + "/challans", {"no": "615", "date": "2026-07-02", "sizes": [sz("1.5", red=3300, yellow=3000, blue=2700, black=9600, green=13800), sz("2.5", red=1500, yellow=1500, blue=1200, black=3900)]})
        ctx = b.new_context(viewport={"width": w, "height": h}, has_touch=(name == "phone"), is_mobile=(name == "phone"), accept_downloads=True)
        ctx.add_init_script("window.__BL_LIVE_MS = 3600000")
        pg = ctx.new_page()
        pg.on("pageerror", lambda e, n=name: errs.append(("pageerror", n, str(e))))
        pg.on("console", lambda m, n=name: errs.append(("console", n, m.text)) if m.type == "error" else None)
        pg.goto(BASE); pg.wait_for_timeout(800)
        pg.click('.spnav button[data-v="billing"]'); pg.wait_for_selector(".bl-kpis")
        # the link is on the point-wiring Work Order only
        assert pg.locator("[data-wire]").count() == 1, "Wire → points on exactly one tile"
        pg.screenshot(path=f"{OUT}/{name}_1_home.png", full_page=True); noscroll(pg, name + " home")
        pg.click("[data-wire]"); pg.wait_for_selector(".bl-wtab")
        pg.screenshot(path=f"{OUT}/{name}_2_wire_two_challans.png", full_page=True); noscroll(pg, name + " wire screen")
        # ---- type the third real challan (617) in the browser
        pg.click("#bl-wadd"); pg.wait_for_selector("#bl-wcno")
        if name == "phone": fingers(pg, "#bl-modal .bl-wcc .ctl, #bl-modal .bl-chform .ctl", "wire challan form", 44)
        pg.fill("#bl-wcno", "617"); pg.fill("#bl-wcdate", "2026-07-16")
        pg.fill('[data-bsz="0"]', "1.5")
        for k, v in dict(RED=3900, YELLOW=3600, BLUE=2400, BLACK=10500, GREEN=14700).items(): pg.fill(f'[data-bv="0:{k}"]', str(v))
        pg.click("#bl-wcaddsz"); pg.fill('[data-bsz="1"]', "2.5")
        for k, v in dict(RED=1500, YELLOW=1500, BLUE=1800, BLACK=4800).items(): pg.fill(f'[data-bv="1:{k}"]', str(v))
        pg.screenshot(path=f"{OUT}/{name}_3_challan_form.png")
        pg.click("#bl-modal .mact .btn.primary"); pg.wait_for_selector("#bl-modal", state="detached")
        pg.wait_for_function("document.querySelectorAll('.bl-wtab thead [data-wch]').length === 3")
        sums = pg.eval_on_selector_all(".bl-wtab tr.sum td.n b", "els => els.map(e => e.textContent.trim())")
        assert sums == ["20,100", "9,900"], sums
        # ---- give the metres to the three real lines, from the browser
        for code, size, metres in (("1.1.5", "1.5", 20100), ("1.1.3", "2.5", 6000), ("1.1.4", "2.5", 3900)):
            pg.fill("#bl-wfind", code); pg.wait_for_selector(f'#bl-wfound .bl-hit[data-c="{code}"]')
            pg.click(f'#bl-wfound .bl-hit[data-c="{code}"]')
            blk = pg.locator(".bl-wb").last
            blk.locator("select[data-sz]").select_option(size)
            blk.locator("[data-m]").fill(str(metres))
            pg.fill("#bl-wfind", "")
        assert pg.inner_text("#bl-wtot").strip() == "₹ 64,28,813", pg.inner_text("#bl-wtot")
        assert pg.inner_text("#bl-wpts").strip() == "4,450"
        pg.screenshot(path=f"{OUT}/{name}_4_allocated.png", full_page=True); noscroll(pg, name + " allocated")
        if name == "phone":
            fingers(pg, ".bl-wrow .ctl, .bl-wrow .btn", "allocation row", 44)
            fingers(pg, "#bl-wsave, #bl-wadd", "wire buttons", 40)
        pg.click("#bl-wsave"); pg.wait_for_function("document.getElementById('bl-toast') && document.getElementById('bl-toast').textContent.includes('Metres saved')")
        pg.wait_for_selector(".bl-wb [data-m]")
        assert pg.locator(".bl-wb").count() == 3
        st = get(W)
        assert {l["item_code"]: l["points"] for l in st["lines"]} == {"1.1.3": 300, "1.1.4": 130, "1.1.5": 4020}
        assert [x["left"] for x in st["size_status"]] == [0, 0]
        # ---- the RA: supply = the points, as a 'Wire RA-01' challan
        pg.click("#bl-wback"); pg.wait_for_selector(".bl-ch")
        assert "Wire RA-01" in pg.inner_text("#billing")
        pg.screenshot(path=f"{OUT}/{name}_5_challans.png", full_page=True); noscroll(pg, name + " challans")
        ra = post(P + "/ra", {"service": PW, "series": "supply"})["id"]
        bill = get(P + f"/ra/{ra}")
        assert abs(bill["this_total"] - 6428812.5) < 0.01, bill["this_total"]
        # ---- the invoice workbook has the WIRE CHALLAN sheet like the site's own
        data = u.urlopen(BASE + P + f"/invoice/{ra}").read()
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data))
        assert "WIRE CHALLAN" in wb.sheetnames and "CHALLAN" in wb.sheetnames, wb.sheetnames
        ws = wb["WIRE CHALLAN"]
        vals = {c.value for r in ws.iter_rows() for c in r if c.value is not None}
        assert {"WIRE 1.5", "WIRE 2.5", "RED", "YELLOW", "BLUE", 20100, 6000, 3900, 4020, 300, 130, "1.1.5", "1.1.3", "1.1.4"} <= vals, sorted(map(str, vals))[:40]
        ctx.close()
    b.close()
assert not errs, errs[:5]
print("billing wire chromium: all passed")
