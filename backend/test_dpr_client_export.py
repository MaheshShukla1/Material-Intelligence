"""The client-facing DPR (export-dpr?summary=none|without|with): sheet "DPR" (progress status, manpower, work done against the client's BOQ,
hold-ups, plan) + optional "Summary" and "BOQ progress". Real flow end to end: a ProjectBase BOQ uploaded through the real route (client refs,
lettered pipe lines, cable-tray i/ii), progress saved on several days through the real progress routes, manpower typed through the real
route, the xlsx read back cell by cell."""
import datetime as _dt
import io
import json
import re
import sys
import types
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from backend import siteprogress as sp, structure

SLUG = "thoth-mall"
U = f"/api/siteprogress/{SLUG}"
HEAD = ["Internal Ref No", "Client Ref No", "BOQ", "BOQ Unit", "Order Quantity", "Design Quantity", "Rate"]
PIPE = ("Supply, Installation, Testing and Commissioning of the following sizes of flange jointed / screwed / welded, Above Ground M.S. "
        "Pipes of Class C (Heavy class) including accessories and laying on the surface complete as required. a) 150 mm dia")
ROWS = [
    ["1", "1", "Point wiring", None, None, None, None],
    ["2", "1.1.4", "One light point controlled by SP MCB. Cost inculdes rate of MS conduit & Wire with average point length of 30 Rmtr.", "Nos", 200, 200, 1000],
    ["3", "10", "Perforated cable tray (heading, no quantity of its own)", None, None, None, None],
    ["4", "i", "14 G.G.I Perforated - 150 X 50", "Rmt", 100, 100, 500],
    ["5", "ii", "14 G.G.I Perforated - 100 X 50", "Rmt", 60, 60, 400],
    ["6", "39", PIPE, "Rmt", 4200, 4200, 1200],
    ["7", None, "b) 100 mm dia", "Rmt", 1000, 1000, 900],
    ["8", "13.2.1", "25 x 6 mm GI Strip", "Rmt", 800, 800, None],           # no rate in the BOQ
]


class Clock:
    day = _dt.date(2026, 10, 12)


class FakeDate(_dt.date):
    @classmethod
    def today(cls):
        return Clock.day


@pytest.fixture
def env(tmp_path, monkeypatch):
    monkeypatch.setattr(sp, "PROJECTS", tmp_path / "projects")
    monkeypatch.setattr(sp, "RUNS", tmp_path / "runs")
    sp.PROJECTS.mkdir(parents=True)
    sp.RUNS.mkdir(parents=True)
    ns = types.SimpleNamespace(**{k: getattr(_dt, k) for k in dir(_dt) if not k.startswith("__")})
    ns.date = FakeDate
    monkeypatch.setattr(sp, "dt", ns)
    Clock.day = _dt.date(2026, 10, 12)
    for c in (sp._read_forecast_parquet_cached, sp._forecast_pool_cached, sp._full_run_rows_cached, sp._linkage_match_cached):
        c.cache_clear()
    app = FastAPI()
    app.include_router(sp.router)
    c = TestClient(app)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Electrical"
    ws.append(HEAD)
    for r in ROWS:
        ws.append(r)
    wb.save(tmp_path / "boq.xlsx")
    with open(tmp_path / "boq.xlsx", "rb") as fh:
        assert c.post(f"{U}/boq", files={"file": ("boq.xlsx", fh)}).status_code == 200
    d = sp._dir(SLUG)
    s = structure.mall("Thoth Mall", levels=["Ground", "First"], zones_per_level=2)
    (d / "structure.json").write_text(s.to_json())
    codes = {r["description"][:12]: str(r["item_code"]) for r in __import__("pandas").read_parquet(d / "boq.parquet").to_dict("records")}
    (d / "activities.json").write_text(json.dumps({"Electrical": ["Point wiring", "Cable tray", "Piping"]}))
    by_ref = {}
    import pandas as pd
    for r in pd.read_parquet(d / "boq.parquet").to_dict("records"):
        by_ref[(r["description"][:14])] = str(r["item_code"])
    pq = pd.read_parquet(d / "boq.parquet")
    code = lambda txt: str(pq[pq.description.str.contains(txt, regex=False)].iloc[0].item_code)
    c.codes = {"light": code("One light point"), "tray1": code("150 X 50"), "tray2": code("100 X 50"), "pipe_a": code("150 mm dia"),
               "pipe_b": code("b) 100 mm dia"), "strip": code("GI Strip")}
    (d / "mapping.json").write_text(json.dumps({"Electrical": {"Point wiring": [c.codes["light"]], "Cable tray": [c.codes["tray1"], c.codes["tray2"]],
                                                              "Piping": [c.codes["pipe_a"], c.codes["pipe_b"], c.codes["strip"]]}}))
    assert c.post(f"{U}/settings", json={"default_install_pct": 15}).status_code == 200
    c.zones = [r["id"] for r in s.rooms()]
    c.zone_names = {r["id"]: r["name"] for r in s.rooms()}
    c.dir = d
    return c


def day(n):
    Clock.day = _dt.date(2026, 10, n)


def slider(c, key, frac, room=None):
    body = {"service": "Electrical", "item_code": c.codes[key], "frac": frac}
    if room:
        body["room"] = room
    r = c.post(f"{U}/progress/item", json=body)
    assert r.status_code == 200, r.text


def export(c, start, mode, **kw):
    q = {"start": start, "summary": mode, **kw}
    r = c.get(f"{U}/export-dpr", params=q)
    assert r.status_code == 200, r.text
    return openpyxl.load_workbook(io.BytesIO(r.content)), r


def rows_of(ws):
    return [[cell for cell in row] for row in ws.iter_rows(values_only=True)]


def find(ws, col, text):
    """All rows (as value lists) whose column `col` (1-based) equals text."""
    return [row for row in rows_of(ws) if len(row) >= col and row[col - 1] == text]


def work_lines(ws):
    """The DPR sheet's section-3 lines as dicts, keyed by their BOQ Ref."""
    out, on = {}, False
    for row in rows_of(ws):
        a = row[0]
        if isinstance(a, str) and a.startswith("3."):
            on = True
            continue
        if isinstance(a, str) and a.startswith("4."):
            break
        if on and isinstance(a, int):
            out[row[1]] = dict(sr=a, ref=row[1], desc=row[2], unit=row[3], boq=row[4], prev=row[5], today=row[6], cum=row[7], bal=row[8],
                               pct=row[9], loc=row[10], rem=row[11])
    return out


# ------------------------------------------------------------------ the reported problem: TODAY shows today's work only
def test_the_dpr_shows_only_that_days_work_not_the_cumulative_quantity(env):
    c = env
    slider(c, "tray1", 0.7)                         # Monday: 70 of 100 Rmt
    day(13)
    slider(c, "tray1", 1.0)                         # Tuesday: the remaining 30
    ws = export(c, "2026-10-13", "none")[0]["DPR"]
    ln = work_lines(ws)["10 (i)"]
    assert (ln["boq"], ln["prev"], ln["today"], ln["cum"], ln["bal"]) == (100, 70, 30, 100, 0)
    assert ln["pct"] == pytest.approx(1.0)
    mon = work_lines(export(c, "2026-10-12", "none")[0]["DPR"])["10 (i)"]
    assert (mon["prev"], mon["today"], mon["cum"], mon["bal"]) == (0, 70, 70, 30)      # Monday re-exported afterwards: its own figures


def test_only_lines_worked_that_day_are_listed_and_untouched_boq_lines_are_not(env):
    c = env
    slider(c, "tray1", 0.5)
    day(13)
    slider(c, "light", 0.25)
    assert list(work_lines(export(c, "2026-10-13", "none")[0]["DPR"])) == ["1.1.4"]
    assert list(work_lines(export(c, "2026-10-12", "none")[0]["DPR"])) == ["10 (i)"]


def test_a_correction_the_same_day_does_not_double_count_and_a_no_change_save_is_not_work(env):
    c = env
    slider(c, "tray1", 0.3)
    slider(c, "tray1", 0.6)
    slider(c, "tray2", 0.0)                         # touched but nothing moved
    lines = work_lines(export(c, "2026-10-12", "none")[0]["DPR"])
    assert lines["10 (i)"]["today"] == 60 and "10 (ii)" not in lines


def test_overall_slider_lines_name_the_levels_and_zone_work_names_the_zones(env):
    c = env
    slider(c, "tray1", 0.4)
    slider(c, "tray2", 1.0, c.zones[1])
    lines = work_lines(export(c, "2026-10-12", "none")[0]["DPR"])
    assert lines["10 (i)"]["loc"] == "First, Ground (all zones)"
    assert c.zone_names[c.zones[1]] in lines["10 (ii)"]["loc"]


# ------------------------------------------------------------------ the client's own BOQ references
def test_boq_refs_are_the_clients_numbers_with_the_letter_of_lettered_lines(env):
    c = env
    for k in ("light", "tray1", "tray2", "pipe_a", "pipe_b", "strip"):
        slider(c, k, 0.5)
    lines = work_lines(export(c, "2026-10-12", "none")[0]["DPR"])
    assert set(lines) == {"1.1.4", "10 (i)", "10 (ii)", "39 (a)", "39 (b)", "13.2.1"}
    assert lines["39 (b)"]["desc"] == "b) 100 mm dia"                     # blank number in the client's sheet: parent 39 + its letter
    assert lines["39 (a)"]["desc"] == "a) 150 mm dia"                    # the parent's text is split off the first lettered line


def test_the_parent_text_sits_once_above_its_lettered_lines(env):
    c = env
    slider(c, "tray1", 0.5)
    slider(c, "tray2", 0.5)
    ws = export(c, "2026-10-12", "none")[0]["DPR"]
    parents = [r[0] for r in rows_of(ws) if isinstance(r[0], str) and r[0].startswith("Perforated cable tray")]
    assert len(parents) == 1


# ------------------------------------------------------------------ sheets, files and the amount split
def test_the_three_choices_give_the_three_files_and_sheet_sets(env):
    c = env
    slider(c, "tray1", 0.5)
    for mode, names, suffix in (("none", ["DPR"], ""), ("without", ["DPR", "Summary", "BOQ progress"], "_summary-without-amount"),
                                ("with", ["DPR", "Summary", "BOQ progress"], "_summary-with-amount")):
        wb, r = export(c, "2026-10-12", mode)
        assert wb.sheetnames == names
        assert f'filename="DPR_thoth-mall_2026-10-12{suffix}.xlsx"' in r.headers["content-disposition"]


def test_without_amount_file_has_no_amount_wording_or_rupee_anywhere(env):
    c = env
    for k in ("light", "tray1", "pipe_a"):
        slider(c, k, 0.5)
    wb, r = export(c, "2026-10-12", "without")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                assert "₹" not in str(cell.value) and "amount" not in str(cell.value).lower() and "rupee" not in str(cell.value).lower(), (ws.title, cell.coordinate, cell.value)
                assert "₹" not in cell.number_format, (ws.title, cell.coordinate)
        assert "amount" not in (ws.oddFooter.center.text or "").lower()
    blob = r.content
    import zipfile
    z = zipfile.ZipFile(io.BytesIO(blob))
    for n in z.namelist():
        if n.endswith(".xml"):
            t = z.read(n).decode("utf8", "ignore").lower()
            assert "amount" not in t and "₹" not in t and "rupee" not in t, n
    ws = wb["Summary"]
    assert ws["A1"].value.startswith("Site progress summary — Thoth Mall")
    assert wb["BOQ progress"]["A1"].value.startswith("BOQ progress to date — Thoth Mall")
    assert "WITHOUT" not in ws["A1"].value.upper()


def test_with_amount_values_tie_everywhere_and_to_the_apps_own_summary(env):
    """Done ₹ is BOQ rate x installation % x quantity done: the same figure the existing Summary export gives."""
    c = env
    slider(c, "light", 0.5)              # 100 Nos x 1000 x 15% = 15,000
    slider(c, "tray1", 0.7)              # 70 x 500 x 15% = 5,250
    slider(c, "pipe_a", 0.1)             # 420 x 1200 x 15% = 75,600
    slider(c, "strip", 0.5)              # no BOQ rate: no ₹
    wb, _r = export(c, "2026-10-12", "with")
    summ, boqp = wb["Summary"], wb["BOQ progress"]
    row = find(summ, 1, "Electrical")[0]
    pct, done, rem, total = row[1], row[2], row[3], row[4]
    assert done == pytest.approx(15000 + 5250 + 75600)
    assert done + rem == pytest.approx(total)
    # the BOQ progress sheet adds up to the same totals
    items = [r for r in rows_of(boqp)[2:] if isinstance(r[0], int)]
    assert sum(r[14] for r in items if isinstance(r[14], (int, float))) == pytest.approx(done)
    assert sum(r[13] for r in items if isinstance(r[13], (int, float))) == pytest.approx(total)
    tot = [r for r in rows_of(boqp) if r[0] == "Total"][0]
    assert tot[13] == pytest.approx(total) and tot[14] == pytest.approx(done)
    assert pct == pytest.approx(done / total)
    # an item without a BOQ rate: "—" in every amount column, never 0
    strip = [r for r in items if r[2] == "13.2.1"][0]
    assert strip[12:16] == ["—", "—", "—", "—"]
    # and it is the existing Summary's number
    legacy = openpyxl.load_workbook(io.BytesIO(c.get(f"{U}/export-dpr", params={"start": "2026-10-12"}).content))["Summary"]
    assert [r for r in rows_of(legacy) if r[0] == "Electrical"][0][2] == pytest.approx(done, abs=1)
    # the material-waste tile has no source in this export: never a made-up 0
    assert summ["G5"].value == "—"


def test_summary_zone_counts_come_from_the_boq_progress_sheet(env):
    c = env
    slider(c, "tray1", 1.0, c.zones[0])
    slider(c, "tray1", 0.5, c.zones[1])
    wb, _r = export(c, "2026-10-12", "without")
    boqp = wb["BOQ progress"]
    items = [r for r in rows_of(boqp)[2:] if isinstance(r[0], int)]
    row = find(wb["Summary"], 1, "Electrical")[0]
    assert row[4:7] == [sum(r[9] for r in items), sum(r[10] for r in items), sum(r[11] for r in items)]
    assert row[4] >= 1 and row[5] >= 1


def test_progress_status_up_to_previous_day_today_and_to_date(env):
    c = env
    slider(c, "light", 0.5)              # Monday: 15,000
    day(13)
    slider(c, "tray1", 0.7)              # Tuesday: 5,250
    wb, _r = export(c, "2026-10-13", "without")
    total = 200 * 1000 * .15 + 100 * 500 * .15 + 60 * 400 * .15 + 4200 * 1200 * .15 + 1000 * 900 * .15
    row = find(wb["DPR"], 2, "Electrical") or [r for r in rows_of(wb["DPR"]) if r[1] == "Electrical"]
    prog = [r for r in rows_of(wb["DPR"]) if r[1] == "Electrical"][0]
    assert prog[4] == pytest.approx(15000 / total) and prog[5] == pytest.approx(5250 / total) and prog[6] == pytest.approx(20250 / total)
    assert [r for r in rows_of(wb["DPR"]) if r[1] == "Overall project"][0][6] == pytest.approx(20250 / total)


# ------------------------------------------------------------------ manpower, hold-ups, plan, parties
def test_manpower_hold_ups_plan_and_parties_land_on_the_dpr_with_man_days(env):
    c = env
    slider(c, "tray1", 0.5)
    day(13)
    slider(c, "tray1", 0.6)
    c.post(f"{U}/dpr/day", json={"date": "2026-10-12", "manpower": {"Electrical": {"supervisor": 1, "skilled": 8, "helper": 12}}})
    r = c.post(f"{U}/dpr/day", json={"date": "2026-10-13", "manpower": {"Electrical": {"supervisor": 1, "skilled": 9, "helper": 14}},
                                     "holdups": "Slab opening not handed over.", "plan": "Tray in B2.",
                                     "client": "Thoth Developers", "contractor": "Mahesh Engg", "prepared_by": "S. Engineer"})
    assert r.status_code == 200 and r.json()["weekday"] == "Tuesday"
    ws = export(c, "2026-10-13", "none")[0]["DPR"]
    mp = [r for r in rows_of(ws) if r[1] == "Electrical"][1]
    assert mp[4:10] == [1, 9, 14, 24, 21, 45]                                   # today 24, up to previous day 21, to date 45
    flat = [str(v) for row in rows_of(ws) for v in row if v is not None]
    assert "Slab opening not handed over." in flat and "Tray in B2." in flat
    assert ws["C3"].value == "Thoth Developers" and ws["G3"].value == "Mahesh Engg" and ws["K3"].value == "S. Engineer"
    assert ws["K2"].value == 2 and "13-Oct-2026 (Tuesday)" in ws["G2"].value      # second day with updates


def test_manpower_not_entered_is_a_dash_never_a_zero(env):
    c = env
    slider(c, "tray1", 0.5)
    ws = export(c, "2026-10-12", "none")[0]["DPR"]
    mp = [r for r in rows_of(ws) if r[1] == "Electrical"][1]
    assert mp[4:10] == ["—"] * 6


def test_save_day_validation_and_the_blank_means_not_entered(env):
    c = env
    bad = lambda body: c.post(f"{U}/dpr/day", json=body).status_code
    assert bad({"manpower": {"Electrical": {"skilled": "x"}}}) == 400
    assert bad({"manpower": {"Electrical": {"skilled": -1}}}) == 400
    assert bad({"manpower": {"Electrical": {"skilled": 2.5}}}) == 400
    assert bad({"manpower": {"Plumbing": {"skilled": 2}}}) == 400            # not a service of this project
    assert bad({"date": "13/10/2026"}) == 400
    r = c.post(f"{U}/dpr/day", json={"date": "2026-10-12", "manpower": {"Electrical": {"skilled": 5, "helper": ""}}})
    assert r.json()["manpower"] == {"Electrical": {"supervisor": 0, "skilled": 5, "helper": 0}}
    r = c.post(f"{U}/dpr/day", json={"date": "2026-10-12", "manpower": {"Electrical": {"supervisor": "", "skilled": "", "helper": None}}})
    assert r.json()["manpower"] == {}
    assert c.get(f"{U}/dpr/day", params={"date": "2026-10-12"}).json()["holdups"] == ""
    assert not (c.dir / "dpr_days.json").read_text().strip("{} \n")           # an emptied day leaves nothing behind


def test_a_range_gives_a_period_report_not_a_wrong_single_day(env):
    c = env
    slider(c, "tray1", 0.5)
    day(13)
    slider(c, "tray1", 0.8)
    ws = export(c, "2026-10-12", "none", end="2026-10-13")[0]["DPR"]
    ln = work_lines(ws)["10 (i)"]
    assert (ln["prev"], ln["today"], ln["cum"]) == (0, 80, 80)
    assert ws["E2"].value == "Period" and "12-Oct-2026 to 13-Oct-2026" in ws["G2"].value and ws["K2"].value == "—"


def test_area_filter_limits_the_lines(env):
    c = env
    slider(c, "tray1", 1.0, c.zones[0])
    slider(c, "tray2", 1.0, c.zones[3])
    ws = export(c, "2026-10-12", "none", rooms=[c.zones[0]])[0]["DPR"]
    assert list(work_lines(ws)) == ["10 (i)"]


def test_without_the_summary_parameter_the_old_export_is_unchanged(env):
    c = env
    slider(c, "tray1", 0.5)
    r = c.get(f"{U}/export-dpr", params={"start": "2026-10-12"})
    assert openpyxl.load_workbook(io.BytesIO(r.content)).sheetnames == ["Summary", "Item detail", "Daily updates"]
    assert 'filename="DPR_thoth-mall_2026-10-12.xlsx"' in r.headers["content-disposition"]
    assert c.get(f"{U}/export-dpr", params={"start": "2026-10-12", "summary": "bogus"}).status_code == 400


def test_a_day_with_no_updates_is_an_honest_empty_report(env):
    c = env
    ws = export(c, "2026-10-12", "none")[0]["DPR"]
    assert work_lines(ws) == {}
    assert any("No BOQ line was updated" in str(v) for row in rows_of(ws) for v in row if v)


def test_unknown_service_is_refused_and_a_service_filter_limits_the_sheets(env):
    c = env
    slider(c, "tray1", 0.5)
    assert c.get(f"{U}/export-dpr", params={"start": "2026-10-12", "summary": "none", "service": "Nope"}).status_code == 400
    wb, r = export(c, "2026-10-12", "without", service="Electrical")
    assert "DPR_thoth-mall_2026-10-12_electrical_summary-without-amount.xlsx" in r.headers["content-disposition"]
    assert wb["DPR"]["B4"].value == "Scope: Electrical"


def test_a_hotel_project_names_floors_and_rooms(env):
    """The same export on a hotel-shaped project (Floor > Room), like Hyatt."""
    c = env
    s = structure.hotel("Hyatt Hotel", floors=["13TH", "14TH"], room_labels=["5", "14"])
    (c.dir / "structure.json").write_text(s.to_json())
    rooms = [r["id"] for r in s.rooms()]
    slider(c, "tray1", 0.5)
    slider(c, "tray2", 1.0, rooms[0])
    ws = export(c, "2026-10-12", "none")[0]["DPR"]
    lines = work_lines(ws)
    assert lines["10 (i)"]["loc"] == "13TH, 14TH (all rooms)"
    assert lines["10 (ii)"]["loc"].startswith("13TH · ")
    assert ws["C2"].value == "Hyatt Hotel"
