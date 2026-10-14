"""DPR quantities by unit: a counted unit (Nos ...) prints WHOLE numbers, and "work done today" is never negative.

The reported DPR (Thoth Mall, 14-Oct-2026, jet fans in Nos) showed 5.25 / 6.75 / 4.286 / 3.999 Nos and -1 / -0.75 Nos for "TODAY". Reproduced
here the way it happened: older entries were a PERCENTAGE of the item (75 % of 7 fans = 5.25), today the engineer typed the real installed
count in the item's quantity box (6), some of them lower than the old percentage said. Through the real routes, the xlsx read back cell by cell."""
import datetime as _dt
import io
import json
import sys
import types
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl
import pandas as pd
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from backend import dpr, siteprogress as sp, structure

SLUG = "thoth-mall"
U = f"/api/siteprogress/{SLUG}"
HEAD = ["Internal Ref No", "Client Ref No", "BOQ", "BOQ Unit", "Order Quantity", "Design Quantity", "Rate"]
FAN = "2500 CFM Normal Mode/ 4300 CFM Fire Mode"
ROWS = [
    ["1", "8.1.3", "Ceiling Mounted Centrifugal Jet-vent fan", None, None, None, None],
    ["2", None, f"c) {FAN} FANC", "Nos", 7, 7, 1000],
    ["3", None, f"d) {FAN} FAND", "Nos", 8, 8, 1000],
    ["4", None, f"e) {FAN} FANE", "Nos", 9, 9, 1000],
    ["5", None, f"f) {FAN} FANF", "Nos", 6, 6, 1000],
    ["6", None, f"j) {FAN} FANJ", "Nos", 10, 10, 1000],
    ["7", "10", "Perforated cable tray heading", None, None, None, None],
    ["8", "i", "14 G.G.I Perforated - 150 X 50", "Rmt", 100, 100, 500],
]


class Clock:
    day = _dt.date(2026, 10, 12)


class FakeDate(_dt.date):
    @classmethod
    def today(cls):
        return Clock.day


def day(n):
    Clock.day = _dt.date(2026, 10, n)


@pytest.fixture
def env(tmp_path, monkeypatch):
    monkeypatch.setattr(sp, "PROJECTS", tmp_path / "projects")
    monkeypatch.setattr(sp, "RUNS", tmp_path / "runs")
    sp.PROJECTS.mkdir(parents=True)
    sp.RUNS.mkdir(parents=True)
    ns = types.SimpleNamespace(**{k: getattr(_dt, k) for k in dir(_dt) if not k.startswith("__")})
    ns.date = FakeDate
    monkeypatch.setattr(sp, "dt", ns)
    Clock.day = _dt.date(2026, 10, 12)
    for c in (sp._read_forecast_parquet_cached, sp._forecast_pool_cached, sp._full_run_rows_cached, sp._linkage_match_cached):
        c.cache_clear()
    app = FastAPI()
    app.include_router(sp.router)
    c = TestClient(app)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "HVAC"
    ws.append(HEAD)
    for r in ROWS:
        ws.append(r)
    wb.save(tmp_path / "boq.xlsx")
    with open(tmp_path / "boq.xlsx", "rb") as fh:
        assert c.post(f"{U}/boq", files={"file": ("boq.xlsx", fh)}).status_code == 200
    d = sp._dir(SLUG)
    s = structure.mall("Thoth Mall", levels=["B1", "B2"], zones_per_level=2)
    (d / "structure.json").write_text(s.to_json())
    pq = pd.read_parquet(d / "boq.parquet")
    code = lambda txt: str(pq[pq.description.str.contains(txt, regex=False)].iloc[0].item_code)
    c.codes = {k: code(f"FAN{k.upper()}") for k in "cdefj"}
    c.codes["tray"] = code("150 X 50")
    (d / "activities.json").write_text(json.dumps({"HVAC": ["Jet fan", "Tray"]}))
    (d / "mapping.json").write_text(json.dumps({"HVAC": {"Jet fan": [c.codes[k] for k in "cdefj"], "Tray": [c.codes["tray"]]}}))
    assert c.post(f"{U}/settings", json={"default_install_pct": 15}).status_code == 200
    c.dir = d
    return c


def percent(c, key, frac):                      # the old way: a percentage of the whole item
    r = c.post(f"{U}/progress/item", json={"service": "HVAC", "item_code": c.codes[key], "frac": frac})
    assert r.status_code == 200, r.text


def typed(c, key, qty):                         # the engineer types the real installed quantity in the item's own box
    r = c.post(f"{U}/progress/item-qty", json={"service": "HVAC", "item_code": c.codes[key], "qty": qty})
    assert r.status_code == 200, r.text


def export(c, d, mode="without"):
    r = c.get(f"{U}/export-dpr", params={"start": f"2026-10-{d:02d}", "summary": mode})
    assert r.status_code == 200, r.text
    return openpyxl.load_workbook(io.BytesIO(r.content))


def dpr_lines(wb):
    """{marker: figures} of the work-done table of the DPR sheet (columns D..L), found by a word of the item's own description."""
    ws = wb["DPR"]
    out = {}
    for row in ws.iter_rows(min_row=1):
        txt = str(row[2].value or "")
        for k in ("FANC", "FAND", "FANE", "FANF", "FANJ", "150 X 50"):
            if k in txt:
                out[k] = {"unit": row[3].value, "boq": row[4].value, "prev": row[5].value, "today": row[6].value, "cum": row[7].value,
                          "bal": row[8].value, "pct": row[9].value, "remarks": row[11].value, "fmt": row[6].number_format}
    return out


def reported_day(c):
    """Old percentages on the 12th, then the engineer's typed counts on the 14th (the reported DPR)."""
    day(12)
    for k, f in (("c", 0.75), ("d", 1.0), ("e", 0.75), ("f", 5 / 7), ("j", 0.7)):
        percent(c, k, f)
    day(14)
    for k, q in (("c", 6), ("d", 7), ("e", 6), ("f", 5), ("j", 10)):
        typed(c, k, q)


def test_the_reported_dpr_prints_whole_nos_and_never_a_negative_today(env):
    c = env
    reported_day(c)
    got = dpr_lines(export(c, 14, "none"))
    assert set(got) >= {"FANC", "FAND", "FANE", "FANF", "FANJ"}, got
    # c: 75 % of 7 = 5.25 earlier, 6 typed today -> 5 + 1 = 6
    assert (got["FANC"]["prev"], got["FANC"]["today"], got["FANC"]["cum"], got["FANC"]["bal"]) == (5, 1, 6, 1)
    # f: 5/7 of 6 = 4.286 earlier, 5 typed -> 4 + 1 = 5
    assert (got["FANF"]["prev"], got["FANF"]["today"], got["FANF"]["cum"]) == (4, 1, 5)
    # j: 70 % of 10 = 7, 10 typed -> 7 + 3 = 10 (full)
    assert (got["FANJ"]["prev"], got["FANJ"]["today"], got["FANJ"]["cum"], got["FANJ"]["bal"]) == (7, 3, 10, 0)
    # d: was 8 of 8, recounted to 7 -> NOT -1 of work: today 0, the row still adds up, and the correction is written
    assert (got["FAND"]["prev"], got["FAND"]["today"], got["FAND"]["cum"]) == (7, 0, 7)
    assert got["FAND"]["remarks"] == "Quantity corrected down by 1 Nos"
    # e: 75 % of 9 = 6.75 earlier, 6 typed -> corrected down by 1 (7 -> 6), not -0.75
    assert (got["FANE"]["prev"], got["FANE"]["today"], got["FANE"]["cum"]) == (6, 0, 6)
    assert got["FANE"]["remarks"] == "Quantity corrected down by 1 Nos"
    for k in ("FANC", "FAND", "FANE", "FANF", "FANJ"):
        g = got[k]
        assert g["unit"] == "Nos"
        for f in ("prev", "today", "cum", "bal"):
            assert float(g[f]) == int(g[f]) and g[f] >= 0, (k, f, g[f])           # whole, never negative
        assert g["prev"] + g["today"] == g["cum"], k                                # every row adds up
        assert g["fmt"] == "#,##0", g["fmt"]                                        # printed as a whole number
    assert got["FANC"]["remarks"] in (None, dpr.DASH)


def test_the_today_percent_counts_only_real_work_not_the_two_corrections(env):
    """Work done today = c +1, f +1, j +3 = 5 fans (d and e were recounted DOWN: corrections, not work). Contract = 40 fans + 100 Rmt tray at
    15 % installation of the rate: 40 x 150 + 100 x 75 = 13,500. Today = 5 x 150 / 13,500, never reduced by the corrections."""
    c = env
    reported_day(c)
    wb = export(c, 14, "without")
    rows = {str(r[1].value): r for r in wb["DPR"].iter_rows(min_row=6, max_row=9)}
    assert rows["HVAC"][5].value == pytest.approx(5 * 150 / 13500)
    assert rows["Overall project"][5].value == pytest.approx(5 * 150 / 13500)
    # and the same for the Summary sheet's tiles
    s = wb["Summary"]
    vals = [c_.value for r in s.iter_rows(min_row=1, max_row=40) for c_ in r
            if c_.number_format == "+0.0%" and isinstance(c_.value, (int, float))]
    assert vals and all(v > -1e-9 for v in vals)


def test_the_boq_progress_sheet_is_whole_too_and_rupees_are_untouched(env):
    c = env
    reported_day(c)
    ws = export(c, 14, "with")["BOQ progress"]
    rows = {}
    for r in ws.iter_rows(min_row=3):
        t = str(r[3].value or "")
        for k in ("FANC", "FAND", "FANE", "FANF", "FANJ"):
            if k in t:
                rows[k] = r
    assert set(rows) == {"FANC", "FAND", "FANE", "FANF", "FANJ"}
    for k, (cum, bal) in {"FANC": (6, 1), "FAND": (7, 1), "FANE": (6, 3), "FANF": (5, 1), "FANJ": (10, 0)}.items():
        r = rows[k]
        assert (r[6].value, r[7].value) == (cum, bal), (k, r[6].value, r[7].value)
        assert r[6].number_format == "#,##0" and r[7].number_format == "#,##0"
    # rupees: the typed counts x rate x installation % (15 %) -- same as before, not computed from the rounded figures
    done = sum(r[14].value for r in rows.values())
    assert done == pytest.approx((6 + 7 + 6 + 5 + 10) * 1000 * 0.15)


def test_a_length_item_keeps_its_decimals_and_a_correction_is_shown_not_negative(env):
    c = env
    day(12)
    percent(c, "tray", 0.8)
    day(13)
    r = c.post(f"{U}/progress/item", json={"service": "HVAC", "item_code": c.codes["tray"], "frac": 0.6})
    assert r.status_code == 200
    got = dpr_lines(export(c, 13, "none"))["150 X 50"]
    assert (got["prev"], got["today"], got["cum"]) == (60, 0, 60)
    assert got["remarks"] == "Quantity corrected down by 20 Rmt"
    day(14)
    percent(c, "tray", 0.625)                    # a real fraction of a length stays a fraction
    got = dpr_lines(export(c, 14, "none"))["150 X 50"]
    assert (got["prev"], got["today"], got["cum"]) == (60, 2.5, 62.5)
    assert got["fmt"] == "#,##0.00"


def test_an_unknown_day_quantity_is_still_a_dash_not_a_guess(env):
    c = env
    d = c.dir
    (d / "dpr_log.json").write_text(json.dumps([{"date": "2026-10-14", "service": "HVAC", "floor": "B1", "activity": "Jet fan",
                                                  "room": "Zone 1", "item": f"c) {FAN} FANC", "qty": 2.0, "unit": "Nos",
                                                  "item_code": c.codes["c"]}]))
    day(14)
    got = dpr_lines(export(c, 14, "none"))
    assert got["FANC"]["today"] == dpr.DASH and got["FANC"]["prev"] == dpr.DASH


# ------------------------------------------------------------------ the pure helper
@pytest.mark.parametrize("unit,prev,today,cum,want", [
    ("Nos", 5.25, 0.75, 6.0, (5.0, 1.0, 6.0, None)),
    ("Nos", 6.75, -0.75, 6.0, (6.0, 0.0, 6.0, "Quantity corrected down by 1 Nos")),
    ("Nos", 8.0, -1.0, 7.0, (7.0, 0.0, 7.0, "Quantity corrected down by 1 Nos")),
    ("Nos", 3.999, 2.001, 6.0, (4.0, 2.0, 6.0, None)),
    ("Nos", 4.4, 0.2, 4.6, (4.0, 1.0, 5.0, None)),                  # the screen rounds 4.6 up to 5; so does the report
    ("Set", 2.5, 1.0, 3.5, (3.0, 1.0, 4.0, None)),                  # half rounds up, like the screen
    ("Rmt", 80.0, -20.0, 60.0, (60.0, 0.0, 60.0, "Quantity corrected down by 20 Rmt")),
    ("Rmt", 60.0, 2.5, 62.5, (60.0, 2.5, 62.5, None)),
    ("Nos", None, None, 6.0, (None, None, 6.0, None)),               # unknown stays unknown
    ("Nos", None, 2.0, None, (None, 2.0, None, None)),
])
def test_present_qty(unit, prev, today, cum, want):
    r = dpr.present_qty(unit, prev, today, cum)
    assert (r["prev"], r["today"], r["cum"], r["remark"]) == want


def test_count_units_are_the_same_set_the_screen_uses():
    js = (Path(__file__).resolve().parent.parent / "frontend" / "siteprogress.js").read_text()
    import re
    m = re.search(r"const COUNT_UNITS = new Set\(\[(.*?)\]\)", js, re.S)
    screen = set(re.findall(r'"([^"]+)"', m.group(1)))
    assert screen == set(dpr.COUNT_UNITS)
