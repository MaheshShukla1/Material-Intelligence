"""Material Forecast export: scoped by service/category/size/status, same
filter logic /api/forecast already uses (shared via _scoped_forecast), a
real polished .xlsx instead of an unfiltered raw CSV dump.

Run from the repo root (same level as backend/):
    pytest tests/test_forecast_export.py -v
"""
import json
import io

import openpyxl
import pandas as pd
import pytest
from fastapi.testclient import TestClient

from backend import api


ROWS = [
    # service,      material,                    subcategory, unit, stock, total_consumed, status,       confidence, order_by,     exhaust_date
    ("Electrical", "1.5 SQMM FLEXIBLE WIRE GREEN", "Wire", "MTR", 0,    13500, "STOCKED_OUT", "HIGH",   "2026-09-01", "2026-09-01"),
    ("Electrical", "2.5 SQMM FLEXIBLE WIRE BLUE",  "Wire", "MTR", 200,  7560,  "RED",         "HIGH",   "2026-09-16", "2026-09-20"),
    ("Electrical", "4.0 SQMM FLEXIBLE WIRE BLACK", "Wire", "MTR", 5000, 5400,  "GREEN",       "HIGH",   None,         None),
    ("Electrical", "300X50MM HDGI LADDER TRAY",    "Cable tray", "MTR", 50, 1200, "AMBER",   "MEDIUM", "2026-09-22", "2026-09-25"),
    ("Plumbing",   "50MM UPVC PIPE",                "Pipe", "MTR", 300,  900,   "OVERSTOCK",   "HIGH",   None,         None),
    ("Plumbing",   "75MM PVC PIPE",                 "Pipe", "MTR", 0,    2000,  "STOCKED_OUT", "LOW",    "2026-09-01", "2026-09-01"),
]
COLS = ["service", "material", "subcategory", "unit", "stock", "total_consumed",
       "status", "confidence", "order_by", "exhaust_date"]


@pytest.fixture
def run(tmp_path, monkeypatch):
    monkeypatch.setattr(api, "RUNS", tmp_path)
    run_id = "run1"
    d = tmp_path / run_id
    d.mkdir()
    df = pd.DataFrame(ROWS, columns=COLS)
    df["order_by"] = pd.to_datetime(df["order_by"])
    df["exhaust_date"] = pd.to_datetime(df["exhaust_date"])
    df["tool_type"] = None   # already present -> add_tooltype() never called
    df.to_parquet(d / "forecast.parquet")
    (d / "meta.json").write_text(json.dumps({"project": "Fixture"}))
    client = TestClient(api.app)
    return client, run_id


def _wb(resp):
    return openpyxl.load_workbook(io.BytesIO(resp.content))


class TestBackwardCompatibility:
    def test_no_filters_returns_every_material(self, run):
        client, run_id = run
        r = client.get(f"/api/export/{run_id}")
        assert r.status_code == 200
        wb = _wb(r)
        ws = wb["Material Forecast"]
        materials = [ws.cell(row, 1).value for row in range(5, ws.max_row + 1)]
        materials = [m for m in materials if m]
        assert len(materials) == len(ROWS)


class TestServiceAndCategoryScoping:
    def test_service_plus_category_scopes_correctly(self, run):
        client, run_id = run
        r = client.get(f"/api/export/{run_id}",
                       params={"service": "Electrical", "subcategory": "Wire"})
        wb = _wb(r)
        ws = wb["Material Forecast"]
        materials = [ws.cell(row, 1).value for row in range(5, ws.max_row + 1) if ws.cell(row, 1).value]
        assert len(materials) == 3   # the 3 real Wire rows
        assert all("WIRE" in m for m in materials)
        assert "300X50MM HDGI LADDER TRAY" not in materials, \
            "FAIL: Cable tray leaked into a Wire-scoped export"
        assert not any("PVC" in m or "UPVC" in m for m in materials), \
            "FAIL: Plumbing leaked into an Electrical-scoped export"

    def test_size_further_narrows_within_category(self, run):
        client, run_id = run
        r = client.get(f"/api/export/{run_id}",
                       params={"service": "Electrical", "subcategory": "Wire", "size": "2.5SQMM"})
        wb = _wb(r)
        ws = wb["Material Forecast"]
        materials = [ws.cell(row, 1).value for row in range(5, ws.max_row + 1) if ws.cell(row, 1).value]
        assert materials == ["2.5 SQMM FLEXIBLE WIRE BLUE"]

    def test_status_bucket_scoping_matches_act_today_definition(self, run):
        """'Act today' on screen = STOCKED_OUT + RED -- the export must use
        the exact same two-status set, not a looser/different one."""
        client, run_id = run
        r = client.get(f"/api/export/{run_id}", params={"status": "STOCKED_OUT,RED"})
        wb = _wb(r)
        ws = wb["Material Forecast"]
        rows = [(ws.cell(row, 1).value, ws.cell(row, 4).value)
               for row in range(5, ws.max_row + 1) if ws.cell(row, 1).value]
        assert len(rows) == 3   # 2 STOCKED_OUT + 1 RED
        assert all(status in ("Already out", "Order now") for _, status in rows)


class TestColumnsAndValues:
    def test_total_received_equals_consumed_plus_stock(self, run):
        """Matches app.js's own totalReceived(): total_consumed + stock --
        the same figure the on-screen 'Total rcvd' column already shows."""
        client, run_id = run
        r = client.get(f"/api/export/{run_id}",
                       params={"service": "Electrical", "subcategory": "Wire", "size": "2.5SQMM"})
        wb = _wb(r)
        ws = wb["Material Forecast"]
        row = [row for row in range(5, ws.max_row + 1) if ws.cell(row, 1).value][0]
        assert ws.cell(row, 5).value == 200     # Stock
        assert ws.cell(row, 6).value == 7760    # Total Received = 7560 + 200
        assert ws.cell(row, 7).value == 7560    # Consumed to date

    def test_dates_are_absolute_not_relative_text(self, run):
        """The real reported design requirement: an exported Order By /
        Runs Out must be a real calendar date, never "in N days" text that
        goes stale after today."""
        client, run_id = run
        r = client.get(f"/api/export/{run_id}",
                       params={"service": "Electrical", "subcategory": "Wire", "size": "2.5SQMM"})
        wb = _wb(r)
        ws = wb["Material Forecast"]
        row = [row for row in range(5, ws.max_row + 1) if ws.cell(row, 1).value][0]
        assert ws.cell(row, 8).value == "2026-09-20"   # Runs Out
        assert ws.cell(row, 9).value == "2026-09-16"   # Order By

    def test_stable_status_labels_no_relative_phrasing(self, run):
        client, run_id = run
        r = client.get(f"/api/export/{run_id}")
        wb = _wb(r)
        ws = wb["Material Forecast"]
        labels = {ws.cell(row, 4).value for row in range(5, ws.max_row + 1) if ws.cell(row, 1).value}
        assert labels == {"Already out", "Order now", "No action", "Order this week", "Stop ordering"}
        for lbl in labels:
            assert "day" not in lbl.lower() and "tomorrow" not in lbl.lower(), \
                f"FAIL: '{lbl}' looks like it carries relative-day phrasing"


class TestSharedFilterConsistency:
    def test_export_and_json_endpoint_agree_on_the_same_scope(self, run):
        """The whole point of _scoped_forecast(): /api/forecast (screen)
        and /api/export (file) must never disagree about which rows a
        given scope includes."""
        client, run_id = run
        params = {"service": "Electrical", "subcategory": "Wire"}
        json_rows = client.get(f"/api/forecast/{run_id}", params=params).json()
        wb = _wb(client.get(f"/api/export/{run_id}", params=params))
        ws = wb["Material Forecast"]
        xlsx_materials = sorted(ws.cell(row, 1).value for row in range(5, ws.max_row + 1) if ws.cell(row, 1).value)
        json_materials = sorted(r["material"] for r in json_rows)
        assert xlsx_materials == json_materials
