"""Measurement grid, Excel export and import preview -- through the real FastAPI router, same Thoth Fire case as test_measure_grid.py.

Rules under test:
  * the export is the grid: one row per BOQ line, one column per place, the hidden key row names each place by its id, the numbers are what Site
    Progress holds (and BOQ qty stays EMPTY where it is not known -- never a guess), locked boxes say why, boxes an item is not in are grey and empty;
  * a file of another service / project / not made by the grid is refused with a plain reason;
  * the import preview judges every box like typing would: locked boxes and boxes the item is not in are skipped, whole-number items refuse decimals,
    a number above the zone's plan is flagged (not silently cut), a blank box means "no change", sums typed in Excel are read, formulas that point
    at other cells are not; and the preview saves NOTHING."""
import io

import pytest
from openpyxl import Workbook, load_workbook

from backend import measurexl
from backend.test_measure_grid import site, grid, row, typed, measured, SVC, SLUG, P, PIPES  # noqa: F401  (site is a fixture)


def export(c, svc=SVC):
    r = c.get(f"{P}/measure-grid/{svc}/export")
    assert r.status_code == 200, r.text
    return r


def wb_of(r):
    return load_workbook(io.BytesIO(r.content))


def preview(c, data, svc=SVC, name="m.xlsx"):
    return c.post(f"{P}/measure-grid/{svc}/import-preview", files={"file": (name, data, "application/octet-stream")})


def save(wb):
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def find_row(ws, code):
    return next(r for r in range(measurexl.FIRST_ROW, ws.max_row + 1) if str(ws.cell(r, 1).value) == code)


def col_of(ws, pid):
    return next(c for c in range(measurexl.FIRST_COL, ws.max_column + 1) if ws.cell(measurexl.KEY_ROW, c).value == pid)


@pytest.fixture
def case(site):
    c, d, ids = site
    measured(c, ids[1], 37.0, "FACP-1", "Riser 2")
    measured(c, ids[1], 9.4, "Riser 2", "Zone 2 end")
    typed(c, ids[4], 58.0)
    return c, d, ids


# ------------------------------------------------------------------ export
def test_export_is_a_real_xlsx_download_with_a_clear_name(case):
    c, d, ids = case
    r = export(c)
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert 'filename="Measurement_thoth-mall_fire_' in r.headers["content-disposition"] and r.headers["content-disposition"].endswith('.xlsx"')
    assert r.content[:2] == b"PK"
    assert load_workbook(io.BytesIO(r.content)).sheetnames == ["Measurement", "Measurement sheet", "Summary", "_meta"]


def test_export_layout_keys_levels_zones_and_lines(case):
    c, d, ids = case
    wb = wb_of(export(c))
    ws = wb["Measurement"]
    assert ws.cell(measurexl.KEY_ROW, 1).value == "#KEY" and ws.row_dimensions[measurexl.KEY_ROW].hidden
    keys = [ws.cell(measurexl.KEY_ROW, measurexl.FIRST_COL + i).value for i in range(21)]
    assert keys == ids                                              # one column per place, in sheet order, named by its id
    assert [ws.cell(measurexl.HEAD2, measurexl.FIRST_COL + i).value for i in range(7)] == [f"Zone {i}" for i in range(1, 8)]
    assert ws.cell(measurexl.HEAD1, measurexl.FIRST_COL).value == "B3" and ws.cell(measurexl.HEAD1, measurexl.FIRST_COL + 7).value == "B2"
    assert {str(m) for m in ws.merged_cells.ranges} >= {"D4:J4", "K4:Q4", "R4:X4"}
    codes = [ws.cell(r, 1).value for r in range(measurexl.FIRST_ROW, measurexl.FIRST_ROW + 7)]
    assert codes == ["39a", "39b", "39c", "39d", "39e", "39f", "N.1"]
    assert ws.cell(find_row(ws, "39b"), 1).number_format == "@"      # codes are text, so Excel cannot turn 9.1.16 or 1.10 into a number
    assert ws.freeze_panes == "D6"


def test_export_numbers_are_the_grids_numbers(case):
    c, d, ids = case
    g = grid(c)
    ws = wb_of(export(c))["Measurement"]
    r = find_row(ws, "39b")
    assert ws.cell(r, col_of(ws, ids[1])).value == pytest.approx(46.4) and ws.cell(r, col_of(ws, ids[4])).value == 58
    assert ws.cell(r, col_of(ws, ids[0])).value == 0
    hdr = {ws.cell(measurexl.HEAD2, k).value: k for k in range(measurexl.FIRST_COL + 21, ws.max_column + 1)}
    assert ws.cell(r, hdr["BOQ qty"]).value == 1947.5 and ws.cell(r, hdr["Site plan"]).value == 1259.5
    assert str(ws.cell(r, hdr["To date"]).value).startswith("=SUM(D")        # a formula, so editing in Excel keeps the total honest
    assert "MAX(" in str(ws.cell(r, hdr["Pending"]).value)
    # every export cell equals the grid's own cell
    for rw in grid(c)["rows"]:
        rr = find_row(ws, rw["code"])
        for pid in ids:
            assert ws.cell(rr, col_of(ws, pid)).value == g["cells"][f"{rw['code']}|{pid}"]


def test_export_locked_boxes_say_why_and_unknown_boq_qty_is_empty_not_invented(case):
    c, d, ids = case
    ws = wb_of(export(c))["Measurement"]
    cell = ws.cell(find_row(ws, "39b"), col_of(ws, ids[1]))
    assert cell.comment is not None and "2 measurement entries" in cell.comment.text
    assert ws.cell(find_row(ws, "39b"), col_of(ws, ids[4])).comment is None         # a typed number is not locked
    h = wb_of(export(c, "HVAC"))["Measurement"]                                       # per-room qty: the whole-project BOQ qty is unknown
    hdr = {h.cell(measurexl.HEAD2, k).value: k for k in range(measurexl.FIRST_COL, h.max_column + 1)}
    assert h.cell(find_row(h, "H.1"), hdr["BOQ qty"]).value is None
    assert h.cell(find_row(h, "H.1"), hdr["Site plan"]).value == 210


def test_export_item_not_in_a_place_is_a_grey_empty_box(case):
    c, d, ids = case
    r = c.post(f"{P}/item-rooms", json={"service": SVC, "item_code": "39c", "rooms": ids[:3]})
    assert r.status_code == 200, r.text
    ws = wb_of(export(c))["Measurement"]
    rr = find_row(ws, "39c")
    assert ws.cell(rr, col_of(ws, ids[0])).value == 0
    out = ws.cell(rr, col_of(ws, ids[5]))
    assert out.value is None and out.fill.fgColor.rgb.endswith("EEECE6")


def test_export_summary_sheet_has_the_sign_off_box(case):
    c, d, ids = case
    sm = wb_of(export(c))["Summary"]
    text = " | ".join(str(x.value) for rw in sm.iter_rows() for x in rw if x.value is not None)
    assert "Thoth Mall" in text and "Fire" in text and "Prepared by" in text and "Client representative" in text and "Name / signature" in text


def test_export_unknown_service_is_404(case):
    c, d, ids = case
    assert c.get(f"{P}/measure-grid/Nope/export").status_code == 404


# ------------------------------------------------------------------ import preview
def test_importing_the_untouched_export_changes_nothing(case):
    c, d, ids = case
    r = preview(c, export(c).content)
    assert r.status_code == 200, r.text
    j = r.json()
    assert j["changes"] == [] and j["skipped"] == [] and j["unknown_items"] == [] and j["unknown_places"] == []
    assert j["counts"]["same"] == 7 * 21 and j["counts"]["change"] == 0


def test_preview_judges_every_box_like_typing_would(case):
    c, d, ids = case
    wb = wb_of(export(c))
    ws = wb["Measurement"]
    ws.cell(find_row(ws, "39c"), col_of(ws, ids[0])).value = 20                         # a normal change
    ws.cell(find_row(ws, "39d"), col_of(ws, ids[0])).value = "10+2.5"                   # a sum typed as text
    ws.cell(find_row(ws, "39b"), col_of(ws, ids[1])).value = 99                         # locked (measurement entries)
    ws.cell(find_row(ws, "39b"), col_of(ws, ids[2])).value = 100                        # above the zone's plan of 59.98
    ws.cell(find_row(ws, "N.1"), col_of(ws, ids[0])).value = 2.5                        # a counted item
    ws.cell(find_row(ws, "39a"), col_of(ws, ids[0])).value = "abc"
    ws.cell(find_row(ws, "39a"), col_of(ws, ids[1])).value = -5
    ws.cell(find_row(ws, "39e"), col_of(ws, ids[0])).value = None                       # blank = no change
    j = preview(c, save(wb)).json()
    ch = {(x["code"], x["room"]): x for x in j["changes"]}
    assert set(ch) == {("39c", ids[0]), ("39d", ids[0]), ("39b", ids[2])}
    assert ch[("39c", ids[0])]["new"] == 20 and ch[("39c", ids[0])]["cur"] == 0 and ch[("39c", ids[0])]["over_plan"] is False
    assert ch[("39d", ids[0])]["new"] == 12.5
    assert ch[("39b", ids[2])]["new"] == 100 and ch[("39b", ids[2])]["over_plan"] is True      # flagged, not cut
    assert ch[("39c", ids[0])]["place"] == "B3 ▸ Zone 1" and ch[("39c", ids[0])]["desc"] == "c) 80 mm dia"
    sk = {(x["code"], x["place"]): x for x in j["skipped"]}
    assert sk[("39b", "B3 ▸ Zone 2")]["kind"] == "locked" and "2 measurement entries" in sk[("39b", "B3 ▸ Zone 2")]["why"]
    assert sk[("N.1", "B3 ▸ Zone 1")]["kind"] == "bad" and "whole numbers" in sk[("N.1", "B3 ▸ Zone 1")]["why"]
    assert sk[("39a", "B3 ▸ Zone 1")]["kind"] == "bad" and "negative" in sk[("39a", "B3 ▸ Zone 2")]["why"]
    k = j["counts"]
    assert (k["change"], k["over_plan"], k["locked"], k["bad"], k["same"]) == (3, 1, 1, 3, 7 * 21 - 3 - 1 - 3 - 1)


def test_preview_saves_nothing(case):
    c, d, ids = case
    before = grid(c)["cells"]
    wb = wb_of(export(c))
    ws = wb["Measurement"]
    ws.cell(find_row(ws, "39c"), col_of(ws, ids[0])).value = 20
    assert preview(c, save(wb)).json()["counts"]["change"] == 1
    assert grid(c)["cells"] == before


def test_free_mode_item_is_not_flagged_above_plan(case):
    c, d, ids = case
    assert c.post(f"{P}/item-plan-off", json={"service": SVC, "item_code": "39b", "off": True}).status_code == 200
    wb = wb_of(export(c))
    ws = wb["Measurement"]
    ws.cell(find_row(ws, "39b"), col_of(ws, ids[2])).value = 100
    j = preview(c, save(wb)).json()
    assert [x["over_plan"] for x in j["changes"]] == [False]


def test_formulas_without_a_stored_result_are_read_only_when_they_are_plain_sums(case):
    c, d, ids = case
    wb = wb_of(export(c))                                          # saved by openpyxl: formulas have no cached value
    ws = wb["Measurement"]
    ws.cell(find_row(ws, "39c"), col_of(ws, ids[0])).value = "=10+2+30"
    ws.cell(find_row(ws, "39d"), col_of(ws, ids[0])).value = "=A6+2"
    j = preview(c, save(wb)).json()
    assert [(x["code"], x["new"]) for x in j["changes"]] == [("39c", 42.0)]
    assert [x for x in j["skipped"] if x["code"] == "39d"][0]["kind"] == "bad"


def test_unknown_lines_and_places_are_reported_not_guessed(case):
    c, d, ids = case
    wb = wb_of(export(c))
    ws = wb["Measurement"]
    ws.cell(measurexl.FIRST_ROW + 7, 1).value = "ZZ9"                                   # a line that is not in this BOQ
    ws.cell(measurexl.FIRST_ROW + 7, col_of(ws, ids[0])).value = 5
    ws.cell(measurexl.KEY_ROW, col_of(ws, ids[3])).value = "roo999"                      # a column whose place id is not in the project
    ws.cell(find_row(ws, "39c"), col_of(ws, "roo999")).value = 7
    ws.cell(find_row(ws, "39c"), col_of(ws, ids[0])).value = 3
    j = preview(c, save(wb)).json()
    assert j["unknown_items"] == ["ZZ9"] and j["unknown_places"] == ["roo999"]
    assert [(x["code"], x["room"]) for x in j["changes"]] == [("39c", ids[0])]


def test_a_box_the_item_is_not_in_is_skipped_with_its_reason(case):
    c, d, ids = case
    c.post(f"{P}/item-rooms", json={"service": SVC, "item_code": "39c", "rooms": ids[:3]})
    wb = wb_of(export(c))
    ws = wb["Measurement"]
    ws.cell(find_row(ws, "39c"), col_of(ws, ids[5])).value = 9
    j = preview(c, save(wb)).json()
    assert j["changes"] == [] and [x["kind"] for x in j["skipped"]] == ["not_in_place"]


# ------------------------------------------------------------------ files that are not ours
def test_file_of_another_service_is_refused(case):
    c, d, ids = case
    r = preview(c, export(c, "HVAC").content)
    assert r.status_code == 400 and "'HVAC'" in r.json()["detail"] and "'Fire'" in r.json()["detail"]
    assert preview(c, export(c).content, "HVAC").status_code == 400


def test_file_of_another_project_is_refused(case):
    c, d, ids = case
    wb = wb_of(export(c))
    wb["_meta"]["A2"] = "some-other-project"
    r = preview(c, save(wb))
    assert r.status_code == 400 and "different project" in r.json()["detail"]


@pytest.mark.parametrize("data,msg", [
    (b"", "empty"),
    (b"just some text", "not an Excel"),
    (b"PK\x03\x04garbage", "could not be read"),
])
def test_files_that_are_not_workbooks_get_a_plain_reason(case, data, msg):
    c, d, ids = case
    r = preview(c, data)
    assert r.status_code == 400 and msg in r.json()["detail"]


def test_a_workbook_without_the_key_row_is_refused(case):
    c, d, ids = case
    w = Workbook()
    w.active.title = "Measurement"
    r = preview(c, save(w))
    assert r.status_code == 400 and "key row" in r.json()["detail"]
    w2 = Workbook()
    r = preview(c, save(w2))
    assert r.status_code == 400 and 'no sheet named "Measurement"' in r.json()["detail"]


# ------------------------------------------------------------------ the reader itself
@pytest.mark.parametrize("v,count,want", [
    (None, False, ("blank",)), ("  ", False, ("blank",)), (12, False, ("ok", 12.0)), (12.5, False, ("ok", 12.5)),
    ("10+2+30", False, ("ok", 42.0)), ("=37+9.4", False, ("ok", 46.4)), ("2x3", False, ("ok", 6.0)), ("(1+2)*4", False, ("ok", 12.0)),
    ("1,200", False, ("ok", 1200.0)), (0, False, ("ok", 0.0)),
    (-1, False, ("bad", "a quantity cannot be negative")), ("5-2", False, ("bad", "a minus is not allowed")),
    (True, False, ("bad", "not a number")), (float("nan"), False, ("bad", "not a number")), (2e9, False, ("bad", "a very large number, check it")),
    (2.5, True, ("bad", "this item is counted in whole numbers")), (3.0, True, ("ok", 3.0)), ("10+", False, ("bad", "the sum is incomplete")),
])
def test_parse_qty(v, count, want):
    assert measurexl.parse_qty(v, count) == want


def test_parse_qty_refuses_cell_references_and_words():
    assert measurexl.parse_qty("=A1+2")[0] == "bad" and measurexl.parse_qty("twelve")[0] == "bad" and measurexl.parse_qty("__import__('os')")[0] == "bad"


def test_a_code_typed_as_a_plain_number_still_matches():
    w = Workbook()
    ws = w.active
    ws.title = "Measurement"
    ws.cell(measurexl.KEY_ROW, 1, "#KEY")
    ws.cell(measurexl.KEY_ROW, measurexl.FIRST_COL, "roo2")
    ws.cell(measurexl.FIRST_ROW, 1, 39.0)
    ws.cell(measurexl.FIRST_ROW, measurexl.FIRST_COL, 5)
    got = measurexl.read(save(w))
    assert got["rows"] == [{"row": measurexl.FIRST_ROW, "code": "39", "cells": {"roo2": 5}, "formulas": {}}] and got["service"] is None


# ------------------------------------------------------------------ the sum behind a box travels with the file
def test_a_box_with_a_sum_is_exported_as_that_formula_and_read_back_with_it(case):
    c, d, ids = case
    r = c.post(f"{P}/progress/rooms", json={"service": SVC, "item_code": "39c", "entries": [{"room": ids[0], "qty": 42, "expr": "=12+30"}]})
    assert r.status_code == 200, r.text
    wb = wb_of(export(c))
    ws = wb["Measurement"]
    assert ws.cell(find_row(ws, "39c"), col_of(ws, ids[0])).value == "=12+30"            # Excel shows 42 and keeps the pieces
    j = preview(c, save(wb)).json()
    assert j["changes"] == [] and j["counts"]["same"] == 7 * 21                               # a round trip changes nothing
    ws.cell(find_row(ws, "39d"), col_of(ws, ids[0])).value = "=10+2.5"
    j = preview(c, save(wb)).json()
    ch = [x for x in j["changes"] if x["code"] == "39d"][0]
    assert ch["new"] == 12.5 and ch["expr"] == "=10+2.5"                                      # the preview hands the sum on, so Apply keeps it
    ws.cell(find_row(ws, "39e"), col_of(ws, ids[0])).value = 7
    assert [x for x in preview(c, save(wb)).json()["changes"] if x["code"] == "39e"][0]["expr"] is None


# ------------------------------------------------------------------ only what was recorded is shown; the From-To lines come with the file
def hidden_rows(ws):
    return {r for r, dim in ws.row_dimensions.items() if dim.hidden and r >= measurexl.FIRST_ROW}


def hidden_cols(ws):
    return {c for c in range(measurexl.FIRST_COL, ws.max_column + 1) if ws.column_dimensions[openpyxl_letter(c)].hidden}


def openpyxl_letter(c):
    from openpyxl.utils import get_column_letter
    return get_column_letter(c)


def test_export_shows_only_the_lines_and_places_with_a_recorded_quantity(case):
    c, d, ids = case
    ws = wb_of(export(c))["Measurement"]
    shown = [str(ws.cell(r, 1).value) for r in range(measurexl.FIRST_ROW, ws.max_row + 1) if r not in hidden_rows(ws)]
    assert shown == ["39b"], shown                                              # only 39b has anything recorded
    assert len(hidden_rows(ws)) == len(PIPES) + 1 - 1                          # the other pipes and the Nos line are hidden, not removed
    live = {ws.cell(measurexl.KEY_ROW, cc).value for cc in range(measurexl.FIRST_COL, ws.max_column + 1)
            if ws.cell(measurexl.KEY_ROW, cc).value and cc not in hidden_cols(ws)}
    assert live == {ids[1], ids[4]}, live                                       # B3 > Zone 2 (entries) and Zone 5 (typed); every other place hidden
    assert ws.row_dimensions[find_row(ws, "39c")].outlineLevel == 1             # grouped: Excel's + brings them back
    assert "hidden" in str(wb_of(export(c))["Summary"]["A9"].value)


def test_a_hidden_line_is_still_in_the_file_and_imports_like_before(case):
    c, d, ids = case
    wb = wb_of(export(c))
    ws = wb["Measurement"]
    r39c = find_row(ws, "39c")
    assert r39c in hidden_rows(ws)
    ws.cell(r39c, col_of(ws, ids[0])).value = 12                                # a person unhides the line and types in it
    p = preview(c, save(wb))
    assert p.status_code == 200, p.text
    assert [(x["code"], x["room"], x["new"]) for x in p.json()["changes"]] == [("39c", ids[0], 12)]


def test_nothing_is_hidden_when_nothing_is_recorded_yet(site):
    c, d, ids = site
    ws = wb_of(export(c))["Measurement"]
    assert not hidden_rows(ws) and not hidden_cols(ws)                          # a fresh project: the whole sheet is there to fill in
    assert "A9" not in ws.parent["Summary"] or ws.parent["Summary"]["A9"].value is None
    assert "Measurement sheet" not in ws.parent.sheetnames


def sheet_blocks(wb):
    """The vertical sheet as {code: {"lines": [row values], "total": H of the total row}}, read like a person reads it."""
    ws = wb["Measurement sheet"]
    out, cur = {}, None
    for r in range(8, ws.max_row + 1):
        a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
        if a and not str(a).startswith(("Note", "Prepared")) and ws.cell(r, 3).value is not None and b and ws.cell(r, 1).font.bold:
            cur = out.setdefault(str(a), {"lines": [], "total": None, "first": r + 1})
        elif cur is not None and b and str(b).startswith("Total to date"):
            cur["total"], cur["last"] = ws.cell(r, 8).value, r - 1
            cur = None
        elif cur is not None and b:
            cur["lines"].append([ws.cell(r, k).value for k in range(2, 10)])
    return out


def test_the_vertical_sheet_lists_each_line_with_its_places_and_the_from_to_readings(case):
    c, d, ids = case
    typed(c, ids[5], 5.0)                                  # a typed quantity in a place without entries
    measured(c, ids[4], 1.5, "Riser 5", "Zone 5 end")      # Zone 5 holds a typed 58: now 59.5 = 58 typed + 1.5 entered
    wb = wb_of(export(c))
    blocks = sheet_blocks(wb)
    assert list(blocks) == ["39b"]                         # only the lines that have something recorded, in BOQ order
    ls = blocks["39b"]["lines"]
    assert [(x[0].strip(), x[2], x[3], x[6]) for x in ls[:2]] == [("B3 — Zone 2", "FACP-1", "Riser 2", 37), ("B3 — Zone 2", "Riser 2", "Zone 2 end", 9.4)]
    assert [(x[0].strip(), x[2], x[3], x[6]) for x in ls[2:]] == [
        ("B3 — Zone 5", "Recorded without From-To detail", None, 58), ("B3 — Zone 5", "Riser 5", "Zone 5 end", 1.5),
        ("B3 — Zone 6", "Recorded without From-To detail", None, 5)]                      # a typed length says it has no From-To behind it
    assert "B3 — Zone 7" not in [x[0].strip() for x in ls]  # nothing recorded there: not listed


def test_the_vertical_sheet_ties_to_the_grid_line_by_line(case):
    c, d, ids = case
    typed(c, ids[5], 5.0)
    typed(c, ids[0], 12.0, "39c")
    typed(c, ids[7], 3.0, "N.1")                          # a Nos line: places and quantities only, no From-To
    wb = wb_of(export(c))
    ws, blocks = wb["Measurement"], sheet_blocks(wb)
    vs = wb["Measurement sheet"]
    for code, b in blocks.items():
        r = find_row(ws, code)
        grid_total = sum((ws.cell(r, cc).value or 0) for cc in range(measurexl.FIRST_COL, measurexl.FIRST_COL + 21))
        listed = sum(vs.cell(rr, 8).value or 0 for rr in range(b["first"], b["last"] + 1))
        assert listed == pytest.approx(grid_total, abs=1e-6), (code, listed, grid_total)
        assert b["total"] == f"=SUM(H{b['first']}:H{b['last']})"                                    # a live formula, so Excel adds it up itself


def test_a_whole_level_total_is_one_line_never_split_into_zones(case):
    c, d, ids = case
    r = c.post(f"{P}/progress/area-total", json={"service": SVC, "item_code": "39d", "rooms": ids[7:14], "qty": 120})
    assert r.status_code == 200, r.text
    ls = sheet_blocks(wb_of(export(c)))["39d"]["lines"]
    assert len(ls) == 1 and ls[0][0].strip() == "B2 — whole level / floor" and ls[0][6] == pytest.approx(120, abs=1e-3)
    assert "Total for the whole area" in ls[0][7]


def test_a_cancelled_entry_is_not_on_the_vertical_sheet(case):
    c, d, ids = case
    runs = c.get(f"{P}/progress/runs", params={"service": SVC, "item_code": "39b"}).json()["rooms"][ids[1]]["runs"]
    r = c.post(f"{P}/progress/runs/cancel", json={"service": SVC, "item_code": "39b", "room": ids[1], "run_id": runs[0]["id"], "reason": "wrong riser"})
    assert r.status_code == 200, r.text
    ls = sheet_blocks(wb_of(export(c)))["39b"]["lines"]
    assert [x[2] for x in ls if x[0].strip() == "B3 — Zone 2"] == ["Riser 2"]


def test_drum_readings_are_in_their_own_columns(case):
    c, d, ids = case
    r = c.post(f"{P}/progress/runs", json={"service": SVC, "item_code": "39c", "scope": "zone", "room": ids[9], "from": "Room WIFI", "to": "ELV Room", "start": 112, "end": 78})
    assert r.status_code == 200, r.text
    ls = sheet_blocks(wb_of(export(c)))["39c"]["lines"]
    assert [(x[2], x[3], x[4], x[5], x[6]) for x in ls] == [("Room WIFI", "ELV Room", 112, 78, 34)]            # From, To, START, END, length = |112 - 78|


def test_the_vertical_sheet_is_read_only_the_import_ignores_it(case):
    c, d, ids = case
    wb = wb_of(export(c))
    vs = wb["Measurement sheet"]
    vs.cell(9, 8).value = 9999                                                    # someone types over a line of the client's sheet
    p = preview(c, save(wb))
    assert p.status_code == 200, p.text
    assert p.json()["changes"] == []
