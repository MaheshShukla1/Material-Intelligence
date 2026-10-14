"""Supporting Measurement Sheet + whole-area totals.

Reported need (billing person): "how will the installation backup be prepared?" -- the client certifies WHERE each installed
quantity was fitted. Site Progress already records a quantity per zone/room, but a site often knows only "B3 total = 420 Rmt".
Rules under test:
  * an area total is recorded as ONE total, spread over the area's rooms exactly like any other quantity (every % and report
    keeps working), and stays a total only while no room of the area has been changed since;
  * the Measurement Sheet is the LAST sheet of the invoice workbook, in the bill's own format, quantities only;
  * it always ties to the bill: previous / this / total of every item equal Annexure B-BOQ;
  * what the records cannot explain is ONE visible Adjustment line; what was billed earlier without locations is said, not guessed.
Harness: test_billing.py's Fire BOQ (GI pipe 100 Rmt over 4 rooms = 25 each; floors "1" and "2", rooms a/b)."""
import io
import json
import subprocess
import sys
import tempfile
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl
import pytest

from backend import billing, locqty, siteprogress as sp
from backend.test_billing import env, prepared, SLUG, line                      # noqa: F401
from backend.test_billing_challan import new_ra, bill, B

P = f"/api/siteprogress/{SLUG}"
PIPE = "GI pipe 100mm"


def plan(client, code="3"):
    return client.get(f"{P}/item-room-plan", params={"service": "Fire", "item_code": code}).json()


def area_total(client, rooms, qty, code="3"):
    return client.post(f"{P}/progress/area-total", json={"service": "Fire", "item_code": code, "rooms": rooms, "qty": qty})


def set_rooms(client, entries, code="3"):
    r = client.post(f"{P}/progress/rooms", json={"service": "Fire", "item_code": code, "entries": entries})
    assert r.status_code == 200, r.text


def meas_book(client, ras, extra="boq"):
    r = client.get(f"{B}/download", params={"ras": ras, "sheets": f"{extra},meas" if extra else "meas"})
    assert r.status_code == 200, r.text
    return openpyxl.load_workbook(io.BytesIO(r.content))


def recalc(client, ras, extra="boq"):
    """The same download, with the draft's formulas calculated (LibreOffice), so the figures can be read."""
    r = client.get(f"{B}/download", params={"ras": ras, "sheets": f"{extra},meas"})
    with tempfile.TemporaryDirectory() as t:
        src = Path(t) / "in.xlsx"
        src.write_bytes(r.content)
        out = Path(t) / "out"
        out.mkdir()
        subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(out), str(src)], check=True,
                       capture_output=True, timeout=120)
        return openpyxl.load_workbook(out / "in.xlsx", data_only=True)


def meas_ws(wb):
    return wb[wb.sheetnames[-1]]


def item_block(ws, desc):
    """(header_row, [location rows], adjustment_row|None, total_row) of one item of the Measurement Sheet."""
    h = next(r for r in range(9, ws.max_row + 1) if ws.cell(r, 2).value == desc)
    r = h + 1
    rows = []
    while not str(ws.cell(r, 2).value or "").startswith("Total qty in bill"):
        rows.append(r)
        r += 1
    adj = next((x for x in rows if str(ws.cell(x, 2).value or "").strip().startswith("Less / add")), None)
    return h, [x for x in rows if x != adj], adj, r


def vals(ws, r):
    return [ws.cell(r, c).value for c in (5, 6, 7)]


def boq_this(ws_boq, desc):
    hdr = next(r for r in range(1, 20) if "THIS BILL QTY" in [ws_boq.cell(r, c).value for c in range(1, 25)])
    col = next(c for c in range(1, 25) if ws_boq.cell(hdr, c).value == "THIS BILL QTY")
    row = next(r for r in range(8, ws_boq.max_row + 1) if ws_boq.cell(r, 2).value == desc)
    return ws_boq.cell(row, col).value, [ws_boq.cell(hdr, c).value for c in range(1, 25)], row, col, hdr


def lock(client, ra):
    r = client.post(f"{B}/ra/{ra}/lock")
    assert r.status_code == 200, r.text


# =================================================================== recording a whole-area total
def test_area_total_is_spread_over_the_area_and_remembered_as_one_total(env):
    client, d, _ = prepared(env)
    r = area_total(client, ["roo2", "roo3"], 30)                      # floor 1 = 30 Rmt of its 50 planned
    assert r.status_code == 200, r.text
    pl = plan(client)
    assert {k: v["frac"] for k, v in pl["rooms"].items()} == {"roo2": 0.6, "roo3": 0.6, "roo5": 0.2, "roo6": 0.2}
    assert pl["pools"] == [{"area": "1", "rooms": ["roo2", "roo3"], "qty": 30.0}]
    loc = sp.item_location_qty(d, "Fire", "3")
    assert loc["pools"]["1"]["q"] == 30 and loc["rest"] == 10 and loc["sum"] == 40
    assert loc["sum"] == pytest.approx(billing._install_qty(d, "Fire")["3"])             # the same quantity the bill uses


def test_an_awkward_total_comes_back_exactly(env):
    """Reported on the real Thoth item 13.2.1 (533.3333 Rmt per zone, 7 zones): 1500 typed must read back 1500, not 1500.0011."""
    client, d, _ = prepared(env)
    (d / "planned.json").write_text(json.dumps({"Fire": {"3": 3733.3331}}))
    assert area_total(client, ["roo2", "roo3"], 1500 / 3.0).status_code == 200
    assert plan(client)["pools"][0]["qty"] == pytest.approx(1500 / 3.0, abs=1e-4)
    assert sp.item_location_qty(d, "Fire", "3")["pools"]["1"]["q"] == pytest.approx(1500 / 3.0, abs=1e-4)


def test_a_total_at_or_above_the_plan_counts_the_area_as_done(env):
    client, d, _ = prepared(env)
    assert area_total(client, ["roo2", "roo3"], 80).status_code == 200
    assert {plan(client)["rooms"][r]["frac"] for r in ("roo2", "roo3")} == {1.0}
    assert plan(client)["pools"][0]["qty"] == 50


def test_zero_clears_the_area_total(env):
    client, d, _ = prepared(env)
    area_total(client, ["roo2", "roo3"], 30)
    assert area_total(client, ["roo2", "roo3"], 0).status_code == 200
    assert plan(client)["pools"] == []


def test_a_room_changed_afterwards_ends_the_total_for_that_area(env):
    client, d, _ = prepared(env)
    area_total(client, ["roo2", "roo3"], 30)
    set_rooms(client, [{"room": "roo2", "qty": 5}])                    # the engineer now types a zone figure
    assert plan(client)["pools"] == []
    loc = sp.item_location_qty(d, "Fire", "3")
    assert loc["pools"] == {} and set(loc["rooms"]) == {"roo2", "roo3"}
    assert loc["rooms"]["roo2"] == 5 and loc["rooms"]["roo3"] == 15


def test_the_overall_slider_ends_the_total_too(env):
    client, d, _ = prepared(env)
    area_total(client, ["roo2", "roo3"], 30)
    progress_all = client.post(f"{P}/progress/item", json={"service": "Fire", "item_code": "3", "frac": 0.5})
    assert progress_all.status_code == 200
    assert plan(client)["pools"] == []
    assert sp.item_location_qty(d, "Fire", "3")["pools"] == {}


def test_a_total_for_rooms_of_two_floors_is_refused_and_nothing_is_written(env):
    client, d, _ = prepared(env)
    before = (d / "item_progress.json").read_text()
    assert area_total(client, ["roo2", "roo5"], 30).status_code == 400
    assert area_total(client, ["roo2", "nope"], 30).status_code == 400
    assert area_total(client, ["roo2", "roo3"], -1).status_code == 400
    assert area_total(client, ["roo2", "roo3"], "x").status_code == 400
    assert area_total(client, [], 5).status_code == 400
    assert (d / "item_progress.json").read_text() == before and not (d / "item_pools.json").exists()


def test_a_room_with_no_plan_cannot_take_a_share_of_a_total(env):
    client, d, _ = prepared(env)
    r = client.post(f"{P}/progress/area-total", json={"service": "Fire", "item_code": "NOPE", "rooms": ["roo2"], "qty": 5})
    assert r.status_code == 400


# =================================================================== the sheet in the workbook
def test_the_measurement_sheet_is_the_last_sheet_and_ties_to_the_bill(env):
    client, d, _ = prepared(env)
    area_total(client, ["roo2", "roo3"], 30)                           # floor 1: whole-area total
    set_rooms(client, [{"room": "roo5", "qty": 10}], code="3")         # floor 2: one zone typed, the other stays overall (5)
    ra = new_ra(client, "install")
    lock(client, ra)
    wb = meas_book(client, ra)
    assert wb.sheetnames[-1] == "Annexure C-Measurement" and "Annexure B-BOQ" in wb.sheetnames
    ws = meas_ws(wb)
    h, rows, adj, tot = item_block(ws, PIPE)
    labels = [ws.cell(r, 2).value.strip() for r in rows]
    assert labels[0] == "1 — whole floor" and "a" in labels[1] and "Remaining areas" in labels[2]
    assert ws.cell(rows[0], 8).value == "Total for the whole area, not split zone-wise"
    assert [vals(ws, rows[0])] == [[0, 30, 30]]                         # first RA: previous 0
    assert vals(ws, tot) == [0, 45, 45] and adj is None                # 30 + 10 + 5
    thisq, *_ = boq_this(wb["Annexure B-BOQ"], PIPE)
    assert thisq == 45                                                 # == Annexure B-BOQ
    assert ws.cell(h, 4).value == 100                                  # BOQ qty


def test_a_draft_sheet_has_working_formulas_and_the_same_numbers(env):
    client, d, _ = prepared(env)
    area_total(client, ["roo2", "roo3"], 30)
    ra = new_ra(client, "install")
    wb = recalc(client, ra)
    ws = meas_ws(wb)
    assert ws["B5"].value.startswith("DRAFT")
    h, rows, adj, tot = item_block(ws, PIPE)
    assert vals(ws, tot) == [0, 40, 40] and sum(ws.cell(r, 7).value for r in rows) == 40
    assert boq_this(wb["Annexure B-BOQ"], PIPE)[0] == 40


def test_next_ra_shows_this_ra_by_location(env):
    client, d, _ = prepared(env)
    area_total(client, ["roo2", "roo3"], 30)
    ra1 = new_ra(client, "install")
    lock(client, ra1)
    area_total(client, ["roo2", "roo3"], 40)                           # floor 1 grew by 10
    set_rooms(client, [{"room": "roo6", "qty": 15}])                   # floor 2 zone b went 5 -> 15
    ra2 = new_ra(client, "install")
    lock(client, ra2)
    wb = meas_book(client, ra2)
    ws = meas_ws(wb)
    h, rows, adj, tot = item_block(ws, PIPE)
    by = {ws.cell(r, 2).value.strip(): vals(ws, r) for r in rows}
    assert by["1 — whole floor"] == [30, 10, 40]
    assert vals(ws, tot) == [40, 20, 60]                               # previous bill qty 40 -> this 20
    assert boq_this(wb["Annexure B-BOQ"], PIPE)[0] == 20 and adj is None
    # the earlier RA keeps its own figures after the later one exists
    ws1 = meas_ws(meas_book(client, ra1))
    assert vals(ws1, item_block(ws1, PIPE)[3]) == [0, 40, 40]


def test_an_item_billed_before_without_locations_says_so_and_still_ties(env):
    client, d, _ = prepared(env)
    ra1 = new_ra(client, "install")
    lock(client, ra1)
    st = billing._load(d)                                              # an RA locked before locations were kept
    for r in st["ras"]:
        for row in r["snap"]["lines"]:
            (row.get("install") or {}).pop("loc", None)
    billing._save(d, st)
    area_total(client, ["roo2", "roo3"], 30)                           # 20 billed earlier -> 40 now
    ra2 = new_ra(client, "install")
    lock(client, ra2)
    ws = meas_ws(meas_book(client, ra2))
    h, rows, adj, tot = item_block(ws, PIPE)
    assert "Location of earlier RAs not recorded" in ws.cell(h, 8).value
    assert vals(ws, tot) == [20, 20, 40]
    assert all(ws.cell(r, 5).value is None and ws.cell(r, 6).value is None for r in rows)       # nothing split that was never recorded
    assert sum(ws.cell(r, 7).value for r in rows) == 40
    # and the old RA itself, opened as a sheet, says its places are not recorded
    ws1 = meas_ws(meas_book(client, ra1))
    h1, rows1, _, tot1 = item_block(ws1, PIPE)
    assert ws1.cell(rows1[0], 2).value.strip() == "Location not recorded" and vals(ws1, tot1) == [0, 20, 20]


def test_a_place_recorded_differently_from_the_last_bill_is_shown_as_one_total(env):
    client, d, _ = prepared(env)
    area_total(client, ["roo2", "roo3"], 30)
    ra1 = new_ra(client, "install")
    lock(client, ra1)
    set_rooms(client, [{"room": "roo2", "qty": 20}, {"room": "roo3", "qty": 15}])     # same floor, now zone by zone
    ra2 = new_ra(client, "install")
    lock(client, ra2)
    ws = meas_ws(meas_book(client, ra2))
    h, rows, adj, tot = item_block(ws, PIPE)
    row = next(r for r in rows if ws.cell(r, 2).value.strip() == "1")
    assert vals(ws, row) == [30, 5, 35] and "differently" in ws.cell(row, 8).value
    assert vals(ws, tot)[2] == 35 + 10                                   # floor 2 stays overall (10)


def test_quantity_capped_at_the_boq_is_one_adjustment_line_and_the_total_ties(env):
    client, d, _ = prepared(env)
    (d / "planned.json").write_text(json.dumps({"Fire": {"3": 150}}))      # site plan 150 Rmt, BOQ says 100
    area_total(client, ["roo2", "roo3"], 75)
    area_total(client, ["roo5", "roo6"], 75)                              # 150 recorded, only 100 can be billed
    ra = new_ra(client, "install")
    lock(client, ra)
    wb = meas_book(client, ra)
    ws = meas_ws(wb)
    h, rows, adj, tot = item_block(ws, PIPE)
    assert adj is not None and "capped at the BOQ quantity" in ws.cell(adj, 2).value
    assert vals(ws, adj) == [0, -50, -50]
    assert vals(ws, tot) == [0, 100, 100] and boq_this(wb["Annexure B-BOQ"], PIPE)[0] == 100


def test_a_quantity_typed_by_hand_is_an_adjustment_not_a_location(env):
    client, d, _ = prepared(env)
    ra = new_ra(client, "install")
    r = client.post(f"{B}/ra/{ra}/line", json={"item_code": "3", "stage": "install", "qty": 30, "reason": "site measurement"})
    assert r.status_code == 200, r.text
    lock(client, ra)
    wb = meas_book(client, ra)
    ws = meas_ws(wb)
    h, rows, adj, tot = item_block(ws, PIPE)
    assert adj is not None and "by hand" in ws.cell(adj, 2).value
    assert vals(ws, tot) == [0, 30, 30] and boq_this(wb["Annexure B-BOQ"], PIPE)[0] == 30


def test_only_items_billed_in_this_ra_are_on_the_sheet_and_no_rate_appears(env):
    client, d, _ = prepared(env)
    ra = new_ra(client, "install")
    lock(client, ra)
    ws = meas_ws(meas_book(client, ra))
    descs = [ws.cell(r, 2).value for r in range(9, ws.max_row + 1)]
    assert PIPE in descs and "Butterfly valve 150mm" in descs and "Diesel engine pump set" not in descs
    heads = [ws.cell(8, c).value for c in range(1, 9)]
    assert heads == ["Sr. No", "Description / Location", "Unit", "BOQ QTY", "PREVIOUS BILL QTY", "THIS BILL QTY", "TOTAL QTY", "Remarks"]
    assert not any("rate" in str(c.value or "").lower() or "amount" in str(c.value or "").lower() for row in ws.iter_rows(min_row=8, max_row=8) for c in row)
    assert any(str(ws.cell(r, 1).value or "").startswith("Prepared") or str(ws.cell(r, 4).value or "").startswith("Checked") for r in range(9, ws.max_row + 1))


def test_without_the_option_the_workbook_is_exactly_as_before(env):
    client, d, _ = prepared(env)
    ra = new_ra(client, "install")
    r = client.get(f"{B}/download", params={"ras": ra})                       # default sheets
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert not any("Measurement" in n for n in wb.sheetnames)
    r = client.get(f"{B}/download", params={"ras": ra, "sheets": "invoice,summary,boq"})
    assert not any("Measurement" in n for n in openpyxl.load_workbook(io.BytesIO(r.content)).sheetnames)


def test_a_supply_ra_has_no_measurement_sheet(env):
    client, d, _ = prepared(env)
    ra = new_ra(client, "supply")
    r = client.get(f"{B}/download", params={"ras": ra, "sheets": "boq,meas"})
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert not any("Measurement" in n for n in wb.sheetnames)
    r = client.get(f"{B}/download", params={"ras": ra, "sheets": "meas"})
    assert r.status_code == 400


def test_all_in_one_workbook_puts_every_measurement_sheet_at_the_end(env):
    client, d, _ = prepared(env)
    ra1 = new_ra(client, "install")
    lock(client, ra1)
    r = client.get(f"{B}/download", params={"ras": ra1, "mode": "one", "sheets": "invoice,boq,meas"})
    names = openpyxl.load_workbook(io.BytesIO(r.content)).sheetnames
    assert names[-1].endswith("MEASUREMENT") and sum(n.endswith("MEASUREMENT") for n in names) == 1


def test_coverage_names_what_was_billed_earlier_without_locations(env):
    client, d, _ = prepared(env)
    ra1 = new_ra(client, "install")
    lock(client, ra1)
    st = billing._load(d)
    for r in st["ras"]:
        for row in r["snap"]["lines"]:
            (row.get("install") or {}).pop("loc", None)
    billing._save(d, st)
    area_total(client, ["roo2", "roo3"], 30)
    ra2 = new_ra(client, "install")
    cov = client.get(f"{B}/measurement/coverage", params={"ras": f"{ra1},{ra2}"}).json()["ras"]
    c1, c2 = cov
    assert c1["legacy_ra"] is True and c1["applies"] and c1["items"] == 2
    assert {(e["desc"], e["qty"]) for e in c2["earlier"]} == {(PIPE, 20.0)}              # the valve has nothing new in RA 2, so it is not on that sheet
    assert c2["legacy_ra"] is False
    sup = new_ra(client, "supply")
    assert client.get(f"{B}/measurement/coverage", params={"ras": sup}).json()["ras"][0]["applies"] is False


def test_the_billing_screens_never_carry_the_location_records(env):
    client, d, _ = prepared(env)
    area_total(client, ["roo2", "roo3"], 30)
    ra = new_ra(client, "install")
    lock(client, ra)
    assert "loc" not in json.dumps(client.get(f"{B}/ra/{ra}").json()["lines"])
    assert any("loc" in (r.get("install") or {}) for r in billing._load(d)["ras"][0]["snap"]["lines"])


# =================================================================== lines_for (pure): hotel rooms, every project shape
INFO = {"name": {f"r{i}": str(300 + i) for i in range(1, 11)}, "loc": {f"r{i}": "Floor 3" for i in range(1, 11)},
        "order": {f"r{i}": i for i in range(1, 11)}, "word": "floor", "leaf": "room", "collapse": True}


def now(rooms=None, pools=None, rest=0.0):
    rooms = rooms or {}
    pools = pools or {}
    return {"rooms": rooms, "pools": {a: {"q": q, "o": 1} for a, q in pools.items()}, "rest": rest,
            "sum": sum(rooms.values()) + sum(pools.values()) + rest}


def test_ten_equal_rooms_become_one_line_and_a_different_one_stands_alone():
    rooms = {f"r{i}": 30.0 for i in range(1, 11)}
    rooms["r5"] = 28.0
    ls = locqty.lines_for(now(rooms), None, "none", INFO)
    assert [l["label"] for l in ls] == ["Floor 3 — 301–304 (4 rooms)", "Floor 3 — 305", "Floor 3 — 306–310 (5 rooms)"]
    assert [l["total"] for l in ls] == [120, 28, 150]


def test_mall_zones_are_never_merged():
    info = {**INFO, "collapse": False, "leaf": "zone", "word": "level"}
    ls = locqty.lines_for(now({f"r{i}": 10.0 for i in range(1, 6)}), None, "none", info)
    assert len(ls) == 5 and ls[0]["label"] == "Floor 3 — 301"


def test_two_equal_rooms_are_not_merged():
    ls = locqty.lines_for(now({"r1": 10.0, "r2": 10.0}), None, "none", INFO)
    assert len(ls) == 2


def test_overall_only_reads_as_one_whole_project_line():
    ls = locqty.lines_for(now(rest=20.0), None, "none", INFO)
    assert ls == [{"label": "Whole project (recorded as overall total)", "total": 20.0, "remark": "Recorded as an overall total, not by location", "prev": 0.0, "this": 20.0}]


def test_pool_and_new_rooms_in_the_same_area_both_show():
    prev = now(pools={"Floor 3": 100.0})
    cur = {**now({"r1": 5.0}, pools={"Floor 3": 120.0}), }
    ls = locqty.lines_for(cur, prev, "loc", INFO)
    assert [(l["label"], l["prev"], l["this"], l["total"]) for l in ls] == [("Floor 3 — whole floor", 100.0, 20.0, 120.0), ("Floor 3 — 301", 0.0, 5.0, 5.0)]


def test_nothing_recorded_gives_no_lines():
    assert locqty.lines_for(now(), None, "none", INFO) == []


# =================================================================== every project shape names its places
@pytest.mark.parametrize("kind,word,leaf", [("hotel", "floor", "room"), ("mall", "level", "zone"), ("hospital", "area", "room"), ("custom", "area", "room")])
def test_location_info_is_generic_across_project_types(env, kind, word, leaf):
    client, d, _ = env
    st = json.loads((d / "structure.json").read_text())
    st["kind"] = kind
    (d / "structure.json").write_text(json.dumps(st))
    info = sp.location_info(d)
    assert info["word"] == word and info["leaf"] == leaf and info["collapse"] is (leaf == "room")
    assert set(info["name"]) == set(info["loc"]) == set(info["order"]) and len(info["name"]) == 4
