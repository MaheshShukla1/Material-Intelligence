"""ProjectBase Billing Request sheet: Connect once per service, then every RA's quantities go into the "Now" column of
that same file. Real exports (Electrical / HVAC, 11-column "BOQ Details", 4 rows per BOQ line) test the reader; the
endpoints run through the real routers on a synthetic export built exactly like them, for the Fire test project."""
import io
import os
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl
import pytest

from backend import projectbase as pb
from backend.test_billing import env, SLUG  # noqa: F401  (env is a fixture)
from backend.test_billing_challan import cs, ch, new_ra, bill, B

UP = "/root/.claude/uploads/cc641dec-fdd2-5679-9e93-268c1cf0e51e/"
REAL = {"electrical": UP + "4e0aaadc-Billing_Request-undefined.xlsx", "hvac": UP + "fcaa3bc4-Billing_Request-VEIPL-GUJ_BR_26-27_003.xlsx"}
real = pytest.mark.skipif(not all(os.path.exists(p) for p in REAL.values()), reason="real ProjectBase exports not on this machine")

BL = {"supply": "75% on Delivery of material at site", "install": "15% on Installation of material",
      "testing": "5% on Testing & commissioning", "handover": "5% on Handing over of the system"}


# --------------------------------------------------------------------------- real exports
@real
def test_real_exports_read_four_rows_per_line_by_the_words_of_the_billing_line():
    for key, n in (("electrical", 339), ("hvac", 551)):
        p = pb.read(open(REAL[key], "rb").read())
        assert len(p["groups"]) == n and not any(g["odd"] for g in p["groups"].values())
        for g in p["groups"].values():
            assert set(g["rows"]) == {"supply", "install", "testing", "handover"}
            assert g["rows"]["install"] == g["rows"]["supply"] + 1 and g["rows"]["handover"] == g["rows"]["supply"] + 3
            assert g["pct"] == {"supply": 75.0, "install": 15.0, "testing": 5.0, "handover": 5.0}


@real
def test_real_fill_writes_only_now_on_the_carried_rows_and_nothing_else_changes():
    raw = open(REAL["hvac"], "rb").read()
    p = pb.read(raw)
    refs = list(p["groups"])[:3]
    ref_of = {f"c{i}": r for i, r in enumerate(refs)}
    out, rep = pb.fill(raw, p, ref_of, {"c0": {"install": 4.5}, "c1": {"install": 2}}, ["install"])
    a, b = openpyxl.load_workbook(io.BytesIO(raw)).active, openpyxl.load_workbook(io.BytesIO(out)).active
    now = p["cols"]["Now"]
    changed = [(c.row, c.column) for row in a.iter_rows() for c in row if c.value != b.cell(c.row, c.column).value]
    want = {(p["groups"][refs[0]]["rows"]["install"], now), (p["groups"][refs[1]]["rows"]["install"], now)}
    # exactly the two Installation/Now cells differ -- unless the export already held a value there that is now cleared
    assert want <= set(changed)
    rows_ = {g["rows"][s] for g in p["groups"].values() for s in ("supply", "install")}
    assert all(col == now and r in rows_ for r, col in changed)                       # only Now cells of Supply / Installation rows
    assert b.cell(p["groups"][refs[0]]["rows"]["install"], now).value == 4.5
    assert b.cell(p["groups"][refs[1]]["rows"]["install"], now).value == 2
    assert [x[:3] for x in rep["filled"]] == [("c0", refs[0], "install"), ("c1", refs[1], "install")]


@real
def test_real_internal_ref_is_enough_to_match_every_line_and_a_foreign_file_matches_none():
    p = pb.read(open(REAL["hvac"], "rb").read())
    lines = [{"item_code": f"x{i}", "internal_ref": r, "desc": "", "unit": "Nos", "boq_qty": 1} for i, r in enumerate(p["groups"])]
    assert pb.match(lines, p)["matched"] == len(lines)
    other = pb.read(open(REAL["electrical"], "rb").read())
    m = pb.match([{"item_code": "z", "internal_ref": "99999", "desc": "", "unit": "Nos", "boq_qty": 1}], other)
    assert m["matched"] == 0 and m["missing"][0]["why"].startswith("its Internal Ref")


# --------------------------------------------------------------------------- a synthetic export of the test Fire BOQ
# internal refs of the test BOQ: item 1 -> "2", item 2 -> "3", item 3 -> "4", item 4 (rate only) -> "5", item 5 -> "8"
FIRE = [("2", "1", "Diesel engine pump set", "Nos", 2), ("3", "2", "Butterfly valve 150mm", "Nos", 4),
        ("4", "3", "GI pipe 100mm", "Rmt", 100), ("5", "4", "Rate only fitting", "Nos", "RO"), ("8", "5", "Hydrant valve", "Nos", 10)]


def export(path=None, now=None, drop=(), swap=None, pcts=(75, 15, 5, 5), client_of=None, extra=()):
    """A Billing Request export. now = {(internal_ref, stage): value} already typed in the Now column."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BOQ Details"
    ws.append(pb.HEADERS)
    ws.append(["PRJ/1", None, "A", "Fire Pump & Accessories", None, None, None, None, None, None, None])        # a heading
    for ref, cl, text, unit, qty in list(FIRE) + list(extra):
        if ref in drop:
            continue
        ref = (swap or {}).get(ref, ref)
        for st, p in zip(("supply", "install", "testing", "handover"), pcts):
            bl = BL[st].replace("75%", f"{p}%") if st == "supply" else BL[st]
            ws.append(["PRJ/1", ref, (client_of or {}).get(ref, cl), text, "Terms", bl, unit, qty, None, None,
                       (now or {}).get((ref, st))])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def upload(client, raw, base="Fire", name="Billing_Request.xlsx"):
    return client.post(f"{B}/projectbase/{base}/inspect", files={"file": (name, raw)})


def connect(client, raw, base="Fire", confirm=None):
    r = upload(client, raw, base)
    assert r.status_code == 200, r.text
    c = client.post(f"{B}/projectbase/{base}/connect", json={"confirm": confirm or {}, "name": "BR.xlsx"})
    assert c.status_code == 200, c.text
    return r.json()


def sheet_now(data):
    ws = openpyxl.load_workbook(io.BytesIO(data)).active
    p = pb.read(data)
    return {(ref, st): ws.cell(r, p["cols"]["Now"]).value for ref, g in p["groups"].items() for st, r in g["rows"].items()}


def pbxlsx(client, ra, **kw):
    r = client.get(f"{B}/download", params={"ras": ra, "pb": 1, "invoice": 0, **kw})
    assert r.status_code == 200, r.text
    return r


def test_inspect_reads_the_file_and_reports_what_matches_and_what_setup_disagrees_on(env):
    client, _, _ = env
    cs(client, sup=70, ins=30)
    r = upload(client, export())
    assert r.status_code == 200, r.text
    j = r.json()
    assert (j["project"], j["lines"], j["matched"], j["total"]) == ("PRJ/1", 5, 4, 4)         # the rate-only line is not billable
    assert j["units"] == [{"unit": "Fire", "package": "Fire", "total": 4, "matched": 4}] or j["units"][0]["matched"] == 4
    assert j["pct"] == {"supply": 75.0, "install": 15.0, "testing": 5.0, "handover": 5.0}
    assert any("Supply 75%" in w and "70%" in w for w in j["warnings"])                        # file 75, MI setup 70: said, not hidden
    assert client.get(B).json()["pb"] == {}                                                    # nothing is connected by inspect alone


def test_a_file_of_another_service_is_refused_and_not_kept(env):
    client, d, _ = env
    cs(client)
    wrong = export(swap={"2": "900", "3": "901", "4": "902", "5": "903", "8": "904"},
                   client_of={"900": "z1", "901": "z2", "902": "z3", "903": "z4", "904": "z5"})
    r = upload(client, wrong)
    assert r.status_code == 400 and "None of Fire" in r.json()["detail"]
    assert client.post(f"{B}/projectbase/Fire/connect", json={}).status_code == 409
    assert upload(client, b"not an excel file").status_code == 400
    wb = openpyxl.Workbook()
    wb.active.append(["a", "b"])
    buf = io.BytesIO()
    wb.save(buf)
    r = upload(client, buf.getvalue())
    assert r.status_code == 400 and "Internal Ref No" in r.json()["detail"]
    assert upload(client, export(), base="Nope").status_code == 404


def test_connect_keeps_the_file_and_disconnect_removes_it(env):
    client, d, _ = env
    cs(client)
    connect(client, export())
    meta = client.get(B).json()["pb"]["Fire"]
    assert meta["project"] == "PRJ/1" and meta["lines"] == 5 and meta["name"] == "BR.xlsx"
    assert (d / "projectbase" / "Fire.xlsx").exists() and not (d / "projectbase" / "Fire.pending.xlsx").exists()
    assert client.delete(f"{B}/projectbase/Fire").status_code == 200
    assert client.get(B).json()["pb"] == {} and not (d / "projectbase" / "Fire.xlsx").exists()


def test_the_supply_ra_fills_only_the_supply_rows_of_the_lines_it_bills(env):
    client, _, _ = env
    cs(client, sup=75, ins=15)
    ch(client, "614", "2026-07-04", [("1", 2), ("2", 3)])
    ra = new_ra(client)
    connect(client, export())
    got = sheet_now(pbxlsx(client, ra).content)
    assert got[("2", "supply")] == 2 and got[("3", "supply")] == 3
    assert [k for k, v in got.items() if v not in (None, "")] == [("2", "supply"), ("3", "supply")]         # nothing else, no install/testing/handover
    prev = client.get(f"{B}/projectbase/preview", params={"ras": ra}).json()["ras"][0]
    assert prev["connected"] and prev["filled"] == 2 and prev["not_found"] == [] and prev["cleared"] == [] and prev["left"] == []
    assert prev["expect_total"] == pytest.approx(100000 * 2 * .75 + 5000 * 3 * .75) == pytest.approx(bill(client, ra)["this_total"])


def test_old_values_in_the_now_column_are_blanked_and_reported_not_imported_into_this_bill(env):
    client, _, _ = env
    cs(client, sup=75, ins=15)
    ch(client, "614", "2026-07-04", [("1", 2)])
    ra = new_ra(client)
    # the export still holds a placeholder on another line's supply row (typed to create the BR), an old install value,
    # and a quantity typed on a Testing row
    connect(client, export(now={("8", "supply"): 1, ("2", "install"): 7, ("4", "testing"): 9}))
    prev = client.get(f"{B}/projectbase/preview", params={"ras": ra}).json()["ras"][0]
    assert {(c["ref"], c["stage"]) for c in prev["cleared"]} == {("8", "supply"), ("2", "install")}    # a Supply RA must not import an old install qty either
    assert [(x["ref"], x["stage"]) for x in prev["left"]] == [("4", "testing")]
    got = sheet_now(pbxlsx(client, ra).content)
    assert got[("8", "supply")] is None and got[("2", "supply")] == 2
    assert got[("2", "install")] is None and got[("4", "testing")] == 9                               # Testing / Handover are never touched, only reported


def test_a_line_that_cannot_be_matched_is_reported_with_its_quantity_and_never_guessed(env):
    client, _, _ = env
    cs(client, sup=75, ins=15)
    ch(client, "614", "2026-07-04", [("1", 2), ("5", 4)])
    ra = new_ra(client)
    connect(client, export(drop=("8",)))                       # hydrant valve is not in this file
    prev = client.get(f"{B}/projectbase/preview", params={"ras": ra}).json()["ras"][0]
    assert prev["filled"] == 1 and [(x["item_code"], x["qty"]) for x in prev["not_found"]] == [("5", 4)]
    assert prev["expect_total"] == pytest.approx(100000 * 2 * .75)                                  # only what really went in
    assert prev["ra_total"] == pytest.approx(100000 * 2 * .75 + 3000 * 4 * .75)


def test_a_near_match_is_offered_for_confirmation_only_when_exactly_one_line_fits(env):
    client, _, _ = env
    cs(client)
    raw = export(swap={"8": "8001"})                            # hydrant valve exported under a different Internal Ref
    j = upload(client, raw).json()
    assert [(a["item_code"], a["file"]["ref"]) for a in j["ask"]] == [("5", "8001")]
    assert j["matched"] == 3
    # two identical candidates -> nothing is offered
    raw2 = export(swap={"8": "8001"}, extra=[("8002", "5", "Hydrant valve", "Nos", 10)])
    assert upload(client, raw2).json()["ask"] == []
    # confirmed -> used, and remembered for the next upload of this service
    upload(client, raw)
    assert client.post(f"{B}/projectbase/Fire/connect", json={"confirm": {"5": "8001"}}).status_code == 200
    ch(client, "614", "2026-07-04", [("5", 2)])
    ra = new_ra(client)
    assert sheet_now(pbxlsx(client, ra).content)[("8001", "supply")] == 2
    again = upload(client, raw).json()
    assert again["ask"] == [] and again["matched"] == 4


def test_negative_and_manual_lines_are_warned_about_not_silently_dropped(env):
    client, _, _ = env
    cs(client, sup=75, ins=15)
    ch(client, "614", "2026-07-04", [("1", 2)])
    ra1 = new_ra(client)
    client.post(f"{B}/ra/{ra1}/lock")
    ch(client, "615", "2026-07-05", [("2", 1)])
    ra2 = new_ra(client)
    connect(client, export())
    # edit RA 2 so line 1 goes below what RA 1 billed
    assert client.post(f"{B}/ra/{ra2}/line", json={"item_code": "1", "stage": "supply", "qty": -1, "reason": "returned"}).status_code == 200
    prev = client.get(f"{B}/projectbase/preview", params={"ras": ra2}).json()["ras"][0]
    assert prev["negative"] == ["1"]


def test_download_names_the_file_and_can_be_combined_with_the_invoice(env):
    client, _, _ = env
    cs(client)
    ch(client, "614", "2026-07-04", [("1", 1)])
    ra = new_ra(client)
    assert client.get(f"{B}/download", params={"ras": ra, "pb": 1, "invoice": 0}).status_code == 409        # not connected yet
    connect(client, export())
    r = client.get(f"{B}/download", params={"ras": ra, "pb": 1, "invoice": 0})
    assert "ProjectBase" in r.headers["content-disposition"] and r.headers["content-disposition"].endswith('.xlsx"')
    r2 = client.get(f"{B}/download", params={"ras": ra, "pb": 1, "invoice": 1})
    assert r2.status_code == 200
    names = zipfile.ZipFile(io.BytesIO(r2.content)).namelist()
    assert len(names) == 2 and sum("ProjectBase" in n for n in names) == 1


def test_the_total_check_says_match_and_otherwise_names_the_line_that_explains_the_gap(env):
    client, _, _ = env
    cs(client, sup=75, ins=15)
    ch(client, "614", "2026-07-04", [("1", 2), ("2", 3), ("3", 69)])         # 69 m of pipe went into the file; ProjectBase showed 66
    ra = new_ra(client)
    connect(client, export())
    mine = client.get(f"{B}/projectbase/preview", params={"ras": ra}).json()["ras"][0]["expect_total"]
    ok = client.post(f"{B}/ra/{ra}/pb-check", json={"subtotal": mine}).json()
    assert ok["match"] and ok["diff"] == 0 and ok["suspects"] == []
    step = 800 * 75 / 100
    bad = client.post(f"{B}/ra/{ra}/pb-check", json={"subtotal": mine - 3 * step}).json()
    assert not bad["match"] and bad["diff"] == pytest.approx(3 * step)
    top = bad["suspects"][0]
    assert (top["item_code"], top["sent"], top["would_be"], top["difference"]) == ("3", 69, 66, 3)
    assert len(bad["suspects"]) <= 3
    assert client.post(f"{B}/ra/{ra}/pb-check", json={"subtotal": "abc"}).status_code == 400


def test_pbras_limits_the_projectbase_files_to_the_connected_services_while_the_invoice_covers_every_tick(env):
    client, _, _ = env
    cs(client)
    cs(client, svc="HVAC")
    ch(client, "614", "2026-07-04", [("1", 1)])
    ch(client, "H1", "2026-07-04", [("1", 1)], svc="HVAC")
    fire, hvac = new_ra(client), new_ra(client, svc="HVAC")
    connect(client, export())                                   # only Fire is connected
    assert client.get(f"{B}/download", params={"ras": f"{fire},{hvac}", "pb": 1}).status_code == 409
    r = client.get(f"{B}/download", params={"ras": f"{fire},{hvac}", "pb": 1, "pbras": fire})
    assert r.status_code == 200
    names = zipfile.ZipFile(io.BytesIO(r.content)).namelist()
    assert len(names) == 3 and sum("ProjectBase" in n for n in names) == 1       # 2 invoices + 1 ProjectBase sheet
    only = client.get(f"{B}/download", params={"pb": 1, "invoice": 0, "pbras": fire})
    assert only.status_code == 200 and "ProjectBase" in only.headers["content-disposition"]
