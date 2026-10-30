"""Sheet-sync / upload housekeeping (disk and memory), reproducing what happened
on the live server: every open tab re-synced the same 17 MB Google Sheet every
10 minutes, each sync left its raw .xlsx behind forever and added one more run.

Run from the repo root:
    pytest backend/test_sync_cleanup.py -v
"""
import datetime as dt
import json
import shutil
import threading
import time

import openpyxl
import pytest
from fastapi.testclient import TestClient

from backend import api

LINK = "https://docs.google.com/spreadsheets/d/e/FAKE/pub?output=xlsx"
OTHER_LINK = "https://docs.google.com/spreadsheets/d/e/OTHER/pub?output=xlsx"


def make_register(path, days=12, with_ppe=True):
    """A small but real-shaped site stock register: dates over IN/OUT/BALANCE
    triples, material rows, plus a PPE sheet (so the PPE parse is exercised)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Site Stock"
    start = dt.datetime(2026, 9, 1)
    hdr = ["Sr No", "Groups", "Material Description", "Unit", "Opening Stock"]
    for d in range(days):
        ws.cell(row=1, column=6 + 3 * d, value=start + dt.timedelta(days=d))
        hdr += ["IN", "OUT", "BALANCE"]
    for c, v in enumerate(hdr, 1):
        ws.cell(row=2, column=c, value=v)
    mats = [("Electrical", "2.5 SQMM FLEXIBLE WIRE BLUE", "COIL", 100, 3),
            ("Electrical", "4.0 SQMM FLEXIBLE WIRE BLACK", "COIL", 80, 2),
            ("Plumbing", "50MM UPVC PIPE", "MTR", 300, 9)]
    for i, (grp, name, unit, opening, out) in enumerate(mats):
        r = 3 + i
        ws.cell(row=r, column=1, value=i + 1)
        ws.cell(row=r, column=2, value=grp)
        ws.cell(row=r, column=3, value=name)
        ws.cell(row=r, column=4, value=unit)
        ws.cell(row=r, column=5, value=opening)
        bal = opening
        for d in range(days):
            bal -= out
            ws.cell(row=r, column=7 + 3 * d, value=out)
            ws.cell(row=r, column=8 + 3 * d, value=bal)
    if with_ppe:
        p = wb.create_sheet("PPE")
        p.append(["Date", "Name", "Shoes Number", "Shoes", "Helmet", "Jacket", "Contractor Name", "Signature"])
        p.append([dt.datetime(2026, 9, 5), "Test Worker One", 8, 1, 1, 1, "Contractor A", "TW1"])
        p.append([dt.datetime(2026, 9, 6), "Test Worker Two", 9, 1, 1, None, "Contractor B", "TW2"])
    wb.save(path)


@pytest.fixture
def env(tmp_path, monkeypatch):
    runs, uploads = tmp_path / "runs", tmp_path / "uploads"
    runs.mkdir()
    uploads.mkdir()
    monkeypatch.setattr(api, "RUNS", runs)
    monkeypatch.setattr(api, "UPLOADS", uploads)
    api.load_run.cache_clear()
    src = tmp_path / "register.xlsx"
    make_register(src)
    fetches = []

    def fake_fetch(link, dest):
        fetches.append(link)
        shutil.copy(src, dest)

    monkeypatch.setattr(api, "_fetch_sheet_xlsx", fake_fetch)
    client = TestClient(api.app)
    yield client, runs, uploads, src, fetches
    api.load_run.cache_clear()


def sync(client, link=LINK, **kw):
    r = client.post("/api/sync-sheet", data={"link": link, **kw})
    assert r.status_code == 200, r.text
    return r.json()


def test_the_synthetic_register_really_builds_a_run(env):
    client, runs, uploads, src, _ = env
    j = sync(client)
    assert j["meta"]["has_ppe"] is True
    assert (runs / j["run_id"] / "forecast.parquet").exists()
    assert j["summary"]


# ---- raw .xlsx is not kept --------------------------------------------------
def test_a_sheet_sync_leaves_no_raw_file_behind_but_the_run_is_complete(env):
    client, runs, uploads, *_ = env
    j = sync(client)
    assert list(uploads.iterdir()) == []
    d = runs / j["run_id"]
    assert {p.name for p in d.iterdir()} >= {"forecast.parquet", "daily.parquet", "meta.json", "ppe.json"}
    assert client.get(f"/api/run/{j['run_id']}").status_code == 200
    assert client.get(f"/api/ppe/{j['run_id']}").json()["records"]       # PPE read BEFORE the file went


def test_a_manual_upload_leaves_no_raw_file_behind(env):
    client, runs, uploads, src, _ = env
    with open(src, "rb") as fh:
        r = client.post("/api/upload", files={"file": ("reg.xlsx", fh)}, data={"project": "Upload Proj"})
    assert r.status_code == 200, r.text
    assert list(uploads.iterdir()) == []
    assert client.get(f"/api/run/{r.json()['run_id']}").status_code == 200


def test_a_bad_file_is_refused_and_leaves_nothing_behind(env):
    client, runs, uploads, *_ = env
    r = client.post("/api/upload", files={"file": ("bad.xlsx", b"not an excel file")})
    assert r.status_code == 422
    assert list(uploads.iterdir()) == []
    assert list(runs.iterdir()) == []


def test_a_failure_after_parsing_still_removes_the_raw_file(env, monkeypatch):
    client, runs, uploads, *_ = env

    def boom(*a, **k):
        raise RuntimeError("forecast blew up")

    monkeypatch.setattr(api.engine, "forecast", boom)
    with pytest.raises(RuntimeError):
        TestClient(api.app, raise_server_exceptions=True).post(
            "/api/sync-sheet", data={"link": LINK})
    assert list(uploads.iterdir()) == []


# ---- one sync per link ------------------------------------------------------
def test_silent_auto_sync_inside_the_window_reuses_the_last_run_and_does_not_download(env):
    client, runs, uploads, src, fetches = env
    first = sync(client, lead_time=7, project="Hyatt Test")
    assert len(fetches) == 1
    again = sync(client, lead_time=7, project="Hyatt Test", auto="1")      # a second tab's timer
    assert again["reused"] is True
    assert again["run_id"] == first["run_id"]
    assert again["meta"]["project"] == "Hyatt Test" and again["summary"]
    assert len(fetches) == 1                                               # no second download
    assert len(list(runs.iterdir())) == 1


def test_the_sync_button_always_pulls_fresh_data(env):
    client, runs, uploads, src, fetches = env
    first = sync(client, project="Hyatt Test")
    second = sync(client, project="Hyatt Test")                            # no auto flag = button
    assert second.get("reused") is None
    assert second["run_id"] != first["run_id"]
    assert len(fetches) == 2


def test_auto_sync_after_the_window_pulls_fresh_data(env, monkeypatch):
    client, runs, uploads, src, fetches = env
    first = sync(client, project="Hyatt Test")
    meta_path = runs / first["run_id"] / "meta.json"
    m = json.loads(meta_path.read_text())
    m["created"] = (dt.datetime.now() - dt.timedelta(seconds=api.AUTO_SYNC_REUSE_SECONDS + 60)).isoformat(timespec="seconds")
    meta_path.write_text(json.dumps(m))
    second = sync(client, project="Hyatt Test", auto="1")
    assert second.get("reused") is None and second["run_id"] != first["run_id"]
    assert len(fetches) == 2


def test_a_different_lead_time_or_asof_is_never_served_from_the_reused_run(env):
    client, runs, uploads, src, fetches = env
    sync(client, lead_time=7, project="Hyatt Test")
    other_lead = sync(client, lead_time=14, project="Hyatt Test", auto="1")
    assert other_lead.get("reused") is None
    with_asof = sync(client, lead_time=14, project="Hyatt Test", auto="1", asof="2026-09-05")
    assert with_asof.get("reused") is None
    assert len(fetches) == 3


def test_tabs_firing_at_the_same_moment_do_one_download_not_three(env, monkeypatch):
    client, runs, uploads, src, fetches = env
    real = api._fetch_sheet_xlsx

    def slow_fetch(link, dest):
        time.sleep(0.3)                                                    # the real download takes seconds
        real(link, dest)

    monkeypatch.setattr(api, "_fetch_sheet_xlsx", slow_fetch)
    results, errs = [], []

    def tab():
        try:
            results.append(sync(TestClient(api.app), lead_time=7, project="Hyatt Test", auto="1"))
        except Exception as e:                                             # pragma: no cover
            errs.append(e)

    threads = [threading.Thread(target=tab) for _ in range(3)]
    [t.start() for t in threads]
    [t.join() for t in threads]
    assert not errs
    assert len(fetches) == 1
    assert len({r["run_id"] for r in results}) == 1
    assert len(list(runs.iterdir())) == 1


def test_different_links_do_not_block_or_reuse_each_other(env):
    client, runs, uploads, src, fetches = env
    a = sync(client, link=LINK, project="A")
    b = sync(client, link=OTHER_LINK, project="B", auto="1")
    assert b.get("reused") is None and b["run_id"] != a["run_id"]


# ---- retention --------------------------------------------------------------
def test_only_the_newest_runs_of_a_sheet_are_kept(env, monkeypatch):
    client, runs, uploads, src, fetches = env
    monkeypatch.setattr(api, "SYNC_RUNS_KEEP", 3)
    ids = []
    for i in range(6):
        ids.append(sync(client, project="Hyatt Test")["run_id"])
        time.sleep(1.05)                                                   # meta.created has 1 s resolution
    left = {p.name for p in runs.iterdir()}
    assert left == set(ids[-3:])
    assert client.get(f"/api/run/{ids[-1]}").status_code == 200
    assert client.get(f"/api/run/{ids[0]}").status_code == 404
    assert list(uploads.iterdir()) == []


def test_retention_never_touches_manual_uploads_or_other_sheets(env, monkeypatch):
    client, runs, uploads, src, fetches = env
    monkeypatch.setattr(api, "SYNC_RUNS_KEEP", 1)
    with open(src, "rb") as fh:
        manual = client.post("/api/upload", files={"file": ("reg.xlsx", fh)}, data={"project": "Hyatt Test"}).json()["run_id"]
    other = sync(client, link=OTHER_LINK, project="Other Sheet")["run_id"]
    time.sleep(1.05)
    sync(client, project="Hyatt Test")
    time.sleep(1.05)
    last = sync(client, project="Hyatt Test")["run_id"]
    left = {p.name for p in runs.iterdir()}
    assert manual in left and other in left and last in left
    assert len(left) == 3


def test_shortage_history_never_reads_a_run_folder_so_pruning_cannot_change_it():
    """shortage_log.json keeps (run_id, service, item_code) as a key and stores the
    numbers themselves. If shortage_history ever starts opening run folders,
    deleting old runs would silently change the case-study numbers - fail then."""
    import inspect
    from backend import shortage_history as sh
    src = inspect.getsource(sh)
    for needle in ("RUNS", "open(", "Path(", "read_parquet", "read_text", "os.listdir"):
        assert needle not in src, needle



# ---- memory: the per-run forecast cache must stay small ---------------------------------
def test_forecast_frames_of_old_runs_do_not_pile_up_in_memory(env, monkeypatch):
    """Live server (Railway, billed per GB-minute): every sync adds a run, and load_run() kept up to 128 of
    their forecast DataFrames for the life of the process (SYNC_RUNS_KEEP = 50 runs per sheet stay on disk,
    so nothing cleared them). Only the newest run of a project is ever read, so a handful is enough."""
    client, runs, uploads, src, fetches = env
    monkeypatch.setattr(api, "SYNC_RUNS_KEEP", 50)
    for i in range(12):
        rid = sync(client, project="Hyatt Test")["run_id"]
        assert client.get(f"/api/run/{rid}").status_code == 200            # the screen opens each new run as it appears
        time.sleep(1.05)                                                   # meta.created has 1 s resolution
    assert len(list(runs.iterdir())) == 12                                  # 12 runs on disk, as the live server keeps up to 50
    assert api.load_run.cache_info().currsize <= 8, api.load_run.cache_info()
