"""Issue 4: the Material Forecast export was missing a UOM column entirely.
forecast.parquet already carries a real per-row 'unit' field (engine.py's
own forecast() output, confirmed) -- this only adds it to the export,
mapping engine.py's own "-" missing-unit placeholder (build_daily(), so a
genuinely unit-less register row doesn't vanish from groupby) to a clear
"\u2014" rather than a blank cell that reads as broken."""
import json

import pandas as pd
import pytest
from fastapi.testclient import TestClient

from backend import api


@pytest.fixture
def run(tmp_path, monkeypatch):
    monkeypatch.setattr(api, "RUNS", tmp_path)
    run_id = "run1"
    d = tmp_path / run_id
    d.mkdir()
    # Reproduces the real reported case: some materials have a real unit
    # (Wire -> Mtr), others have engine.py's own "-" placeholder (fittings
    # whose register rows never had a Unit value) or a genuinely missing/NaN
    # unit -- all three must be handled, never a silent blank that looks
    # like the column itself is broken.
    df = pd.DataFrame([
        {"service": "Electrical", "material": "1.5 SQMM WIRE BLACK", "subcategory": "Wire",
         "unit": "MTR", "stock": 0, "total_consumed": 22410, "status": "STOCKED_OUT",
         "confidence": "HIGH", "order_by": None, "exhaust_date": None},
        {"service": "Fire & HVAC", "material": "25MM BENDING SPRING", "subcategory": "Pipe",
         "unit": "-", "stock": 5, "total_consumed": 10, "status": "GREEN",
         "confidence": "HIGH", "order_by": None, "exhaust_date": None},
        {"service": "Fire & HVAC", "material": "32X25MM REDUCER COUPLING", "subcategory": "Pipe",
         "unit": None, "stock": 3, "total_consumed": 8, "status": "GREEN",
         "confidence": "HIGH", "order_by": None, "exhaust_date": None},
    ])
    df.to_parquet(d / "forecast.parquet")
    (d / "meta.json").write_text(json.dumps({"project": "Hyatt Hotel"}))
    return TestClient(api.app), run_id


def _load(resp):
    import io
    import openpyxl
    return openpyxl.load_workbook(io.BytesIO(resp.content))


class TestUomColumn:
    def test_header_includes_uom_right_after_material(self, run):
        client, run_id = run
        r = client.get(f"/api/export/{run_id}")
        wb = _load(r)
        ws = wb["Material Forecast"]
        headers = [ws.cell(4, c).value for c in range(1, 12)]
        assert headers[0] == "Material"
        assert headers[1] == "UOM"

    def test_real_unit_shows_correctly(self, run):
        client, run_id = run
        r = client.get(f"/api/export/{run_id}")
        wb = _load(r)
        ws = wb["Material Forecast"]
        rows = {ws.cell(row, 1).value: ws.cell(row, 2).value for row in range(5, ws.max_row + 1) if ws.cell(row, 1).value}
        assert rows["1.5 SQMM WIRE BLACK"] == "MTR"

    def test_placeholder_dash_shows_as_em_dash_not_blank(self, run):
        """engine.py's own '-' missing-unit placeholder -- must read as
        honestly unknown ("\u2014"), never a blank cell that looks broken."""
        client, run_id = run
        r = client.get(f"/api/export/{run_id}")
        wb = _load(r)
        ws = wb["Material Forecast"]
        rows = {ws.cell(row, 1).value: ws.cell(row, 2).value for row in range(5, ws.max_row + 1) if ws.cell(row, 1).value}
        assert rows["25MM BENDING SPRING"] == "\u2014"
        assert rows["25MM BENDING SPRING"] != "", "FAIL: must not be a blank cell"

    def test_genuinely_missing_unit_also_shows_as_em_dash(self, run):
        client, run_id = run
        r = client.get(f"/api/export/{run_id}")
        wb = _load(r)
        ws = wb["Material Forecast"]
        rows = {ws.cell(row, 1).value: ws.cell(row, 2).value for row in range(5, ws.max_row + 1) if ws.cell(row, 1).value}
        assert rows["32X25MM REDUCER COUPLING"] == "\u2014"

    def test_all_other_columns_still_correct_after_the_shift(self, run):
        """The real risk of inserting a column: every column index after it
        shifts by one -- confirms Service/Category/Status/Stock etc. still
        land in the right place, and the bold-red 'Already out' styling
        still targets the real Status column (now index 5, not 4)."""
        client, run_id = run
        r = client.get(f"/api/export/{run_id}")
        wb = _load(r)
        ws = wb["Material Forecast"]
        wire_row = next(row for row in range(5, ws.max_row + 1) if ws.cell(row, 1).value == "1.5 SQMM WIRE BLACK")
        assert ws.cell(wire_row, 3).value == "Electrical"     # Service
        assert ws.cell(wire_row, 4).value == "Wire"           # Category
        assert ws.cell(wire_row, 5).value == "Already out"    # Status
        assert ws.cell(wire_row, 6).value == 0                # Stock
        assert ws.cell(wire_row, 7).value == 22410             # Total Received
        # bold-red styling must follow Status to its new column (5)
        assert ws.cell(wire_row, 5).font.bold is True
        assert ws.cell(wire_row, 5).font.color.rgb == "00B23A2E"

    def test_no_filters_still_returns_every_material_unchanged(self, run):
        """Backward compatibility: adding a column must not change WHICH
        rows come back for the existing no-filter caller."""
        client, run_id = run
        r = client.get(f"/api/export/{run_id}")
        wb = _load(r)
        ws = wb["Material Forecast"]
        materials = [ws.cell(row, 1).value for row in range(5, ws.max_row + 1) if ws.cell(row, 1).value]
        assert len(materials) == 3
